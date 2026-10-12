"""
生成请假审批表 Word 文档

直接拿模板（backend/templates/请假审批表.docx）改里面的表格单元格。
模板的结构是 12 行 × 6 列的表格，前 8 行是学生填的内容，
后面是班导师意见、学院意见、销假记录（这三块系统不管，管理员下载后打印出来手签）。

模板放在 backend/templates/ 里（进 git、也随镜像走），不再用 data/downloads/ 那份：
data/ 是挂载进来的数据卷，跑一次 scripts/setup_resources.py 就会被素材目录里的
同名文件覆盖回去，改在这里的排版会无声丢失。「资料下载」里那份是给学生下载手填的
空白表（本次没跟着改），与这份生成模板各管各的。
"""

from pathlib import Path
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

from ..config import BASE_DIR
from ..models import Student, LeaveRequest

# 模板路径：仓库里那份生成模板（data/downloads 那份是给学生手填的空白表，不是它）
TEMPLATE_PATH = BASE_DIR / "templates" / "请假审批表.docx"

# 表格里的单元格坐标（行, 列），按模板的实际结构来
# 模板是 6 列，很多单元格是合并的，实际写入时找对位置就行
#
# 行0: 学生姓名(0,0) 值(0,1) 性别(0,2) 值(0,3) 班级(0,4) 值(0,5)
# 行1: 学号(1,0) 值(1,1) 电话(1,2) 值(1,3) 邮箱(1,4) 值(1,5)
# 行2: 身份证号(2,0) 值(2,1~2,2 合并) QQ/微信(2,3) 值(2,4~2,5 合并)
# 行3: 家庭地址(3,0) 值(3,1~3,2 合并) 家人联系电话(3,3) 值(3,4~3,5 合并)
# 行4: 请假时间(4,0) 值(4,1~4,5 合并，其实是 5 个重复列)
# 行5: 请假事由(5,0) 值(5,1~5,5 合并)
# 行6: 出行方式(6,0) 值(6,1~6,2) 行程(6,3) 值(6,4~6,5)
# 行7: 去向地址(7,0) 值(7,1~7,2) 去向联系电话(7,3) 值(7,4~7,5)


def _set_cell_text(cell, text: str, font_size: float = 11) -> None:
    """把单元格内容替换成指定文字（保留第一个段落，其余删掉）"""
    # 清空原有内容
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    if cell.paragraphs:
        p = cell.paragraphs[0]
        # 清空原有 run
        for run in p.runs:
            run.text = ""
        # 写入新内容
        if p.runs:
            p.runs[0].text = text
            p.runs[0].font.size = Pt(font_size)
        else:
            run = p.add_run(text)
            run.font.size = Pt(font_size)


def _adaptive_font_size(text: str, base_size: float = 11, min_size: float = 8) -> float:
    """根据文本长度自适应字号，长文本缩小以避免换页

    以中文每个字约占 2 个英文字符宽度来估算。
    每超过阈值就缩小 1pt，直到 min_size。
    """
    if not text:
        return base_size
    # 估算「等效字符数」：中文算 2，英文/数字算 1
    length = sum(2 if ord(c) > 127 else 1 for c in text)
    # 单元格大约能容纳 40 个等效字符不换行（按 11pt 估算）
    # 每多 15 个字符缩小 1pt
    shrink = max(0, (length - 40) // 15)
    return max(min_size, base_size - shrink)


def generate_leave_docx(student: Student, leave: LeaveRequest) -> bytes:
    """根据学生信息和请假申请，生成填好的请假审批表 Word

    返回 docx 文件的二进制内容（可以直接写响应）。
    """
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"请假审批表模板不存在：{TEMPLATE_PATH}")

    doc = Document(str(TEMPLATE_PATH))
    table = doc.tables[0]

    # —— 第 0 行：姓名 / 性别 / 班级 ——
    _set_cell_text(table.cell(0, 1), student.name or "")
    _set_cell_text(table.cell(0, 3), student.gender or "")
    _set_cell_text(table.cell(0, 5), student.class_name or "")

    # —— 第 1 行：学号 / 电话 / 邮箱 ——
    _set_cell_text(table.cell(1, 1), student.student_no or "")
    _set_cell_text(table.cell(1, 3), student.phone or "")
    _set_cell_text(table.cell(1, 5), student.email or "")

    # —— 第 2 行：身份证号 / QQ微信 ——
    _set_cell_text(table.cell(2, 1), student.id_card or "")
    _set_cell_text(table.cell(2, 4), student.qq or "")

    # —— 第 3 行：家庭地址 / 家人联系电话 ——
    # 地址经常很长，用自适应字号，避免撑高单元格导致换页
    addr_size = _adaptive_font_size(student.family_address or "", base_size=10)
    _set_cell_text(table.cell(3, 1), student.family_address or "", font_size=addr_size)
    _set_cell_text(table.cell(3, 4), student.family_phone or "")

    # —— 第 4 行：请假时间（精确到「时」，与模板「年 月 日 时」一致）——
    start = leave.start_at
    end = leave.end_at
    # 保险起见转 datetime（有时可能是字符串）
    if isinstance(start, str):
        start = datetime.fromisoformat(start)
    if isinstance(end, str):
        end = datetime.fromisoformat(end)
    # 跨几个日历天（含起止当天），同日不同时也算 1 天
    days = (end.date() - start.date()).days + 1
    time_text = (
        f"{start.year}年{start.month}月{start.day}日 {start.hour}时 至 "
        f"{end.year}年{end.month}月{end.day}日 {end.hour}时（共 {days} 天）"
    )
    _set_cell_text(table.cell(4, 1), time_text)

    # —— 第 5 行：请假事由 ——
    reason_text = f"[{leave.leave_type}] {leave.reason}" if leave.leave_type else leave.reason
    # 事由也可能很长，同样自适应字号
    reason_size = _adaptive_font_size(reason_text or "", base_size=10.5)
    _set_cell_text(table.cell(5, 1), reason_text or "", font_size=reason_size)

    # —— 第 6 行：出行方式 / 行程 ——
    _set_cell_text(table.cell(6, 1), leave.travel_mode or "")
    _set_cell_text(table.cell(6, 4), leave.travel_detail or "")

    # —— 第 7 行：去向地址 / 去向联系电话 ——
    # 去向地址也可能很长，自适应字号
    dest_size = _adaptive_font_size(leave.destination or "", base_size=10)
    _set_cell_text(table.cell(7, 1), leave.destination or "", font_size=dest_size)
    _set_cell_text(table.cell(7, 4), leave.destination_phone or "")

    # 班导师意见、学院意见、销假记录 这三块留空，管理员下载后手签

    # 写到内存里返回
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
