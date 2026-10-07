"""成绩的共用逻辑：Excel 解析、预览令牌、写入与组装

分三层，接口只负责权限和中文错误：

- **解析**（parse_grade_excel）：只读不写。认不出的人和填错的格子都**不阻断导入**，
  全攒在预览里给班委看 —— 「上传后先确认」这个流程存在的意义就在这里。
- **令牌**（save_preview / load_preview / drop_preview）：解析结果落临时文件，
  确认时按 token 取回，不用重新上传一遍 Excel。
- **写入与组装**（import_grade / set_cell / list_items / detail / my_grades）。

本模块不做角色判断、不抛 HTTPException：解析类错误抛 ValueError 由接口转 400，
业务规则类错误由接口自己判（同 api/students.py、api/polls.py 的分工）。
"""

import io
import json
import re
import tempfile
import time
import uuid
from pathlib import Path

import openpyxl
from sqlalchemy import func
from sqlalchemy.orm import Session

from ..models import GradeItem, GradeSubject, GradeTest, Student, User
from ..models.grade import (
    DEFAULT_FULL_SCORE,
    MARK_VALUES,
    MAX_SCORE,
    PASS_RATIO,
    PASS_VALUES,
)

# 允许的标记，拼成给人看的一句话（错误文案和后端回显共用一处）
MARK_TEXT = " / ".join(MARK_VALUES)
PASS_TEXT = " / ".join(PASS_VALUES)

# 两种形态各自「这格能填什么」。**只写这两份**：解析预览里每一行的原因、
# 单格保存时的 400 文案，用的都是它们，改词表不会再漂
SCORE_HINT = f"只能填 0~{MAX_SCORE} 之间的数字（最多两位小数），或 {MARK_TEXT}"
PASS_HINT = f"只能填 {PASS_TEXT}，或 {MARK_TEXT}"
# 在通过制科目里填了数字 / 日期之类的非文本值
_PASS_MODE_MSG = f"这是通过制科目，{PASS_HINT}"

# 分数段：按**满分的百分比**切，与满分到底是多少无关。
# 用百分比而不是分数区间 —— 150 分制下「90~104 分」没人看得懂，「60~69%」一眼就明白。
# 最后一档敞开：万一满分被设小了、有人考出超过 100%，那一格也不会凭空消失
SCORE_BUCKETS = [
    ("不足 60%", 0.0, 0.6),
    ("60~69%", 0.6, 0.7),
    ("70~79%", 0.7, 0.8),
    ("80~89%", 0.8, 0.9),
    ("90% 以上", 0.9, float("inf")),
]

# 预览令牌的临时目录。**独立目录**，不和 Excel 导入（class-oa-import-preview）、
# 照片导入（class-oa-photo-preview）共用 —— 各自清理自己的令牌，
# 不会出现「清理顺手把另一个功能的预览删了」
_PREVIEW_DIR = Path(tempfile.gettempdir()) / "class-oa-grade-preview"
_PREVIEW_TTL = 600  # 10 分钟

# 令牌是 uuid4().hex，必然是 32 位十六进制。**拼进路径之前先过这一道**：
# student_import_service 那边是 f-string 直接拼的，token 里塞 "../" 就能穿出目录
# （那份只有管理员够得着，所以一直没爆）；这份照 student_photo_service 的做法
# 加一道校验，成本一行
_TOKEN_RE = re.compile(r"[0-9a-f]{32}")

# 解析上限。真实班级远到不了，纯粹是挡「传错文件」（比如把整本成绩册传上来）
MAX_ROWS = 500
MAX_SUBJECTS = 20

# 数字形态：最多四位整数、两位小数。写成文本的「85分」「八十五」都不认 ——
# 与其猜，不如在预览里指名道姓地报出来让班委回 Excel 改
_NUMBER_RE = re.compile(r"^\d{1,4}(\.\d{1,2})?$")

# 这些表头看着就不是科目，预览里**默认不勾**（界面上随时能勾回来）。
# 「剩下的列都当科目」是用户口径，但一张成绩表里捎带一列「总分」太常见了，
# 不问一声就录进来就多出一列假成绩。列名精确匹配。
NON_SUBJECT_HEADERS = frozenset(
    {"序号", "编号", "班级", "专业", "备注", "总分", "总成绩", "平均分", "排名", "名次"}
)


# ---------- 取值归一化（解析和单格写入共用同一套判据）----------

def normalize_student_no(value) -> str:
    """学号归一化

    Excel 把学号存成**数字**时 openpyxl 可能给回 float（20260101.0），直接 str()
    会得到 "20260101.0"，跟库里的学号一个都对不上 —— 这是成绩表最常见的坑
    （学生信息那张表的学号列通常是文本，所以那边没这问题）。
    """
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float):
        # 20260101.0 → 20260101；真的带小数的学号不存在
        return str(int(value)) if value.is_integer() else str(value).strip()
    return str(value).strip()


def fmt_score(value: float) -> str:
    """分数 → 显示文本：90.0 → "90"，85.50 → "85.5"

    统一成这一种写法，格子文本才不会「同一个 90 分一会儿显示 90 一会儿 90.0」。
    """
    text = ("%.2f" % float(value)).rstrip("0").rstrip(".")
    return text or "0"


def cell_value(item: GradeItem | None) -> str:
    """一格成绩的显示文本；没有记录（或空）就是空串

    **必须用 `is not None` 判**：0 分是合法成绩却是假值，写成 `if item.score`
    会把 0 分当成没录。
    """
    if item is None:
        return ""
    if item.score is not None:
        return fmt_score(item.score)
    return item.mark or ""


def normalize_cell(raw, mode: str = "score") -> tuple[float | None, str | None]:
    """一格成绩 → (分数, 标记)，两者必有一个是 None

    (None, None) 表示**空格**（这格没录，不产生记录）。不合法抛 ValueError，
    消息里带上允许的取法 —— 它既是预览里那一行的原因，也是单格保存时的 400 文案。

    `mode` 是这一科的形态：通过制科目只认 通过 / 不通过，计分科目只认数字。
    **MARK_VALUES 两种形态都收** —— 通过制科目照样会有人缺考。两种形态都不收
    对方的取值，而且「在计分科目里填了通过」要明说是计分科目：班委看到这句才知道
    是自己敲错了词、还是科目形态配错了，光说「格式不对」他两边都得猜。

    默认 `score` 是为兼容只关心数字的调用方（测试）；真正落库的两处 ——
    解析和单格写入 —— 都会把科目形态传进来。
    """
    if raw is None:
        return None, None

    if isinstance(raw, str):
        text = raw.strip()
        if not text:
            return None, None
        if text in MARK_VALUES:
            return None, text
        if text in PASS_VALUES:
            if mode != "pass_fail":
                raise ValueError(f"「{text}」是通过制科目的成绩，这一科按分数记，{SCORE_HINT}")
            return None, text
        if mode == "pass_fail":
            raise ValueError(_PASS_MODE_MSG)
        if not _NUMBER_RE.match(text):
            raise ValueError(SCORE_HINT)
        value = float(text)
    elif isinstance(raw, bool):
        # isinstance(True, int) 在 Python 里为真，不显式拦的话 TRUE 会被当成 1 分
        # （同 services/poll_service.read_choices 那条注释）
        raise ValueError("TRUE / FALSE 不是成绩")
    elif isinstance(raw, (int, float)):
        if mode == "pass_fail":
            raise ValueError(_PASS_MODE_MSG)
        value = float(raw)
    else:
        # 日期格式的格子会走到这儿：与其猜不如让班委自己看
        raise ValueError(_PASS_MODE_MSG if mode == "pass_fail" else SCORE_HINT)

    if not 0 <= value <= MAX_SCORE:
        raise ValueError(f"分数要在 0~{MAX_SCORE} 之间")
    return round(value, 2), None


def _display(raw) -> str:
    """格子的原样回显（预览里给班委看的）"""
    if raw is None:
        return ""
    if isinstance(raw, float) and raw.is_integer():
        return str(int(raw))
    return str(raw).strip()


# ---------- 解析 ----------

class GradePreview:
    """一次上传的解析结果

    名字里带 public 的那些字段可以安全返回给前端；`_rows` 是内部数据，
    确认入库时才用得上。
    """

    def __init__(
        self,
        sheet_name: str,
        available_sheets: list[str],
        columns: list[dict],
        sample_rows: list[dict],
        student_count: int,
        matched_count: int,
        unmatched: list[dict],
        name_mismatch: list[dict],
        duplicates: list[dict],
        problems: list[dict],
        warnings: list[str],
    ):
        self.sheet_name = sheet_name
        self.available_sheets = available_sheets
        self.columns = columns
        self.sample_rows = sample_rows
        self.student_count = student_count
        self.matched_count = matched_count
        self.unmatched = unmatched
        self.name_mismatch = name_mismatch
        self.duplicates = duplicates
        self.problems = problems
        self.warnings = warnings
        # 内部：待入库的行 [{row, student_no, student_name, cells: {科目名: [分数, 标记]}}]
        self._rows: list[dict] = []

    @property
    def column_names(self) -> list[str]:
        return [c["name"] for c in self.columns]

    def to_public(self) -> dict:
        """返回给前端的字段（不含待入库的行数据）"""
        return {
            "sheet_name": self.sheet_name,
            "available_sheets": self.available_sheets,
            "columns": self.columns,
            "sample_rows": self.sample_rows,
            "student_count": self.student_count,
            "matched_count": self.matched_count,
            "unmatched": self.unmatched,
            "name_mismatch": self.name_mismatch,
            "duplicates": self.duplicates,
            "problems": self.problems,
            "mark_values": list(MARK_VALUES),
            "warnings": self.warnings,
        }


def _pick_sheet(workbook, sheet_name: str | None) -> str:
    """定用哪张工作表：指定了就用指定的，否则取第一张第一行带「学号」的

    默认**不是「第一张」而是「第一张像成绩表的」**：实际收到的成绩表常把封面、
    填表说明放在第一个工作表，数据在第二张。按第一张解析只会得到一句「缺「学号」列」，
    班委除了回 Excel 删表没别的路。显式指定了工作表就照指定的来，那张表里没有
    「学号」列照旧报错 —— 他要的就是那张。
    """
    available = workbook.sheetnames
    if sheet_name:
        if sheet_name not in available:
            raise ValueError(f"工作表「{sheet_name}」不存在")
        return sheet_name
    for name in available:
        header = [str(c.value).strip() if c.value is not None else "" for c in workbook[name][1]]
        if "学号" in header:
            return name
    if len(available) > 1:
        # 单张表时报「必须有「学号」列」就够清楚了，多张才需要把名字列给他看
        raise ValueError("这些工作表的第一行里都没有「学号」列：" + "、".join(available))
    return available[0]


def parse_grade_excel(
    file_bytes: bytes, sheet_name: str | None = None, db: Session | None = None
) -> GradePreview:
    """解析成绩 Excel，返回预览（**不写库**）

    表头在第 1 行：`学号` 必须有（精确匹配，没有就报错），`姓名` 可选，
    其余非空列一律当科目候选。传了 db 就顺手核对学号在不在名单里。
    """
    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = _pick_sheet(workbook, sheet_name)
    worksheet = workbook[sheet]

    # 表头。行号从 1 开始，报错文案里用的都是**Excel 里的真实行号**
    header_texts = [
        str(cell.value).strip() if cell.value is not None else "" for cell in worksheet[1]
    ]

    if "学号" not in header_texts:
        raise ValueError("表格第一行里必须有「学号」列")
    no_col = header_texts.index("学号")
    name_col = header_texts.index("姓名") if "姓名" in header_texts else None

    warnings: list[str] = []
    if name_col is None:
        warnings.append("没有找到「姓名」列，已跳过姓名核对")

    # 科目列 = 除学号/姓名外的所有非空表头。重名的只留第一列（否则「科目名」
    # 就没法当身份用，确认时传回来的 subjects 会有歧义）
    subject_cols: list[tuple[int, str]] = []
    seen: set[str] = set()
    for index, text in enumerate(header_texts):
        if not text or index in (no_col, name_col):
            continue
        if text in seen:
            warnings.append(f"表头里有重复的「{text}」，只取第一列")
            continue
        seen.add(text)
        subject_cols.append((index, text))

    if not subject_cols:
        raise ValueError("除「学号」「姓名」外没有别的列，没找到科目")
    if len(subject_cols) > MAX_SUBJECTS:
        raise ValueError(f"科目最多 {MAX_SUBJECTS} 个，这张表有 {len(subject_cols)} 个")

    # 逐行读原始值。**迭代次数硬夹在 MAX_ROWS 上**：Excel 的幽灵空行会让
    # max_row 虚高到上百万，不夹住就会白转一百万次
    raw_rows: list[dict] = []
    for row_index, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if row_index > 1 + MAX_ROWS:
            break
        values = list(row)

        def at(index: int):
            return values[index] if index < len(values) else None

        student_no = normalize_student_no(at(no_col))
        name = _display(at(name_col)) if name_col is not None else ""
        cells = [(text, at(index)) for index, text in subject_cols]
        has_any = student_no or name or any(_display(v) for _, v in cells)
        if not has_any:
            continue  # 整行空的（表格末尾）静默跳过
        raw_rows.append({"row": row_index, "student_no": student_no, "name": name, "cells": cells})

    # 夹住迭代次数之后再看一眼下一行，确认是被截断还是真的到头了
    if worksheet.max_row > 1 + MAX_ROWS:
        following = next(
            worksheet.iter_rows(min_row=2 + MAX_ROWS, max_row=2 + MAX_ROWS, values_only=True),
            None,
        )
        if following and any(_display(v) for v in following):
            raise ValueError(f"表格超过 {MAX_ROWS} 行，请分次上传")

    if not raw_rows:
        raise ValueError("表格里没有数据行")

    # 列预览：每列挑前几个非空值给班委看「这列到底是不是科目」。
    # 顺便预判这一列是不是通过制 —— 整列的取值全是「通过 / 不通过」就建议按通过制录
    # （确认页上默认选中，班委可改回来）。这个预判**同时是解析判据**：不拿它去解析的话，
    # 通过制那一列的值会在解析阶段被逐格判成非法，班委看到满屏红字，
    # 其实只是形态还没选而已
    col_values = [
        [v for v in (_display(r["cells"][i][1]) for r in raw_rows) if v]
        for i in range(len(subject_cols))
    ]
    columns = [
        {
            "name": subject_cols[i][1],
            "keep_default": subject_cols[i][1] not in NON_SUBJECT_HEADERS,
            "suggest_pass_fail": bool(values) and all(v in PASS_VALUES for v in values),
            "sample": values[:3],
        }
        for i, values in enumerate(col_values)
    ]
    col_modes = {
        columns[i]["name"]: ("pass_fail" if columns[i]["suggest_pass_fail"] else "score")
        for i in range(len(subject_cols))
    }

    sample_rows = [
        {
            "学号": r["student_no"],
            "姓名": r["name"],
            **{text: _display(value) for text, value in r["cells"]},
        }
        for r in raw_rows[:5]
    ]

    unmatched: list[dict] = []
    name_mismatch: list[dict] = []
    duplicates: list[dict] = []
    problems: list[dict] = []
    warnings_out = list(warnings)
    rows: list[dict] = []
    rounded = 0

    # 名单核对只查表里出现过的学号，不把整张表拉进内存
    live_names: dict[str, str] = {}
    if db is not None:
        wanted = {r["student_no"] for r in raw_rows if r["student_no"]}
        if wanted:
            live_names = {
                no: (name or no)
                for no, name in db.query(Student.student_no, Student.name)
                .filter(Student.student_no.in_(wanted))
                .all()
            }

    seen_nos: set[str] = set()
    for raw in raw_rows:
        student_no = raw["student_no"]
        if not student_no:
            unmatched.append(
                {"row": raw["row"], "student_no": "", "student_name": raw["name"], "reason": "这一行没有学号"}
            )
            continue
        if student_no in seen_nos:
            duplicates.append({"row": raw["row"], "student_no": student_no})
            continue
        if db is not None and student_no not in live_names:
            unmatched.append(
                {
                    "row": raw["row"],
                    "student_no": student_no,
                    "student_name": raw["name"],
                    "reason": "学号不在班级名单里",
                }
            )
            continue
        seen_nos.add(student_no)

        # 学号才是身份，姓名只用来核对：对不上照收，但要让班委看见
        db_name = live_names.get(student_no, "")
        if student_no in live_names:
            db_name = live_names[student_no]
            if raw["name"] and raw["name"] != db_name:
                name_mismatch.append(
                    {
                        "row": raw["row"],
                        "student_no": student_no,
                        "excel_name": raw["name"],
                        "db_name": db_name,
                    }
                )

        cells: dict[str, list] = {}
        for text, value in raw["cells"]:
            try:
                score, mark = normalize_cell(value, col_modes[text])
            except ValueError as exc:
                # 一格填错只丢这一格，不连累整行、更不连累整张表
                problems.append(
                    {
                        "row": raw["row"],
                        "student_no": student_no,
                        "column": text,
                        "raw": _display(value),
                        "reason": str(exc),
                    }
                )
                continue
            if score is None and mark is None:
                continue  # 空格：不产生记录
            if score is not None and _display(value) != fmt_score(score):
                rounded += 1
            cells[text] = [score, mark]

        rows.append(
            {
                "row": raw["row"],
                "student_no": student_no,
                # 入库用库里的名字（学号才是身份），查不到库就退回表里写的
                "student_name": db_name or raw["name"] or student_no,
                "cells": cells,
            }
        )

    if rounded:
        warnings_out.append(f"有 {rounded} 个格子的小数超过两位，已四舍五入到两位小数")
    if unmatched:
        warnings_out.append(f"有 {len(unmatched)} 行学号对不上，不会入库")
    if problems:
        warnings_out.append(f"有 {len(problems)} 个格子没读懂，那些格子会留空")

    preview = GradePreview(
        sheet_name=sheet,
        available_sheets=workbook.sheetnames,
        columns=columns,
        sample_rows=sample_rows,
        student_count=len(raw_rows),
        matched_count=len(rows),
        unmatched=unmatched,
        name_mismatch=name_mismatch,
        duplicates=duplicates,
        problems=problems,
        warnings=warnings_out,
    )
    preview._rows = rows
    return preview


# ---------- 预览令牌（临时文件）----------

def _preview_path(token: str) -> Path:
    """令牌 → 文件路径。**格式不对就给一个不可能存在的路径**

    不做路径穿越防护的话，"../../x" 能拼到目录外面去（见 _TOKEN_RE 的说明）。
    """
    if not _TOKEN_RE.fullmatch(token or ""):
        return _PREVIEW_DIR / "invalid"
    return _PREVIEW_DIR / f"preview_{token}.json"


def save_preview(preview: GradePreview) -> str:
    """把解析结果存成临时文件，返回令牌

    存的是**解析后的行数据**（不含 Excel 字节）：确认阶段完全不需要 openpyxl，
    一个班几十人几百个格子，JSON 也就几十 KB。
    """
    # 顺手清一遍过期的。**不这么做的话弃用的预览会一直躺在 /tmp**：
    # 只有「确认录入」那条路径会删令牌，传上来又退出去的永远不会被删
    # （学生导入、照片导入那两份就是只定义了没调用）。上传是低频动作，
    # 这里 glob 一次几个文件的开销可以忽略
    cleanup_expired_previews()

    token = uuid.uuid4().hex
    data = {
        "public": preview.to_public(),
        "rows": preview._rows,
        "created_at": time.time(),
    }
    _PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    _preview_path(token).write_text(
        json.dumps(data, ensure_ascii=False), encoding="utf-8"
    )
    return token


def load_preview(token: str) -> GradePreview | None:
    """按令牌取回解析结果。过期或不存在返回 None"""
    path = _preview_path(token)
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None
    if time.time() - data.get("created_at", 0) > _PREVIEW_TTL:
        drop_preview(token)
        return None

    public = data["public"]
    preview = GradePreview(
        sheet_name=public["sheet_name"],
        available_sheets=public["available_sheets"],
        columns=public["columns"],
        sample_rows=public["sample_rows"],
        student_count=public["student_count"],
        matched_count=public["matched_count"],
        unmatched=public["unmatched"],
        name_mismatch=public["name_mismatch"],
        duplicates=public["duplicates"],
        problems=public["problems"],
        warnings=public["warnings"],
    )
    preview._rows = data["rows"]
    return preview


def drop_preview(token: str) -> None:
    """删掉令牌文件（用完就删，不用等 TTL）。静默失败"""
    try:
        _preview_path(token).unlink()
    except OSError:
        pass


def cleanup_expired_previews() -> None:
    """清掉过期的令牌文件。不是必须调用，偶尔跑一下就行"""
    if not _PREVIEW_DIR.exists():
        return
    now = time.time()
    for path in _PREVIEW_DIR.glob("preview_*.json"):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            expired = now - data.get("created_at", 0) > _PREVIEW_TTL
        except (json.JSONDecodeError, OSError):
            expired = True
        if expired:
            try:
                path.unlink()
            except OSError:
                pass


# ---------- 写入 ----------

def import_grade(
    db: Session,
    preview: GradePreview,
    *,
    name: str,
    tested_at,
    remark: str,
    subjects: list[str],
    operator: User,
    ip: str | None = None,
    subject_modes: dict[str, str] | None = None,
    subject_full_scores: dict[str, float] | None = None,
) -> dict:
    """按预览把成绩写进库，返回统计（**不发通知**，通知由接口在 commit 之后发）

    subjects 是**保留**的科目名（顺序即列序），勾掉的列连数据一起不要。
    测试、科目、明细、操作日志在同一个事务里，只 commit 一次。

    调用方负责校验 subjects 是预览里出现过的科目名（这里只做防御性过滤），
    以及预览里至少匹配到了一个人。
    """
    from .log_service import log_operation
    from .user_service import account_name

    kept = [text for text in subjects if text in preview.column_names]
    if not kept:
        raise ValueError("至少要保留一个科目")

    subject_modes = subject_modes or {}
    subject_full_scores = subject_full_scores or {}

    # 形态和数据要对得上，**两个方向都要挡**。解析时用的是预览给的预判形态，
    # 而班委在确认页上能把它改回来；改反了的话写进去的数据就和这一科的形态对不上，
    # 事后谁也解释不清该怎么看，所以宁可当场拒掉并点名哪一科
    for text in kept:
        mode = subject_modes.get(text, "score")
        cells = [row["cells"][text] for row in preview._rows if text in row["cells"]]
        if mode == "pass_fail":
            if any(score is not None for score, _ in cells):
                raise ValueError(
                    f"「{text}」这一列里有分数，不能按通过制录入"
                    "（把它改成计分，或回 Excel 把这一列改成 通过 / 不通过）"
                )
        elif any(mark in PASS_VALUES for _, mark in cells):
            raise ValueError(
                f"「{text}」这一列里有通过 / 不通过，不能按计分录入"
                "（把它改成通过制，或回 Excel 把这一列改成数字）"
            )

    test = GradeTest(
        name=name,
        tested_at=tested_at,
        remark=remark,
        creator_no=operator.student_no or "",
        creator_name=account_name(db, operator),
    )
    db.add(test)
    db.flush()  # 先拿到自增 id，科目和明细都要用

    db.add_all(
        [
            GradeSubject(
                test_id=test.id,
                name=text,
                sort=index,
                mode=subject_modes.get(text, "score"),
                full_score=subject_full_scores.get(text, DEFAULT_FULL_SCORE),
            )
            for index, text in enumerate(kept)
        ]
    )
    db.flush()
    subject_ids = {
        s.name: s.id for s in db.query(GradeSubject).filter(GradeSubject.test_id == test.id)
    }

    items = []
    for row in preview._rows:
        for text in kept:
            cell = row["cells"].get(text)
            if not cell:
                continue
            score, mark = cell
            items.append(
                GradeItem(
                    test_id=test.id,
                    subject_id=subject_ids[text],
                    student_no=row["student_no"],
                    student_name=row["student_name"],
                    score=score,
                    mark=mark,
                )
            )
    db.add_all(items)

    log_operation(
        db,
        operator,  # 只 add 不 commit，随上面那笔事务一起提交
        action="导入成绩",
        target=f"{tested_at:%Y-%m-%d} {name}",
        detail=f"{len({i.student_no for i in items})} 人 {len(kept)} 科，"
        f"跳过 {len(preview.unmatched) + len(preview.duplicates)} 行",
        ip=ip,
    )
    db.commit()

    return {
        "test_id": test.id,
        "name": test.name,
        "subject_count": len(kept),
        "student_count": len({i.student_no for i in items}),
        "cell_count": len(items),
        "skipped": len(preview.unmatched) + len(preview.duplicates),
    }


def set_cell(
    db: Session, test: GradeTest, subject: GradeSubject, student_no: str, value: str
) -> tuple[GradeItem | None, str, str]:
    """写一个格子（没有记录就新建，清空就把记录删掉）

    返回 (记录, 旧文本, 新文本)。**值没变就原样返回、不碰库** —— 同
    api/attendance.py「状态没变就不写日志」的口径，否则操作日志里会出现一串
    假的修改记录。不合法由 normalize_cell 抛 ValueError，接口转 400。

    「清空」是**删行**而不是留一行空记录：模型里「没有行 = 这格没录」，
    留空行会让「录入人数」这类聚合数字虚高。
    """
    # 按这一科的形态判取值：通过制科目填数字、计分科目填「通过」都要当场拒掉
    score, mark = normalize_cell(value, subject.mode or "score")
    item = (
        db.query(GradeItem)
        .filter(GradeItem.subject_id == subject.id, GradeItem.student_no == student_no)
        .first()
    )
    old = cell_value(item)
    new = fmt_score(score) if score is not None else (mark or "")
    if old == new:
        return item, old, new

    if score is None and mark is None:
        if item is None:
            return None, old, new
        db.delete(item)
        return None, old, new

    if item is None:
        item = GradeItem(
            test_id=test.id,
            subject_id=subject.id,
            student_no=student_no,
            student_name="",
            score=score,
            mark=mark,
        )
        db.add(item)
    else:
        item.score = score
        item.mark = mark
    return item, old, new


# ---------- 组装 ----------

def _counts_of(db: Session, test_ids: list[int]) -> dict:
    """一次查清一批测试的两个计数，供列表接口用（逐条查就是 2N 次）"""
    if not test_ids:
        return {"subjects": {}, "students": {}}
    return {
        "subjects": dict(
            db.query(GradeSubject.test_id, func.count(GradeSubject.id))
            .filter(GradeSubject.test_id.in_(test_ids))
            .group_by(GradeSubject.test_id)
            .all()
        ),
        # 一个人有多科的成绩，所以人数要数**人头**不是数行
        "students": dict(
            db.query(GradeItem.test_id, func.count(func.distinct(GradeItem.student_no)))
            .filter(GradeItem.test_id.in_(test_ids))
            .group_by(GradeItem.test_id)
            .all()
        ),
    }


def list_items(db: Session, tests: list[GradeTest]) -> list[dict]:
    """把一批测试组装成列表项（两个计数是现算的，表里没有这两列）"""
    counts = _counts_of(db, [test.id for test in tests])
    return [
        {
            "id": test.id,
            "name": test.name,
            "tested_at": test.tested_at,
            "remark": test.remark or "",
            "creator_no": test.creator_no or "",
            "creator_name": test.creator_name or "",
            "created_at": test.created_at,
            "subject_count": counts["subjects"].get(test.id, 0),
            "student_count": counts["students"].get(test.id, 0),
        }
        for test in tests
    ]


def subjects_of(db: Session, test_id: int) -> list[GradeSubject]:
    """一次测试的科目，按列序排

    排序带 id 兜底：sort 撞上时顺序不会在两次请求之间跳来跳去（同
    poll_service._questions_of）。
    """
    return (
        db.query(GradeSubject)
        .filter(GradeSubject.test_id == test_id)
        .order_by(GradeSubject.sort.asc(), GradeSubject.id.asc())
        .all()
    )


def detail(db: Session, test: GradeTest) -> dict:
    """一次测试的明细：科目 + 每个学生的每格成绩

    行 = **学籍正常的学生** ∪ **这次测试已经有成绩记录的人**。前者保证新转入的学生
    也有格子可补录；后者保证毕业删档（物理删除）之后，历史成绩还看得见 ——
    名字退回 grade_items 里的快照。

    每行每科都给键（空串表示没录），前端拿 `row.values[String(s.id)]` 才不会
    踩到 undefined。
    """
    subjects = subjects_of(db, test.id)
    roster = {s.student_no: s for s in db.query(Student).all()}
    items = db.query(GradeItem).filter(GradeItem.test_id == test.id).all()

    rows: dict[str, dict] = {}
    for student in roster.values():
        if (student.school_status or "") == "正常":
            rows[student.student_no] = {
                "student_no": student.student_no,
                "student_name": student.name or student.student_no,
                "in_roster": True,
                "school_status": student.school_status or "",
                "values": {str(s.id): "" for s in subjects},
            }

    for item in items:
        row = rows.get(item.student_no)
        if row is None:
            student = roster.get(item.student_no)
            row = {
                "student_no": item.student_no,
                "student_name": (student.name if student else None)
                or item.student_name
                or item.student_no,
                "in_roster": student is not None,
                "school_status": (student.school_status if student else "") or "",
                "values": {str(s.id): "" for s in subjects},
            }
            rows[item.student_no] = row
        row["values"][str(item.subject_id)] = cell_value(item)

    data = list_items(db, [test])[0]
    data["subjects"] = [
        {
            "id": s.id,
            "name": s.name,
            "sort": s.sort or 0,
            # 形态和满分跟着科目一起下发：明细表头要立刻显示「150 分」「通过制」，
            # 不能等统计接口回来才知道
            "mode": s.mode or "score",
            "full_score": float(s.full_score or DEFAULT_FULL_SCORE) or DEFAULT_FULL_SCORE,
        }
        for s in subjects
    ]
    data["rows"] = [rows[no] for no in sorted(rows)]
    data["mark_values"] = list(MARK_VALUES)
    return data


def subject_stats(db: Session, test: GradeTest) -> dict:
    """一场测试的**单科**统计（班委及以上看）

    刻意只有单科、没有跨科合计 —— 这个模块里不存在「总分」这个概念
    （用户口径），所以别在这里顺手加个总分列。

    三类格子分开数：**有分数的**（平均/最高/最低/分数段/及格率的分母）、
    **有标记词的**（缺考/免修/缓考/违纪）、**没录的**。混着数的话缺考会被当成 0 分
    拽低平均分 —— 而缺考的人根本没参加这场考试，不该出现在任何分数里。
    """
    subjects = subjects_of(db, test.id)
    # 一场最多 20 科 × 500 行，一次全捞出来在内存里分科；按科目发 20 次查询没必要
    items = db.query(GradeItem).filter(GradeItem.test_id == test.id).all()
    by_subject: dict[int, list[GradeItem]] = {}
    for item in items:
        by_subject.setdefault(item.subject_id, []).append(item)

    # 统计范围内的人 = 学籍正常的学生 ∪ 这次已有记录的人（同 detail() 的口径，
    # 后者保证毕业删档的人也算得进去）。「没录」这个数要靠它才数得出来
    roster = {
        s.student_no for s in db.query(Student).all() if (s.school_status or "") == "正常"
    }
    everyone = roster | {item.student_no for item in items}

    rows = []
    for subject in subjects:
        group = by_subject.get(subject.id, [])
        mode = subject.mode or "score"
        # 满分为 0 或没设（`or` 一并兜住）就退回默认值，免得下面拿它做除数
        full_score = float(subject.full_score or DEFAULT_FULL_SCORE) or DEFAULT_FULL_SCORE

        # 判空一律 is not None：0 分是合法成绩，写成 if item.score 会把它当成没录
        scores = [item.score for item in group if item.score is not None]
        # 通过 / 不通过也存在 mark 列里，但它们是**成绩**、不是「没参加考试」的原因，
        # 得跟缺考那几个分开数 —— 混在一起数的话，通过制科目的「标记人数」
        # 会把全部有成绩的人都算进去
        passed = [i for i in group if i.mark == PASS_VALUES[0]]
        failed = [i for i in group if i.mark == PASS_VALUES[1]]
        marked = [i for i in group if i.score is None and i.mark in MARK_VALUES]

        # 四类互不重叠，加起来就是「有记录的人」，剩下的才是没录的
        recorded = len(scores) + len(passed) + len(failed) + len(marked)
        row = {
            "subject_id": subject.id,
            "name": subject.name,
            "mode": mode,
            "full_score": full_score,
            # 有成绩的人数：计分科目是有分数的，通过制是记了通过 / 不通过的。
            # 两种情况都正好是下面那个比例的分母
            "scored_count": len(scores) if mode == "score" else len(passed) + len(failed),
            "marked_count": len(marked),
            "blank_count": len(everyone) - recorded,
        }

        if mode == "pass_fail":
            row["passed_count"] = len(passed)
            row["failed_count"] = len(failed)
            # 通过率的分母只有真记了通过/不通过的人，缺考的两种都不是
            row["pass_rate"] = (
                round(len(passed) / len(passed + failed) * 100, 1)
                if passed or failed else None
            )
        else:
            # 显式 float()：AVG / 除法在 SQLite 上给 float、在 MySQL 上可能是 Decimal
            row["avg"] = round(float(sum(scores) / len(scores)), 1) if scores else None
            row["max"] = max(scores) if scores else None
            row["min"] = min(scores) if scores else None
            # 及格线 = 满分 × 比例。分数段按满分的百分比切，与满分具体是多少无关
            row["pass_count"] = sum(1 for s in scores if s >= full_score * PASS_RATIO)
            row["pass_rate"] = (
                round(row["pass_count"] / len(scores) * 100, 1) if scores else None
            )
            row["ranges"] = [
                {
                    "label": label,
                    "count": sum(1 for s in scores if low <= s / full_score < high),
                }
                for label, low, high in SCORE_BUCKETS
            ]

        rows.append(row)

    return {
        "test_id": test.id,
        "name": test.name,
        "tested_at": test.tested_at,
        "subjects": rows,
        "roster_count": len(everyone),
    }


def _my_entry(subject: GradeSubject, value: str, pool: dict[int, list[float]]) -> dict:
    """我的一格成绩 + 它在这一科的位次（pool 是这一科全部有分数的取值）

    位次口径：**排在前 X%** = 名次 ÷ 有成绩人数。名次取「严格比我高的人数 + 1」,
    并列同分时按最好名次算，不因为有人同分就把自己挤到后面。

    分母（有成绩人数）只用来算这个比例，**不返回给学生** —— 学生端只给百分位，
    不给「一共多少人考了」（用户口径）。

    只在**计分科目**且**我这一格确实是分数**时才算：
    - 通过制科目不算 —— 二元值的百分位没有意义
    - 缺考/免修/缓考/违纪不算 —— 这些值 float() 会抛，本来也不是成绩
    - 空格不算 —— 没录哪来的位次

    班级层面的统计（平均分、及格率、分数段）一律不给学生，只有自己这一格。
    """
    entry = {
        "subject": subject.name,
        "value": value,
        "mode": subject.mode or "score",
        "rank_percent": None,
    }
    if entry["mode"] != "score" or not value:
        return entry
    try:
        score = float(value)
    except ValueError:
        return entry  # 标记词（缺考等）走到这儿

    group = pool.get(subject.id, [])
    if group:
        ahead = sum(1 for other in group if other > score)
        entry["rank_percent"] = round((ahead + 1) / len(group) * 100, 1)
    return entry


def my_grades(db: Session, student_no: str) -> list[dict]:
    """我参加过的每次测试的各科成绩

    **只列我确实有成绩记录的测试**：Excel 里压根没有我的行（没参加、或录入时漏了）
    的测试不出现 —— 与其显示一整屏空表格，不如不显示（用户口径：学生只看自己的分数）。

    有的科目上面列**全部科目**、我没录到的那格给空串：只列我有记录的科目会让
    「这次到底考了几科」看不出来。
    """
    items = db.query(GradeItem).filter(GradeItem.student_no == student_no).all()
    if not items:
        return []

    test_ids = {item.test_id for item in items}
    tests = {t.id: t for t in db.query(GradeTest).filter(GradeTest.id.in_(test_ids)).all()}
    subjects: dict[int, GradeSubject] = {
        s.id: s
        for s in db.query(GradeSubject).filter(GradeSubject.test_id.in_(test_ids)).all()
    }

    # 位次：分母是**这一科实际有分数的人**（缺考/免修/缓考/违纪不算，与及格率同口径）。
    # 一次把这些科目的所有分数捞回来在内存里比 —— 一场几十人，内存里算比逐科发
    # 一条 count 查询省事得多
    pool: dict[int, list[float]] = {}
    if subjects:
        for subject_id, score in (
            db.query(GradeItem.subject_id, GradeItem.score)
            .filter(GradeItem.subject_id.in_(list(subjects)), GradeItem.score.isnot(None))
            .all()
        ):
            pool.setdefault(subject_id, []).append(score)

    mine: dict[int, dict[int, str]] = {}
    for item in items:
        mine.setdefault(item.test_id, {})[item.subject_id] = cell_value(item)

    result = []
    for test_id, values in mine.items():
        test = tests.get(test_id)
        if test is None:
            continue  # 脏数据：测试没了只剩余数，跳过而不是把接口打成 500
        group = [s for s in subjects.values() if s.test_id == test_id]
        group.sort(key=lambda s: (s.sort or 0, s.id))
        result.append(
            {
                "test_id": test.id,
                "name": test.name,
                "tested_at": test.tested_at,
                "remark": test.remark or "",
                "scores": [_my_entry(s, values.get(s.id, ""), pool) for s in group],
            }
        )
    result.sort(key=lambda row: (row["tested_at"], row["test_id"]), reverse=True)
    return result
