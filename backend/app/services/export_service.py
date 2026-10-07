"""学生信息导出服务

两种导出：
1. Excel 名单（excel_list）：选字段，一张表
2. 个人信息表 PDF（pdf_individual）：每人一页完整信息，打包 zip

PDF 导出打不可见水印（复用 pdf_watermark_service），溯源是谁导出的。
"""

import io
import json
import zipfile
from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from ..models.student import Student
from ..utils.student_fields import FIELD_GROUPS, FIELD_MAP
from .pdf_watermark_service import add_invisible_watermark

# 字段清单 FIELD_GROUPS / FIELD_MAP 放在 utils/student_fields.py：操作日志也要用
# 同一份中文名，各存一份迟早会走散（加字段只改一边，另一边冒英文 prop）。

# 默认选中的字段（Excel 导出时前端默认勾上的）
DEFAULT_FIELDS = [
    "student_no", "name", "gender", "ethnicity", "political_status",
    "dormitory", "bed", "phone", "class_position", "gaokao_score",
]


def get_field_groups():
    """返回字段分组信息（给前端渲染复选框用），只含 prop 和 label，不含 getter"""
    return [
        {
            "key": g["key"],
            "label": g["label"],
            "fields": [{"prop": prop, "label": label} for prop, label, _ in g["fields"]],
        }
        for g in FIELD_GROUPS
    ]


def get_default_fields():
    """默认选中的字段列表"""
    return DEFAULT_FIELDS


# 导出范围的中文名。日志、审核列表、水印溯源都要展示，
# 集中在这里，免得几个地方各写一份对不上（把 scope_type 原样抛出去，
# 用户看到的就是「selected」这种内部代号）
SCOPE_LABELS = {
    "all": "全部",
    "filtered": "当前筛选",
    "selected": "手动勾选",
    "id_range": "学号范围",
}


def scope_label(scope_type: str, count: int | None = None) -> str:
    """范围的中文描述，给了人数就带上：手动勾选（3 人）"""
    label = SCOPE_LABELS.get(scope_type, scope_type or "未知")
    return f"{label}（{count} 人）" if count is not None else label


def _get_student_value(student: Student, prop: str) -> str:
    """从 Student 对象取字段值，格式化后返回字符串"""
    info = FIELD_MAP.get(prop)
    if not info:
        return ""
    val = getattr(student, prop, None)
    return info["getter"](val)


def _markdown_to_paragraph_html(text: str) -> str:
    """把简易 Markdown 转成 reportlab Paragraph 能用的 HTML 标签

    只处理最常用的几种语法，够用就行：
    - **粗体** → <b>...</b>
    - *斜体* → <i>...</i>
    - 行内代码 `code` → 等宽字体
    - 链接 [text](url) → <a href="url" color="blue">text</a>
    - 换行 → <br/>
    - 列表（- 或 1. 开头）→ 缩进 + 项目符号

    不处理标题（#）——标题一般写在段落开头，PDF 里用列表或加粗就够了。
    """
    if not text:
        return ""

    lines = text.replace("\r\n", "\n").split("\n")
    result_lines = []

    for line in lines:
        if not line.strip():
            result_lines.append("")
            continue

        # 行内格式化（粗体 / 斜体 / 代码 / 链接）
        s = line
        # 链接：[text](url)
        import re as _re
        s = _re.sub(
            r"\[([^\]]+)\]\(([^)]+)\)",
            r'<a href="\2" color="blue"><u>\1</u></a>',
            s,
        )
        # 粗体：**text**
        s = _re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
        # 斜体：*text*（注意：粗体已经处理过了，剩下的 * 就是斜体）
        s = _re.sub(r"(?<!\*)\*([^*\n]+?)\*(?!\*)", r"<i>\1</i>", s)
        # 行内代码：`text`
        s = _re.sub(r"`([^`]+)`", r'<font face="Courier">\1</font>', s)

        # 无序列表：- / * / + 开头
        list_match = _re.match(r"^\s*[-*+]\s+", s)
        if list_match:
            content = s[list_match.end():]
            result_lines.append(f"&nbsp;&nbsp;•&nbsp;&nbsp;{content}")
            continue

        # 有序列表：1. / 2. 开头
        ol_match = _re.match(r"^\s*\d+\.\s+", s)
        if ol_match:
            prefix = s[:ol_match.end()].strip()
            content = s[ol_match.end():]
            result_lines.append(f"&nbsp;&nbsp;{prefix}&nbsp;{content}")
            continue

        result_lines.append(s)

    return "<br/>".join(result_lines)


def _parse_remark_attachments_pdf(student: Student) -> list[dict]:
    """从学生对象读备注附件列表（PDF 导出用）"""
    if not getattr(student, "remark_attachments", ""):
        return []
    try:
        items = json.loads(student.remark_attachments or "")
    except (ValueError, TypeError):
        return []
    return items if isinstance(items, list) else []


# ---------------------------------------------------------------------------
# Excel 导出
# ---------------------------------------------------------------------------

def generate_excel(students: List[Student], fields: List[str]) -> bytes:
    """生成学生名单 Excel，返回二进制

    fields 是 prop 列表，顺序就是列的顺序。
    """
    # 过滤掉不在白名单里的字段
    valid_fields = [f for f in fields if f in FIELD_MAP]
    if not valid_fields:
        # 如果一个都没选（不应该发生），退回默认字段
        valid_fields = DEFAULT_FIELDS

    wb = Workbook()
    ws = wb.active
    ws.title = "学生名单"

    # 表头行
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
    for col_idx, prop in enumerate(valid_fields, 1):
        cell = ws.cell(row=1, column=col_idx, value=FIELD_MAP[prop]["label"])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 数据行
    for row_idx, stu in enumerate(students, 2):
        for col_idx, prop in enumerate(valid_fields, 1):
            val = _get_student_value(stu, prop)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    # 列宽：按内容估算，中文占 2 个字符宽度
    for col_idx, prop in enumerate(valid_fields, 1):
        label_len = len(FIELD_MAP[prop]["label"]) * 2
        # 取「表头宽度」和「最长内容宽度」的较大值，再加 2 个 padding
        max_content = max(
            (len(_get_student_value(s, prop)) for s in students),
            default=0
        ) * 2
        width = min(max(label_len, max_content) + 2, 40)  # 最宽 40，别太夸张
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 首行冻结
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF 个人信息表
# ---------------------------------------------------------------------------

# 中文字体：用 reportlab 内置的 CID 字体（Adobe STSong-Light，宋体）
# 不需要额外字体文件，任何环境都能用，中文不会变方块
def _register_cjk_font():
    """注册中文字体，返回字体名。注册过一次就直接返回。"""
    if hasattr(_register_cjk_font, "_font_name"):
        return _register_cjk_font._font_name

    font_name = "STSong-Light"  # reportlab 内置 CID 字体，宋体，中文零依赖
    try:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    except Exception:
        font_name = "Helvetica"  # 极端情况下的兜底
    _register_cjk_font._font_name = font_name
    return font_name


def generate_student_pdf(student: Student, photos_dir: Path) -> bytes:
    """生成单个学生的个人信息表 PDF（一页 A4），返回二进制"""
    font_name = _register_cjk_font()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title=f"{student.name} 个人信息表",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCN", parent=styles["Title"],
        fontName=font_name, fontSize=18, leading=24, alignment=1, spaceAfter=16,
    )
    h2_style = ParagraphStyle(
        "H2CN", parent=styles["Heading2"],
        fontName=font_name, fontSize=13, leading=18, spaceBefore=10, spaceAfter=6,
    )
    normal_style = ParagraphStyle(
        "NormalCN", parent=styles["Normal"],
        fontName=font_name, fontSize=10, leading=15,
    )

    story = []

    # 标题
    story.append(Paragraph("学生个人信息表", title_style))
    story.append(Spacer(1, 0.3 * cm))

    # ---- 照片区：左右两张照片 ----
    photo_items = []
    for label, attr in [("高考照片", "gaokao_photo"), ("实地照片", "on_site_photo")]:
        photo_path = getattr(student, attr, None)
        if photo_path:
            # 路径可能是 /static/photos/xxx，实际文件在 photos_dir
            fname = Path(photo_path).name
            full_path = photos_dir / fname
            if full_path.exists():
                try:
                    img = Image(str(full_path), width=4 * cm, height=5.6 * cm, kind="proportional")
                    photo_items.append([img])
                    photo_items.append([Paragraph(label, normal_style)])
                except Exception:
                    photo_items.append([Paragraph("（无照片）", normal_style)])
                    photo_items.append([Paragraph(label, normal_style)])
            else:
                photo_items.append([Paragraph("（无照片）", normal_style)])
                photo_items.append([Paragraph(label, normal_style)])
        else:
            photo_items.append([Paragraph("（无照片）", normal_style)])
            photo_items.append([Paragraph(label, normal_style)])

    if photo_items:
        # 两列：左边基本信息，右边两张照片
        info_data = _build_info_table_data(student, font_name, normal_style)
        info_table = Table(info_data, colWidths=[3 * cm, 5.5 * cm])
        info_table.setStyle(_info_table_style())

        photo_table_data = [[photo_items[0][0], photo_items[2][0]],
                            [photo_items[1][0], photo_items[3][0]]]
        photo_table = Table(photo_table_data, colWidths=[4.2 * cm, 4.2 * cm])
        photo_table.setStyle(TableStyle([
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))

        top_layout = Table([[info_table, photo_table]], colWidths=[9 * cm, 9 * cm])
        top_layout.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(top_layout)
    else:
        # 没有照片时直接放信息表
        info_data = _build_info_table_data(student, font_name, normal_style)
        info_table = Table(info_data, colWidths=[3 * cm, 12 * cm])
        info_table.setStyle(_info_table_style())
        story.append(info_table)

    story.append(Spacer(1, 0.5 * cm))

    # ---- 分组信息 ----
    sections = [
        ("联系方式", [
            ("本人电话", "phone"), ("QQ", "qq"), ("电子邮箱", "email"),
            ("家庭电话", "family_phone"), ("家庭通信地址", "family_address"),
            ("邮政编码", "postal_code"),
        ]),
        ("学习信息", [
            ("考生号", "exam_no"), ("考生类别", "exam_category"),
            ("科类", "subject_type"), ("考试类型", "exam_type"),
            ("录取批次", "admission_batch"), ("高考总分", "gaokao_score"),
            ("毕业中学", "high_school"), ("报到状态", "registration_status"),
            ("报到时间", "registration_time"), ("在校状态", "school_status"),
        ]),
        ("宿舍与职务", [
            ("宿舍", "dormitory"), ("床位", "bed"),
            ("班委职务", "class_position"), ("团学会职务", "league_position"),
            ("是否贫困生", "is_poor"),
        ]),
    ]

    for sec_title, sec_fields in sections:
        story.append(Paragraph(sec_title, h2_style))
        data = []
        row = []
        for i, (label, prop) in enumerate(sec_fields):
            val = _get_student_value(student, prop)
            row.append(Paragraph(f"<b>{label}</b>", normal_style))
            row.append(Paragraph(str(val), normal_style))
            if (i + 1) % 2 == 0:  # 每行两列（标签+值 × 2 组 = 4 列）
                data.append(row)
                row = []
        if row:
            # 奇数个字段，最后一行补空
            row.extend(["", ""])
            data.append(row)

        if data:
            t = Table(data, colWidths=[2.5 * cm, 4.5 * cm, 2.5 * cm, 4.5 * cm])
            t.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D7DE")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F5F7FA")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F5F7FA")),
            ]))
            story.append(t)

    # ---- 备注：单独一个 section，支持 Markdown 渲染 + 附件清单 ----
    remark_text = _get_student_value(student, "remark")
    remark_attachments = _parse_remark_attachments_pdf(student)
    if remark_text or remark_attachments:
        story.append(Paragraph("备注", h2_style))
        remark_content = []
        if remark_text:
            remark_html = _markdown_to_paragraph_html(remark_text)
            remark_content.append([
                Paragraph(remark_html, normal_style),
            ])
        else:
            remark_content.append([Paragraph("（无）", normal_style)])

        # 备注用单列表格，占满整行宽度
        remark_table = Table(remark_content, colWidths=[17 * cm])
        remark_table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D7DE")),
        ]))
        story.append(remark_table)

        # 附件清单
        if remark_attachments:
            story.append(Spacer(1, 0.2 * cm))
            story.append(Paragraph(
                f"附件（共 {len(remark_attachments)} 个）：",
                ParagraphStyle(
                    "AttachTitle", parent=normal_style,
                    fontName=font_name, fontSize=9, textColor=colors.HexColor("#666666"),
                ),
            ))
            for idx, att in enumerate(remark_attachments, 1):
                name = att.get("name", "未知文件")
                size = att.get("size", 0)
                size_str = _format_file_size(size)
                story.append(Paragraph(
                    f"&nbsp;&nbsp;{idx}. {name} ({size_str})",
                    ParagraphStyle(
                        "AttachItem", parent=normal_style,
                        fontName=font_name, fontSize=9, textColor=colors.HexColor("#666666"),
                        leading=14,
                    ),
                ))

    doc.build(story)
    return buf.getvalue()


def _format_file_size(size: int) -> str:
    """把字节数格式化成易读的字符串（PDF 附件清单用）"""
    if size < 1024:
        return f"{size} B"
    if size < 1024 * 1024:
        return f"{size / 1024:.1f} KB"
    return f"{size / 1024 / 1024:.1f} MB"


def _build_info_table_data(student, font_name, style):
    """顶部基本信息表（左侧）"""
    basic_fields = [
        ("学号", "student_no"), ("姓名", "name"), ("性别", "gender"),
        ("出生日期", "birth_date"), ("民族", "ethnicity"),
        ("政治面貌", "political_status"), ("生源地", "hometown"),
        ("身份证号", "id_card"), ("专业", "major"), ("班级", "class_name"),
    ]
    data = []
    for label, prop in basic_fields:
        val = _get_student_value(student, prop)
        data.append([
            Paragraph(f"<b>{label}</b>", style),
            Paragraph(str(val), style),
        ])
    return data


def _info_table_style():
    return TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), _register_cjk_font()),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D7DE")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F5F7FA")),
    ])


def generate_individual_pdfs(
    students: List[Student],
    photos_dir: Path,
    watermark_token: str,
) -> bytes:
    """生成多个学生的个人信息表 PDF，打包成 zip 返回二进制

    每个 PDF 单独打不可见水印（用同一个 token），然后打包 zip。
    zip 内文件名：「学号-姓名-个人信息表.pdf」
    """
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for stu in students:
            pdf_bytes = generate_student_pdf(stu, photos_dir)
            # 打不可见水印
            try:
                pdf_bytes = add_invisible_watermark(pdf_bytes, watermark_token)
            except Exception:
                # 水印失败不影响导出，至少文件要能拿到
                pass
            safe_name = f"{stu.student_no}-{stu.name}-个人信息表.pdf"
            zf.writestr(safe_name, pdf_bytes)
    return zip_buf.getvalue()


# ---------------------------------------------------------------------------
# 工具：按 scope_type + scope_param 从数据库取学生列表
# ---------------------------------------------------------------------------

def resolve_students(db, scope_type: str, scope_param) -> List[Student]:
    """根据范围类型和参数，解析出学生列表

    scope_param 是 dict 或 None：
    - all: 全部学生
    - filtered: { keyword, gender, ethnicity, political_status, dormitory, is_poor }
    - selected: { ids: [1, 2, 3] }
    - id_range: { start_no: "20260000001", end_no: "20260000010" }
    """
    from .student_service import get_students  # 延迟导入，避免循环依赖

    if scope_type == "all":
        result, _ = get_students(db, skip=0, limit=9999)
        return result

    if scope_type == "selected":
        ids = (scope_param or {}).get("ids", [])
        if not ids:
            return []
        return db.query(Student).filter(Student.id.in_(ids)).order_by(Student.student_no).all()

    if scope_type == "id_range":
        start = (scope_param or {}).get("start_no", "")
        end = (scope_param or {}).get("end_no", "")
        q = db.query(Student)
        if start:
            q = q.filter(Student.student_no >= start)
        if end:
            q = q.filter(Student.student_no <= end)
        return q.order_by(Student.student_no).all()

    if scope_type == "filtered":
        params = scope_param or {}
        result, _ = get_students(
            db,
            skip=0,
            limit=9999,
            keyword=params.get("keyword") or None,
            gender=params.get("gender") or None,
            ethnicity=params.get("ethnicity") or None,
            political_status=params.get("political_status") or None,
            dormitory=params.get("dormitory") or None,
            is_poor=params.get("is_poor") or None,
            sort_by=params.get("sort_by") or None,
            sort_order=params.get("sort_order") or "asc",
        )
        return result

    return []
