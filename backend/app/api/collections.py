"""投票收集：班委发起，范围内的人按题型填写

和「投票表决」的区别：收集允许每人多份、六种题型（单选/多选/短文本/长文本/
图片/评分）、竞争名额、管理员删除答卷、按提交时间排序和导出 Excel。
数据从 Collection 开始另起一套，poll_answers 一行不动。

几个贯穿全模块的口径：
1. **越权一律 404**：范围外的人不该知道有这场收集（和投票一致）。
2. **答卷的管理权限**：发起人 + 管理员看详情、提醒、导出；**删除答卷仅管理员**。
3. **匿名只在接口层**：库里存学号用于配额/提醒；管理端详情、导出不返回身份。
4. **竞争名额用原子计数器**：`UPDATE ... WHERE response_count < max_responses`，
   不能用「先 COUNT 再 INSERT」，并发会超收；删除时同事务减一（释放名额）。
"""

import base64
import hashlib
import hmac
import json
import shutil
import uuid
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from fastapi import (
    APIRouter,
    BackgroundTasks,
    Depends,
    File,
    HTTPException,
    Query,
    UploadFile,
)
from fastapi.responses import FileResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import case, or_, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..config import BASE_DIR, resolve_session_secret
from ..database import get_db
from ..models import (
    Collection,
    CollectionAnswer,
    CollectionQuestion,
    CollectionSubmission,
    CollectionTarget,
    Student,
    User,
)
from ..schemas import (
    CollectionAnswerIn,
    CollectionBatchDeleteIn,
    CollectionCloseResponse,
    CollectionCreate,
    CollectionDetailResponse,
    CollectionFileResponse,
    CollectionListResponse,
    CollectionPendingResponse,
    CollectionQuestionIn,
    CollectionResponseList,
    CollectionShareResponse,
    CollectionStatisticsResponse,
    CollectionSubmitIn,
    CollectionSubmitResponse,
    CollectionUpdate,
    CollectionWithdrawResponse,
)
from ..services import account_name, log_operation, notify, recipients
from ..services.collection_service import (
    CHOICE_TYPES,
    IMAGE_MAX_FILES,
    IMAGE_SUFFIXES,
    QTYPE_IMAGE,
    QTYPE_MULTIPLE,
    QTYPE_RATING,
    QTYPE_SINGLE,
    QUESTION_TYPES,
    RATING_MAX,
    RATING_MIN,
    STATUS_CLOSED,
    STATUS_OPEN,
    STATUS_PAUSED,
    STATUS_PENDING,
    TEXT_LIMITS,
    TEXT_TYPES,
    can_manage,
    can_submit,
    cap_reached,
    collection_share_token,
    collection_status,
    detail as collection_detail,
    filter_visible,
    has_submissions,
    is_target,
    list_items,
    my_submission_count,
    parse_collection_share_token,
    pending_rows,
    question_options,
    question_payload,
    questions_of,
    read_files,
    responses_page,
    scope_label,
    statistics,
    submission_payload,
)
from ..services.list_filters import keyword_like
from ..services.share_service import share_path
from ..utils.media_type import guess_media_type
from .deps import (
    admin_required,
    committee_required,
    get_client_ip,
    get_current_user,
    get_optional_user,
    is_admin,
)
# 和投票共用同一套起止时间校验，避免两套「时间先后 / 跨度」口径漂移
from .polls import _check_time

router = APIRouter(prefix="/api/collections", tags=["投票收集"])

MAX_PAGE_SIZE = 100
MAX_FILE_BYTES = 10 * 1024 * 1024
COLLECTION_FILE_DIR = BASE_DIR / "data" / "collection_files"
COLLECTION_FILE_DIR.mkdir(parents=True, exist_ok=True)


# ---------- 取收集 ----------

def _get_collection(db: Session, collection_id: int) -> Collection:
    collection = db.query(Collection).filter(Collection.id == collection_id).first()
    if not collection:
        raise HTTPException(status_code=404, detail="收集不存在")
    return collection


def _load_visible(
    db: Session, collection_id: int, user: User, admin: bool
) -> Collection:
    """看得见 = 在范围内，或者我是发起人 / 管理员；其余一律 404"""
    collection = _get_collection(db, collection_id)
    if not can_manage(collection, user.student_no, admin) and not is_target(
        db, collection_id, user.student_no
    ):
        raise HTTPException(status_code=404, detail="收集不存在")
    return collection


def _load_manageable(
    db: Session, collection_id: int, user: User, admin: bool
) -> Collection:
    """看详情 / 提醒 / 导出 / 关闭：发起人和管理员"""
    collection = _get_collection(db, collection_id)
    if not can_manage(collection, user.student_no, admin):
        raise HTTPException(status_code=404, detail="收集不存在")
    return collection


# ---------- 出卷校验 ----------

def _check_questions(items: list[CollectionQuestionIn]) -> list[dict]:
    """按题型校验并归一化

    - 选择题：至少 2 个选项、非空、不重复；多选 2..选项数
    - 文本/图片/评分：不需要选项；长度、张数、分值范围是常量（见 collection_service）
    """
    normalized = []
    for index, item in enumerate(items, start=1):
        title = (item.title or "").strip()
        if not title:
            raise HTTPException(status_code=400, detail=f"第 {index} 题的题干不能为空")
        qtype = (item.qtype or "").strip()
        if qtype not in QUESTION_TYPES:
            raise HTTPException(status_code=400, detail=f"第 {index} 题的题型不支持")

        options = [option.strip() for option in item.options]
        max_choices = None
        if qtype in CHOICE_TYPES:
            if len(options) < 2:
                raise HTTPException(status_code=400, detail=f"第 {index} 题至少要有两个选项")
            if any(not option for option in options):
                raise HTTPException(status_code=400, detail=f"第 {index} 题有空白选项")
            if len(set(options)) != len(options):
                raise HTTPException(status_code=400, detail=f"第 {index} 题有重复的选项")
            if qtype == QTYPE_SINGLE:
                max_choices = 1
            else:
                max_choices = item.max_choices if item.max_choices is not None else 2
                if not 2 <= max_choices <= len(options):
                    raise HTTPException(
                        status_code=400,
                        detail=f"第 {index} 题「最多选几项」要在 2 到 {len(options)} 之间",
                    )
        else:
            options = []

        normalized.append(
            {
                "title": title,
                "qtype": qtype,
                "required": bool(item.required),
                "options": options,
                "max_choices": max_choices,
            }
        )
    return normalized


def _check_targets(db: Session, targets: list[str]) -> list[Student]:
    unique = list(dict.fromkeys(targets))
    students = (
        db.query(Student)
        .filter(Student.student_no.in_(unique), Student.school_status == "正常")
        .all()
    )
    if len(students) != len(unique):
        raise HTTPException(
            status_code=400,
            detail=f"选中的 {len(unique)} 人里有 {len(unique) - len(students)} 人不存在或不在正常在校状态",
        )
    return students


def _check_settings(submit_quota: int, max_responses: int | None) -> None:
    if max_responses is not None and max_responses < submit_quota:
        raise HTTPException(
            status_code=400,
            detail="竞争总份数不能小于每人可填份数",
        )


# ---------- 图片文件 ----------

def _file_dir(collection_id: int) -> Path:
    path = COLLECTION_FILE_DIR / str(collection_id)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _file_signature(payload: bytes) -> str:
    key = resolve_session_secret().encode("utf-8")
    return hmac.new(key, payload, hashlib.sha256).hexdigest()[:32]


def _file_token(
    collection_id: int,
    student_no: str,
    stored: str,
    original: str,
    size: int,
) -> str:
    """图片上传凭证：绑定收集、学号、临时文件名、原始名和大小

    提交答卷时凭它把临时文件改名成这份答卷专属的文件。**一个凭证只能用一次**：
    提交时是 rename 而不是 copy，第二次拿同一个 token 来会找不到临时文件。
    """
    payload = json.dumps(
        {
            "c": collection_id,
            "s": student_no,
            "stored": stored,
            "original": original,
            "size": size,
        },
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    body = base64.urlsafe_b64encode(payload).decode("ascii").rstrip("=")
    return f"{body}.{_file_signature(payload)}"


def _parse_file_token(token: str, collection_id: int, student_no: str) -> dict | None:
    try:
        body, sig = token.rsplit(".", 1)
        payload = base64.urlsafe_b64decode(body + "=" * (-len(body) % 4))
        if not hmac.compare_digest(_file_signature(payload), sig):
            return None
        data = json.loads(payload)
        if data.get("c") != collection_id or data.get("s") != student_no:
            return None
        if not data.get("stored"):
            return None
        return data
    except (AttributeError, ValueError, TypeError, json.JSONDecodeError):
        return None


def _delete_files(collection_id: int, names: list[str]) -> None:
    """删文件只在数据库 commit 之后做；文件已经不在也算正常"""
    directory = COLLECTION_FILE_DIR / str(collection_id)
    for name in names:
        try:
            (directory / Path(name).name).unlink(missing_ok=True)
        except OSError:
            pass


# ---------- 列表 / 详情 ----------

@router.get("", response_model=CollectionListResponse, summary="收集列表")
def get_list(
    keyword: str | None = Query(None, description="按标题/描述搜索"),
    status: str | None = Query(None, description="pending / open / closed"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=MAX_PAGE_SIZE),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """只列我该看得见的；搜索和状态筛选都在 SQL 里，分页才对"""
    admin = is_admin(user)
    query = filter_visible(db, db.query(Collection), user.student_no, admin)

    match = keyword_like(keyword, Collection.title, Collection.description)
    if match is not None:
        query = query.filter(match)

    now = datetime.now()
    if status == STATUS_PENDING:
        query = query.filter(Collection.closed_at.is_(None), Collection.start_at > now)
    elif status == STATUS_OPEN:
        query = query.filter(
            Collection.closed_at.is_(None),
            Collection.start_at <= now,
            Collection.end_at >= now,
        )
    elif status == STATUS_CLOSED:
        query = query.filter(
            or_(Collection.closed_at.is_not(None), Collection.end_at < now)
        )

    total = query.count()
    collections = (
        query.order_by(Collection.created_at.desc(), Collection.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": list_items(db, collections, user.student_no, admin),
    }


@router.get(
    "/share/{token}/info",
    response_model=CollectionShareResponse,
    summary="收集分享页公开信息（无需登录）",
)
def share_info(
    token: str,
    user: User | None = Depends(get_optional_user),
    db: Session = Depends(get_db),
):
    """公开：标题、完整描述、范围、状态、起止、记名、份数、剩余名额

    题目只在「登录且在范围内」时下发；游客和范围外的人看不到题目。
    不返回 creator_no。
    """
    collection_id = parse_collection_share_token(token)
    collection = (
        db.query(Collection).filter(Collection.id == collection_id).first()
        if collection_id
        else None
    )
    if not collection:
        raise HTTPException(status_code=404, detail="分享链接无效或收集已删除")

    in_target = bool(user and is_target(db, collection.id, user.student_no))
    manage = bool(user and can_manage(collection, user.student_no, is_admin(user)))
    my_count = (
        my_submission_count(db, collection.id, user.student_no) if user else 0
    )
    target_count = (
        db.query(CollectionTarget)
        .filter(CollectionTarget.collection_id == collection.id)
        .count()
    )
    question_count = (
        db.query(CollectionQuestion)
        .filter(CollectionQuestion.collection_id == collection.id)
        .count()
    )
    max_responses = collection.max_responses
    remaining_slots = (
        max(0, max_responses - collection.response_count)
        if max_responses is not None
        else None
    )
    is_full = cap_reached(collection)
    # 自己交过的答卷回显给本人：配额 1 时答过就以「查看自己的答卷」为主；
    # 多份时也能看到已交的每一份，允许撤回时可逐份撤。
    my_responses = []
    if user and in_target:
        submissions = (
            db.query(CollectionSubmission)
            .filter(
                CollectionSubmission.collection_id == collection.id,
                CollectionSubmission.student_no == user.student_no,
            )
            .order_by(
                CollectionSubmission.submitted_at.asc(),
                CollectionSubmission.id.asc(),
            )
            .all()
        )
        my_responses = [
            submission_payload(db, collection, submission)
            for submission in submissions
        ]
    return {
        "id": collection.id,
        "share_token": token,
        "title": collection.title or "",
        "description": collection.description or "",
        "anonymous": bool(collection.anonymous),
        "allow_withdraw": bool(collection.allow_withdraw),
        "status": collection_status(collection),
        "start_at": collection.start_at,
        "end_at": collection.end_at,
        "closed_at": collection.closed_at,
        "scope_label": scope_label(db, collection.id),
        "question_count": question_count,
        "target_count": target_count,
        "submitted_count": collection.response_count,
        "submit_quota": collection.submit_quota,
        "max_responses": max_responses,
        "remaining_slots": remaining_slots,
        "cap_reached": is_full,
        "is_login": user is not None,
        "in_scope": in_target,
        "my_submitted_count": my_count,
        "my_remaining": max(0, collection.submit_quota - my_count),
        "quota_reached": my_count >= collection.submit_quota,
        "can_submit": bool(
            user and can_submit(collection, in_target, my_count)
        ),
        "can_manage": manage,
        "questions": [
            question_payload(question)
            for question in questions_of(db, collection.id)
        ]
        if in_target
        else [],
        "my_responses": my_responses,
    }


@router.get(
    "/{collection_id}",
    response_model=CollectionDetailResponse,
    summary="收集详情（范围内或发起人/管理员）",
)
def get_detail(
    collection_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    collection = _load_visible(db, collection_id, user, is_admin(user))
    return collection_detail(db, collection, user.student_no, is_admin(user))


# ---------- 创建 / 编辑 / 关闭 / 删除 / 延长 ----------

@router.post(
    "", response_model=CollectionDetailResponse, summary="发起收集（班委及以上）"
)
def create_collection(
    data: CollectionCreate,
    user: User = Depends(committee_required),
    db: Session = Depends(get_db),
    ip: str = Depends(get_client_ip),
):
    """创建时不发通知；notify_on_start 决定开始时间到点是否提醒"""
    _check_time(data.start_at, data.end_at)
    questions = _check_questions(data.questions)
    students = _check_targets(db, data.targets)
    max_responses = data.max_responses or None
    _check_settings(data.submit_quota, max_responses)

    collection = Collection(
        title=data.title.strip(),
        description=(data.description or "").strip(),
        anonymous=bool(data.anonymous),
        allow_withdraw=bool(data.allow_withdraw),
        start_at=data.start_at,
        end_at=data.end_at,
        creator_no=user.student_no,
        creator_name=account_name(db, user),
        created_at=datetime.now(),
        submit_quota=data.submit_quota,
        max_responses=max_responses,
        response_count=0,
        notify_on_start=bool(data.notify_on_start),
    )
    db.add(collection)
    db.flush()
    db.add_all(
        [
            CollectionQuestion(
                collection_id=collection.id,
                sort=index,
                title=question["title"],
                qtype=question["qtype"],
                required=question["required"],
                options=json.dumps(question["options"], ensure_ascii=False),
                max_choices=question["max_choices"],
            )
            for index, question in enumerate(questions)
        ]
    )
    db.add_all(
        [
            CollectionTarget(
                collection_id=collection.id,
                student_no=student.student_no,
                student_name=student.name,
            )
            for student in students
        ]
    )
    log_operation(
        db,
        user,
        action="发起收集",
        target=collection.title,
        detail=f"范围 {len(students)} 人，{len(questions)} 题，"
        f"每人 {collection.submit_quota} 份"
        + (f"，限收 {collection.max_responses} 份" if collection.max_responses else ""),
        ip=ip,
    )
    db.commit()
    return collection_detail(db, collection, user.student_no, is_admin(user))


@router.patch(
    "/{collection_id}",
    response_model=CollectionDetailResponse,
    summary="编辑收集（发起人/管理员）",
)
def update_collection(
    collection_id: int,
    data: CollectionUpdate,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    ip: str = Depends(get_client_ip),
):
    """标题/描述随时改；时间按状态分档；题目、必答、份数、记名、开始提醒有答卷后锁定"""
    admin = is_admin(user)
    collection = _load_manageable(db, collection_id, user, admin)
    title = data.title.strip()
    if not title:
        raise HTTPException(status_code=400, detail="标题不能为空")
    description = (data.description or "").strip()

    status = collection_status(collection)
    answered = has_submissions(db, collection.id)

    changed = []
    if title != collection.title:
        changed.append("标题")
    if description != (collection.description or ""):
        changed.append("描述")

    if data.start_at is not None or data.end_at is not None:
        if status == STATUS_CLOSED:
            raise HTTPException(status_code=400, detail="收集已经结束，时间改不了了")
        if data.start_at is not None and status != STATUS_PENDING:
            raise HTTPException(status_code=400, detail="收集已经开始，开始时间改不了了")
        _check_time(
            data.start_at or collection.start_at,
            data.end_at or collection.end_at,
        )

    if data.anonymous is not None:
        if status != STATUS_PENDING or answered:
            raise HTTPException(status_code=400, detail="已经开始或有答卷后，记名方式改不了了")
        if data.anonymous != bool(collection.anonymous):
            collection.anonymous = data.anonymous
            changed.append("记名方式")

    if data.notify_on_start is not None:
        if status != STATUS_PENDING or collection.start_notified_at is not None:
            raise HTTPException(
                status_code=400, detail="开始提醒只能在收集开始前设置"
            )
        if data.notify_on_start != bool(collection.notify_on_start):
            collection.notify_on_start = data.notify_on_start
            changed.append("开始提醒")

    # 撤回设置不影响已有答案的内容，所以不随「有答卷」锁定；
    # 结束后后端会拒绝撤回，这里改不改都不产生实际影响
    if data.allow_withdraw is not None:
        if data.allow_withdraw != bool(collection.allow_withdraw):
            collection.allow_withdraw = data.allow_withdraw
            changed.append("撤回设置")

    if data.submit_quota is not None or data.max_responses is not None:
        if status != STATUS_PENDING or answered:
            raise HTTPException(
                status_code=400, detail="已经开始或有答卷后，份数设置改不了了"
            )
        new_quota = data.submit_quota if data.submit_quota is not None else collection.submit_quota
        new_max = (
            (data.max_responses or None)
            if data.max_responses is not None
            else collection.max_responses
        )
        _check_settings(new_quota, new_max)
        if new_quota != collection.submit_quota:
            collection.submit_quota = new_quota
            changed.append("每人份数")
        if new_max != collection.max_responses:
            collection.max_responses = new_max
            changed.append("竞争份数")

    if data.questions is not None:
        if status != STATUS_PENDING or answered:
            raise HTTPException(status_code=400, detail="已经开始或有答卷后，题目改不了了")
        questions = _check_questions(data.questions)
        db.query(CollectionQuestion).filter(
            CollectionQuestion.collection_id == collection.id
        ).delete(synchronize_session=False)
        db.add_all(
            [
                CollectionQuestion(
                    collection_id=collection.id,
                    sort=index,
                    title=question["title"],
                    qtype=question["qtype"],
                    required=question["required"],
                    options=json.dumps(question["options"], ensure_ascii=False),
                    max_choices=question["max_choices"],
                )
                for index, question in enumerate(questions)
            ]
        )
        changed.append(f"{len(questions)} 道题")

    if data.start_at is not None and data.start_at != collection.start_at:
        collection.start_at = data.start_at
        changed.append("开始时间")
    if data.end_at is not None and data.end_at != collection.end_at:
        collection.end_at = data.end_at
        changed.append("截止时间")

    collection.title = title
    collection.description = description
    log_operation(
        db,
        user,
        action="编辑收集",
        target=title,
        detail="改了" + "、".join(changed) if changed else "内容没变",
        ip=ip,
    )
    db.commit()
    return collection_detail(db, collection, user.student_no, admin)


@router.post(
    "/{collection_id}/close",
    response_model=CollectionCloseResponse,
    summary="立即结束收集（发起人/管理员）",
)
def close_collection(
    collection_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    ip: str = Depends(get_client_ip),
):
    collection = _load_manageable(db, collection_id, user, is_admin(user))
    if collection.closed_at is not None:
        return {"status": STATUS_CLOSED, "response_count": collection.response_count}
    collection.closed_at = datetime.now()
    log_operation(
        db,
        user,
        action="结束收集",
        target=collection.title,
        detail=f"已收 {collection.response_count} 份",
        ip=ip,
    )
    db.commit()
    return {"status": STATUS_CLOSED, "response_count": collection.response_count}


@router.post(
    "/{collection_id}/pause",
    response_model=CollectionDetailResponse,
    summary="暂停收集（发起人/管理员）",
)
def pause_collection(
    collection_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    ip: str = Depends(get_client_ip),
):
    """暂停不是结束：恢复后还能继续填，到点仍照常自动结束；重复调用幂等"""
    collection = _load_manageable(db, collection_id, user, is_admin(user))
    status = collection_status(collection)
    if status == STATUS_PAUSED:
        return collection_detail(db, collection, user.student_no, is_admin(user))
    if status != STATUS_OPEN:
        raise HTTPException(status_code=400, detail="只有进行中的收集才能暂停")
    collection.paused_at = datetime.now()
    log_operation(
        db, user, action="暂停收集", target=collection.title, ip=ip
    )
    db.commit()
    return collection_detail(db, collection, user.student_no, is_admin(user))


@router.post(
    "/{collection_id}/resume",
    response_model=CollectionDetailResponse,
    summary="恢复收集（发起人/管理员）",
)
def resume_collection(
    collection_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    ip: str = Depends(get_client_ip),
):
    """清掉暂停时刻；已经结束的收集清不清都仍是已结束，重复调用幂等"""
    collection = _load_manageable(db, collection_id, user, is_admin(user))
    if collection.paused_at is None:
        return collection_detail(db, collection, user.student_no, is_admin(user))
    collection.paused_at = None
    log_operation(
        db, user, action="恢复收集", target=collection.title, ip=ip
    )
    db.commit()
    return collection_detail(db, collection, user.student_no, is_admin(user))


@router.post(
    "/{collection_id}/extend",
    response_model=CollectionDetailResponse,
    summary="延长收集截止时间 5 分钟（发起人/管理员）",
)
def extend_collection(
    collection_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    ip: str = Depends(get_client_ip),
):
    """和签到的「延长5分钟」同一交互：截止前可反复点；到点/结束后不能再延"""
    collection = _load_manageable(db, collection_id, user, is_admin(user))
    if collection.closed_at is not None:
        raise HTTPException(status_code=400, detail="收集已经结束")
    now = datetime.now()
    if now >= collection.end_at:
        raise HTTPException(status_code=400, detail="已经过了截止时间，不能再延长")

    before = collection.end_at
    collection.end_at = before + timedelta(minutes=5)
    log_operation(
        db,
        user,
        action="延长收集",
        target=collection.title,
        detail=f"截止 {before:%m-%d %H:%M} → {collection.end_at:%m-%d %H:%M}",
        ip=ip,
    )
    db.commit()
    return collection_detail(db, collection, user.student_no, is_admin(user))


@router.delete("/{collection_id}", summary="删除收集（发起人/管理员）")
def remove_collection(
    collection_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    ip: str = Depends(get_client_ip),
):
    collection = _load_manageable(db, collection_id, user, is_admin(user))
    submission_ids = [
        submission_id
        for (submission_id,) in db.query(CollectionSubmission.id)
        .filter(CollectionSubmission.collection_id == collection.id)
        .all()
    ]
    if submission_ids:
        db.query(CollectionAnswer).filter(
            CollectionAnswer.submission_id.in_(submission_ids)
        ).delete(synchronize_session=False)
        db.query(CollectionSubmission).filter(
            CollectionSubmission.collection_id == collection.id
        ).delete(synchronize_session=False)
    db.query(CollectionQuestion).filter(
        CollectionQuestion.collection_id == collection.id
    ).delete(synchronize_session=False)
    db.query(CollectionTarget).filter(
        CollectionTarget.collection_id == collection.id
    ).delete(synchronize_session=False)
    log_operation(db, user, action="删除收集", target=collection.title, ip=ip)
    db.delete(collection)
    db.commit()
    shutil.rmtree(COLLECTION_FILE_DIR / str(collection.id), ignore_errors=True)
    return {"ok": True}


# ---------- 提交答卷 ----------

def _answer_is_empty(qtype: str, answer: CollectionAnswerIn | None) -> bool:
    if answer is None:
        return True
    if qtype in CHOICE_TYPES:
        return not answer.choices
    if qtype in TEXT_TYPES:
        return not (answer.text or "").strip()
    if qtype == QTYPE_IMAGE:
        return not answer.files
    if qtype == QTYPE_RATING:
        return answer.rating is None
    return True


@router.post(
    "/{collection_id}/responses",
    response_model=CollectionSubmitResponse,
    summary="提交一份答卷",
)
def submit_response(
    collection_id: int,
    data: CollectionSubmitIn,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """一次提交 = 一份答卷；每人最多 submit_quota 份，全局最多 max_responses 份"""
    collection = _load_visible(db, collection_id, user, is_admin(user))
    status = collection_status(collection)
    if status == STATUS_CLOSED:
        raise HTTPException(status_code=400, detail="这次收集已经结束")
    if status == STATUS_PENDING:
        raise HTTPException(status_code=400, detail="收集还没开始")
    if status == STATUS_PAUSED:
        raise HTTPException(status_code=400, detail="收集已暂停，等发起人恢复后才能填写")
    if not is_target(db, collection.id, user.student_no):
        raise HTTPException(status_code=400, detail="你不在本次收集的范围内")

    count = my_submission_count(db, collection.id, user.student_no)
    if count >= collection.submit_quota:
        raise HTTPException(
            status_code=400,
            detail=f"每人最多填写 {collection.submit_quota} 份，你已经填满了",
        )
    if cap_reached(collection):
        raise HTTPException(status_code=400, detail="名额已满，这场收集不再接收答卷")

    questions = questions_of(db, collection.id)
    question_map = {question.id: question for question in questions}
    answers_by_question: dict[int, CollectionAnswerIn] = {}
    for answer in data.answers:
        if answer.question_id in answers_by_question:
            raise HTTPException(status_code=400, detail="同一道题提交了两次")
        if answer.question_id not in question_map:
            raise HTTPException(status_code=400, detail="提交了不属于本场收集的题目")
        answers_by_question[answer.question_id] = answer

    prepared = []
    any_answered = False
    for index, question in enumerate(questions, start=1):
        qtype = question.qtype or QTYPE_SINGLE
        answer = answers_by_question.get(question.id)
        empty = _answer_is_empty(qtype, answer)
        if empty:
            if question.required:
                raise HTTPException(status_code=400, detail=f"第 {index} 题是必答题")
            continue
        any_answered = True
        options = question_options(question)

        if qtype == QTYPE_SINGLE:
            choices = answer.choices
            if len(choices) != 1:
                raise HTTPException(status_code=400, detail=f"第 {index} 题是单选题，只能选一项")
            if not 0 <= choices[0] < len(options):
                raise HTTPException(status_code=400, detail=f"第 {index} 题的选项不存在")
            prepared.append(
                {"question": question, "choices": choices, "text": "", "rating": None, "files": []}
            )
        elif qtype == QTYPE_MULTIPLE:
            choices = answer.choices
            if len(set(choices)) != len(choices):
                raise HTTPException(status_code=400, detail=f"第 {index} 题有重复的选项")
            if any(not 0 <= choice < len(options) for choice in choices):
                raise HTTPException(status_code=400, detail=f"第 {index} 题的选项不存在")
            max_choices = question.max_choices or 2
            if len(choices) > max_choices:
                raise HTTPException(
                    status_code=400, detail=f"第 {index} 题最多选 {max_choices} 项"
                )
            prepared.append(
                {"question": question, "choices": choices, "text": "", "rating": None, "files": []}
            )
        elif qtype in TEXT_TYPES:
            text = (answer.text or "").strip()
            limit = TEXT_LIMITS[qtype]
            if len(text) > limit:
                raise HTTPException(
                    status_code=400, detail=f"第 {index} 题最多 {limit} 字"
                )
            prepared.append(
                {"question": question, "choices": [], "text": text, "rating": None, "files": []}
            )
        elif qtype == QTYPE_RATING:
            rating = answer.rating
            if rating is None or not RATING_MIN <= rating <= RATING_MAX:
                raise HTTPException(
                    status_code=400,
                    detail=f"第 {index} 题的评分要在 {RATING_MIN}–{RATING_MAX} 之间",
                )
            prepared.append(
                {"question": question, "choices": [], "text": "", "rating": rating, "files": []}
            )
        elif qtype == QTYPE_IMAGE:
            if len(answer.files) > IMAGE_MAX_FILES:
                raise HTTPException(
                    status_code=400, detail=f"第 {index} 题最多 {IMAGE_MAX_FILES} 张图片"
                )
            payloads = []
            for token in answer.files:
                payload = _parse_file_token(token, collection.id, user.student_no)
                if not payload:
                    raise HTTPException(
                        status_code=400, detail=f"第 {index} 题有图片已失效，请重新上传"
                    )
                source = _file_dir(collection.id) / Path(payload["stored"]).name
                if not source.is_file():
                    raise HTTPException(
                        status_code=400, detail=f"第 {index} 题有图片已失效，请重新上传"
                    )
                payloads.append(payload)
            prepared.append(
                {"question": question, "choices": [], "text": "", "rating": None, "file_payloads": payloads}
            )
        else:
            raise HTTPException(status_code=400, detail=f"第 {index} 题的题型不支持")

    if not any_answered:
        raise HTTPException(status_code=400, detail="至少回答一道题")

    # 竞争名额：原子条件更新，先 COUNT 再 INSERT 会超收
    result = db.execute(
        update(Collection)
        .where(
            Collection.id == collection.id,
            or_(
                Collection.max_responses.is_(None),
                Collection.response_count < Collection.max_responses,
            ),
        )
        .values(response_count=Collection.response_count + 1)
    )
    if result.rowcount == 0:
        db.rollback()
        raise HTTPException(status_code=400, detail="名额已满，这场收集不再接收答卷")

    moved: list[Path] = []
    try:
        submission = CollectionSubmission(
            collection_id=collection.id,
            student_no=user.student_no,
            seq_no=count + 1,
            submitted_at=datetime.now(),
        )
        db.add(submission)
        db.flush()

        for item in prepared:
            file_meta = []
            for payload in item.get("file_payloads", []):
                source = _file_dir(collection.id) / Path(payload["stored"]).name
                suffix = source.suffix.lower()
                dest_name = f"resp-{uuid.uuid4().hex}{suffix}"
                destination = source.with_name(dest_name)
                source.rename(destination)
                moved.append(destination)
                file_meta.append(
                    {
                        "name": dest_name,
                        "original": payload.get("original") or dest_name,
                        "size": int(payload.get("size") or 0),
                    }
                )

            db.add(
                CollectionAnswer(
                    submission_id=submission.id,
                    question_id=item["question"].id,
                    choices=json.dumps(item["choices"]),
                    text_value=item["text"],
                    rating=item["rating"],
                    files=json.dumps(file_meta, ensure_ascii=False),
                )
            )
        db.commit()
    except IntegrityError:
        db.rollback()
        for path in moved:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass
        raise HTTPException(status_code=409, detail="提交冲突，请重试") from None
    except HTTPException:
        db.rollback()
        for path in moved:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass
        raise
    except Exception:
        db.rollback()
        for path in moved:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass
        raise

    db.refresh(collection)
    my_count = count + 1
    return {
        "id": submission.id,
        "seq_no": submission.seq_no,
        "submitted_at": submission.submitted_at,
        "my_submitted_count": my_count,
        "my_remaining": max(0, collection.submit_quota - my_count),
        "can_submit": can_submit(collection, True, my_count),
    }


# ---------- 图片上传 / 读取 ----------

@router.post(
    "/{collection_id}/files",
    response_model=CollectionFileResponse,
    summary="上传答卷图片（返回提交凭证）",
)
def upload_file(
    collection_id: int,
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    collection = _load_visible(db, collection_id, user, is_admin(user))
    if collection_status(collection) != STATUS_OPEN:
        raise HTTPException(status_code=400, detail="现在不能上传图片")
    if not is_target(db, collection.id, user.student_no):
        raise HTTPException(status_code=400, detail="你不在本次收集的范围内")
    count = my_submission_count(db, collection.id, user.student_no)
    if count >= collection.submit_quota:
        raise HTTPException(status_code=400, detail="你已经填满了")
    if cap_reached(collection):
        raise HTTPException(status_code=400, detail="名额已满，这场收集不再接收答卷")

    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in IMAGE_SUFFIXES:
        raise HTTPException(status_code=400, detail="只支持 JPG / PNG / WEBP 图片")
    data = file.file.read(MAX_FILE_BYTES + 1)
    if not data:
        raise HTTPException(status_code=400, detail="图片是空文件")
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(status_code=400, detail="图片不能超过 10MB")

    stored = f"tmp-{uuid.uuid4().hex}{suffix}"
    (_file_dir(collection.id) / stored).write_bytes(data)
    original = Path(file.filename or stored).name
    return {
        "token": _file_token(collection.id, user.student_no, stored, original, len(data)),
        "name": original,
        "size": len(data),
    }


@router.get(
    "/{collection_id}/responses/{response_id}/files/{index}",
    summary="读取答卷图片（发起人/管理员）",
)
def get_response_file(
    collection_id: int,
    response_id: int,
    index: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """发起人/管理员可看；答题人也能看**自己**答卷里的图片（回显用）"""
    collection = _get_collection(db, collection_id)
    submission = (
        db.query(CollectionSubmission)
        .filter(
            CollectionSubmission.id == response_id,
            CollectionSubmission.collection_id == collection.id,
        )
        .first()
    )
    # 越权一律 404，和答卷列表同一口径
    authorized = submission and (
        submission.student_no == user.student_no
        or can_manage(collection, user.student_no, is_admin(user))
    )
    if not authorized:
        raise HTTPException(status_code=404, detail="图片不存在")
    answers = {
        answer.question_id: answer
        for answer in db.query(CollectionAnswer)
        .filter(CollectionAnswer.submission_id == submission.id)
        .all()
    }
    # 按题目顺序展开，和答卷列表里的 files 顺序保持一致 —— 前端用同一个下标拼 URL
    for question in questions_of(db, collection.id):
        files = read_files(answers.get(question.id))
        if 0 <= index < len(files):
            meta = files[index]
            path = COLLECTION_FILE_DIR / str(collection.id) / Path(meta["name"]).name
            if not path.is_file():
                raise HTTPException(status_code=404, detail="图片不存在")
            return FileResponse(
                path,
                media_type=guess_media_type(meta.get("original"), meta["name"]),
                filename=meta.get("original") or meta["name"],
                headers={"Cache-Control": "no-store"},
            )
        index -= len(files)
    raise HTTPException(status_code=404, detail="图片不存在")


# ---------- 管理端答卷 ----------

@router.get(
    "/{collection_id}/responses",
    response_model=CollectionResponseList,
    summary="答卷列表（发起人/管理员，按提交时间排序）",
)
def get_responses(
    collection_id: int,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=MAX_PAGE_SIZE),
    order: str = Query("asc", pattern="^(asc|desc)$"),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    collection = _load_manageable(db, collection_id, user, is_admin(user))
    total, items = responses_page(db, collection, page, page_size, order)
    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "order": order,
        "items": items,
    }


@router.get(
    "/{collection_id}/pending",
    response_model=CollectionPendingResponse,
    summary="未填名单（发起人/管理员）",
)
def get_pending(
    collection_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """按「提交份数 < 每人配额」算人；匿名时只返回是否未完成，不返回份数"""
    collection = _load_manageable(db, collection_id, user, is_admin(user))
    rows = [
        row
        for row in pending_rows(db, collection)
        if row["student_no"] != user.student_no
    ]
    items = []
    for row in rows:
        items.append(
            {
                "student_no": row["student_no"],
                "name": row["name"],
                "submitted_count": None if collection.anonymous else row["submitted_count"],
                "remaining": None if collection.anonymous else row["remaining"],
                "frozen": row["frozen"],
            }
        )
    return {"total": len(items), "items": items}


@router.delete(
    "/{collection_id}/responses/{response_id}",
    summary="删除一份答卷（仅管理员）",
)
def delete_response(
    collection_id: int,
    response_id: int,
    user: User = Depends(admin_required),
    db: Session = Depends(get_db),
    ip: str = Depends(get_client_ip),
):
    """销毁数据：只有系统管理员能删；删除后竞争名额释放"""
    collection = _get_collection(db, collection_id)
    submission = (
        db.query(CollectionSubmission)
        .filter(
            CollectionSubmission.id == response_id,
            CollectionSubmission.collection_id == collection.id,
        )
        .first()
    )
    if not submission:
        raise HTTPException(status_code=404, detail="答卷不存在")

    file_names = _files_of_submission(db, submission.id)
    db.query(CollectionAnswer).filter(
        CollectionAnswer.submission_id == submission.id
    ).delete(synchronize_session=False)
    db.delete(submission)
    db.execute(
        update(Collection)
        .where(Collection.id == collection.id)
        .values(
            response_count=case(
                (Collection.response_count >= 1, Collection.response_count - 1),
                else_=0,
            )
        )
    )
    log_operation(
        db,
        user,
        action="删除答卷",
        target=collection.title,
        detail=f"答卷 #{response_id}",
        ip=ip,
    )
    db.commit()
    _delete_files(collection.id, file_names)
    db.refresh(collection)
    return {"deleted": 1, "response_count": collection.response_count}


@router.post(
    "/{collection_id}/responses/{response_id}/withdraw",
    response_model=CollectionWithdrawResponse,
    summary="答题人撤回自己的答卷",
)
def withdraw_response(
    collection_id: int,
    response_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    ip: str = Depends(get_client_ip),
):
    """出卷人允许时，本人可撤回；竞争名额同事务释放，撤回后还能再填

    只有进行中/暂停中能撤：结束后答卷定稿，撤回答卷等于改最终数据。
    """
    collection = _get_collection(db, collection_id)

    # 先确认「这份答卷是我的」再谈规则：不是自己的统一 404，不暴露收集的设置
    submission = (
        db.query(CollectionSubmission)
        .filter(
            CollectionSubmission.id == response_id,
            CollectionSubmission.collection_id == collection.id,
            CollectionSubmission.student_no == user.student_no,
        )
        .first()
    )
    if not submission:
        raise HTTPException(status_code=404, detail="答卷不存在")
    if not collection.allow_withdraw:
        raise HTTPException(status_code=400, detail="这场收集不允许撤回答卷")
    if collection_status(collection) not in (STATUS_OPEN, STATUS_PAUSED):
        raise HTTPException(status_code=400, detail="现在不能撤回答卷")

    file_names = _files_of_submission(db, submission.id)
    db.query(CollectionAnswer).filter(
        CollectionAnswer.submission_id == submission.id
    ).delete(synchronize_session=False)
    db.delete(submission)
    # 和删除答卷同一套：撤回也释放竞争名额
    db.execute(
        update(Collection)
        .where(Collection.id == collection.id)
        .values(
            response_count=case(
                (Collection.response_count >= 1, Collection.response_count - 1),
                else_=0,
            )
        )
    )
    log_operation(
        db,
        user,
        action="撤回答卷",
        target=collection.title,
        detail=f"答卷 #{response_id}",
        ip=ip,
    )
    db.commit()
    _delete_files(collection.id, file_names)
    db.refresh(collection)

    my_count = my_submission_count(db, collection.id, user.student_no)
    in_target = is_target(db, collection.id, user.student_no)
    return {
        "deleted": 1,
        "response_count": collection.response_count,
        "my_submitted_count": my_count,
        "can_submit": bool(can_submit(collection, in_target, my_count)),
    }


@router.post(
    "/{collection_id}/responses/batch-delete",
    summary="批量删除答卷（仅管理员）",
)
def batch_delete_responses(
    collection_id: int,
    data: CollectionBatchDeleteIn,
    user: User = Depends(admin_required),
    db: Session = Depends(get_db),
    ip: str = Depends(get_client_ip),
):
    collection = _get_collection(db, collection_id)
    submissions = (
        db.query(CollectionSubmission)
        .filter(
            CollectionSubmission.collection_id == collection.id,
            CollectionSubmission.id.in_(data.response_ids),
        )
        .all()
    )
    if not submissions:
        raise HTTPException(status_code=404, detail="没有可删除的答卷")
    ids = [submission.id for submission in submissions]
    file_names = []
    for submission_id in ids:
        file_names.extend(_files_of_submission(db, submission_id))

    db.query(CollectionAnswer).filter(
        CollectionAnswer.submission_id.in_(ids)
    ).delete(synchronize_session=False)
    db.query(CollectionSubmission).filter(
        CollectionSubmission.id.in_(ids)
    ).delete(synchronize_session=False)
    count = len(ids)
    db.execute(
        update(Collection)
        .where(Collection.id == collection.id)
        .values(
            response_count=case(
                (Collection.response_count >= count, Collection.response_count - count),
                else_=0,
            )
        )
    )
    log_operation(
        db,
        user,
        action="批量删除答卷",
        target=collection.title,
        detail=f"删除 {count} 份",
        ip=ip,
    )
    db.commit()
    _delete_files(collection.id, file_names)
    db.refresh(collection)
    return {"deleted": count, "response_count": collection.response_count}


def _files_of_submission(db: Session, submission_id: int) -> list[str]:
    answers = (
        db.query(CollectionAnswer)
        .filter(CollectionAnswer.submission_id == submission_id)
        .all()
    )
    names = []
    for answer in answers:
        names.extend(meta["name"] for meta in read_files(answer))
    return names


@router.get(
    "/{collection_id}/statistics",
    response_model=CollectionStatisticsResponse,
    summary="收集统计（发起人/管理员）",
)
def get_statistics(
    collection_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    collection = _load_manageable(db, collection_id, user, is_admin(user))
    return {"items": statistics(db, collection)}


@router.post(
    "/{collection_id}/remind",
    summary="提醒还没填满的人（发起人/管理员）",
)
def remind_collection(
    collection_id: int,
    background: BackgroundTasks,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    ip: str = Depends(get_client_ip),
):
    collection = _load_manageable(db, collection_id, user, is_admin(user))
    if collection_status(collection) != STATUS_OPEN:
        raise HTTPException(status_code=400, detail="只有进行中的收集才能提醒")
    if cap_reached(collection):
        raise HTTPException(status_code=400, detail="名额已满，不用再提醒了")

    pending = [
        row
        for row in pending_rows(db, collection)
        if not row["frozen"] and row["student_no"] != user.student_no
    ]
    page = share_path("collect", collection_share_token(collection.id))
    for row in pending:
        if collection.anonymous:
            body = f"还未填写，{collection.end_at:%m-%d %H:%M} 截止"
        else:
            body = (
                f"已填 {row['submitted_count']}/{collection.submit_quota} 份，"
                f"还差 {row['remaining']} 份，{collection.end_at:%m-%d %H:%M} 截止"
            )
        notify(
            db,
            background,
            recipients(db, only_no=row["student_no"]),
            f"提醒填写：{collection.title}",
            body,
            level="timeSensitive",
            page=page,
        )

    log_operation(
        db,
        user,
        action="提醒填写",
        target=collection.title,
        detail=f"已提醒 {len(pending)} 人",
        ip=ip,
    )
    db.commit()
    return {"reminded": len(pending)}


# ---------- Excel 导出 ----------

def _excel_text(value) -> str:
    """文本写进 Excel 前防公式注入：以 = + - @ 开头的当成文本

    openpyxl 会把以 = 开头的字符串写成公式；学生填的文本直接进表，
    不处理的话班委一打开表就可能执行对方写的公式。
    """
    text = "" if value is None else str(value)
    if text[:1] in ("=", "+", "-", "@"):
        return "'" + text
    return text


def _answer_text(question, answer: dict | None) -> str:
    """把一份答案格式化成单元格文本；没答的选答题返回空串"""
    if answer is None:
        return ""
    qtype = question.qtype or QTYPE_SINGLE
    if qtype in CHOICE_TYPES:
        options = question_options(question)
        picked = answer.get("choices") or []
        return "；".join(
            f"{chr(65 + index)}. {options[index]}"
            for index in picked
            if 0 <= index < len(options)
        )
    if qtype in TEXT_TYPES:
        return answer.get("text") or ""
    if qtype == QTYPE_RATING:
        return answer.get("rating") if answer.get("rating") is not None else ""
    if qtype == QTYPE_IMAGE:
        return "、".join(meta.get("original") or meta.get("name") for meta in answer.get("files") or [])
    return ""


def _safe_filename(title: str) -> str:
    cleaned = "".join(
        "_" if char in '/\\:*?"<>|' or char.isspace() else char for char in (title or "")
    ).strip("_")
    return cleaned[:40] or "收集"


@router.get(
    "/{collection_id}/responses/export",
    summary="导出答卷 Excel（发起人/管理员）",
)
def export_responses(
    collection_id: int,
    order: str = Query("asc", pattern="^(asc|desc)$"),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    collection = _load_manageable(db, collection_id, user, is_admin(user))
    questions = questions_of(db, collection.id)
    _, responses = responses_page(db, collection, 1, 100000, order)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "答卷"

    header = ["序号", "提交时间"]
    if not collection.anonymous:
        header += ["姓名", "学号"]
    header += [question.title or "" for question in questions]
    sheet.append(header)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, response in enumerate(responses, start=1):
        answers = {answer["question_id"]: answer for answer in response["answers"]}
        values = [row_index, response["submitted_at"].strftime("%Y-%m-%d %H:%M:%S")]
        if not collection.anonymous:
            values += [
                _excel_text(response.get("student_name") or ""),
                _excel_text(response.get("student_no") or ""),
            ]
        for question in questions:
            value = _answer_text(question, answers.get(question.id))
            values.append(value if isinstance(value, int) else _excel_text(value))
        sheet.append(values)

    widths = [8, 20] + ([] if collection.anonymous else [12, 16]) + [
        min(40, max(12, len(question.title or "") * 2 + 4)) for question in questions
    ]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"

    # 第二张表：未填名单。匿名时只列谁还没完成，不列份数
    pending_sheet = workbook.create_sheet("未填名单")
    pending_header = ["姓名", "学号"] + ([] if collection.anonymous else ["已填份数", "还差份数"])
    pending_sheet.append(pending_header)
    for cell in pending_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in pending_rows(db, collection):
        values = [_excel_text(row["name"]), _excel_text(row["student_no"])]
        if not collection.anonymous:
            values += [row["submitted_count"], row["remaining"]]
        pending_sheet.append(values)
    for index, width in enumerate([14, 16, 12, 12][: len(pending_header)], start=1):
        pending_sheet.column_dimensions[get_column_letter(index)].width = width
    pending_sheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    filename = f"收集_{_safe_filename(collection.title)}_{datetime.now():%Y%m%d}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=utf-8''{quote(filename, safe='')}"
        },
    )
