"""我的日程模块测试

覆盖：CRUD、重复日程展开、导入解析、iCal 生成、权限隔离、通知默认关闭

cd backend && python -m pytest tests/test_schedules.py -q
"""
import re
from io import BytesIO
from datetime import date, datetime, time
from typing import Optional

import httpx
import openpyxl
import pytest

PW = "initpass123"

# 3 个学生：甲（管理员）、乙（普通）、丙（普通）
ROSTER = [
    ("2024000001", "甲", "010001"),  # 管理员
    ("2024000002", "乙", "010002"),  # 普通学生
    ("2024000003", "丙", "010003"),  # 普通学生
]
ADMIN_NO = ROSTER[0][0]


# ---------- fixture ----------

@pytest.fixture
def b(backend):
    """初始化完毕、管理员已登录的后端"""
    inst = backend(db="schedules.db", BOOTSTRAP_PASSWORD=PW)
    assert inst.login(password=PW).status_code == 200

    # 建 3 个学生
    for no, name, id_tail in ROSTER:
        r = inst.client.post("/api/students", json={
            "student_no": no, "name": name,
            "id_card": f"11010120060101{id_tail}",
        })
        assert r.status_code == 200, r.text

    # 指认甲为管理员
    r = inst.client.post("/api/setup/finish", json={"admin_student_no": ADMIN_NO})
    assert r.status_code == 200, r.text

    # 用管理员登录
    inst.logout()
    r = inst.login(ADMIN_NO, ROSTER[0][2])
    assert r.status_code == 200, r.text
    return inst


def as_student(b, idx):
    """以第 idx 个学生身份登录的新客户端"""
    no, _, tail = ROSTER[idx]
    c = httpx.Client(base_url=str(b.client.base_url), timeout=10)
    r = c.post("/api/auth/login", json={"student_no": no, "password": tail})
    assert r.status_code == 200, r.text
    return c


def _make_schedule(title="测试日程", start="2026-10-12T08:00:00",
                   end="2026-10-12T09:00:00", **kw):
    """构造日程创建参数"""
    return {
        "title": title,
        "start_time": start,
        "end_time": end,
        "location": "",
        "description": "",
        "all_day": False,
        "repeat_type": "none",
        "repeat_end_date": None,
        "notify_enabled": False,
        "notify_minutes_before": 10,
        "target_all": False,
        "target_student_ids": [],
        **kw,
    }


# ---------- 基础 CRUD ----------

def test_create_and_list(b):
    """创建单次日程，列表能查到"""
    data = _make_schedule(title="晨跑", start="2026-10-12T07:00:00", end="2026-10-12T07:30:00")
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 200, r.text
    assert r.json()["title"] == "晨跑"

    # 列表查询
    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-18",
    })
    assert r.status_code == 200, r.text
    items = r.json()
    assert len(items) == 1
    assert items[0]["title"] == "晨跑"
    assert items[0]["repeat_type"] == "none"


def test_search_by_title(b):
    """按名称搜自己的日程：同名很多条就一行一条列出来，各带一个具体日期

    页面上把这些结果放在弹窗里让用户挑，所以 instance_date 必须准 —— 学生搜
    「晚自习」看到两个日期，点哪个就跳到哪一天。
    """
    # 管理员自己也有一条同名日程，学生那边不该搜到
    r = b.client.post("/api/schedules", json=_make_schedule(
        title="晚自习", start="2026-10-15T19:00:00", end="2026-10-15T21:00:00"))
    assert r.status_code == 200, r.text

    c = as_student(b, 1)  # 乙
    for day in ("2026-10-12", "2026-10-13"):
        r = c.post("/api/schedules", json=_make_schedule(
            title="晚自习", start=f"{day}T19:00:00", end=f"{day}T21:00:00"))
        assert r.status_code == 200, r.text
    r = c.post("/api/schedules", json=_make_schedule(
        title="高等数学", start="2026-10-14T08:00:00", end="2026-10-14T09:40:00"))
    assert r.status_code == 200, r.text

    r = c.get("/api/schedules/search", params={"keyword": "晚自习"})
    assert r.status_code == 200, r.text
    rows = r.json()
    assert [row["instance_date"] for row in rows] == ["2026-10-12", "2026-10-13"], \
        "两条同名的都出来，按日期排；管理员那条不在里面"
    assert {row["title"] for row in rows} == {"晚自习"}

    # 片段匹配，不要求从头开始
    assert [row["title"] for row in c.get(
        "/api/schedules/search", params={"keyword": "数学"}).json()] == ["高等数学"]

    # 空关键词不回整屏日程（弹窗刚打开时不该把全部日程都拉出来）
    assert c.get("/api/schedules/search", params={"keyword": ""}).json() == []
    assert c.get("/api/schedules/search", params={"keyword": "  "}).json() == []
    assert c.get("/api/schedules/search", params={"keyword": "查无此程"}).json() == []


def test_notify_default_off(b):
    """通知默认关闭——建日程时不传 notify_enabled，必须是 False"""
    data = _make_schedule(title="默认通知测试", notify_enabled=None)
    del data["notify_enabled"]  # 完全不传
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 200
    assert r.json()["notify_enabled"] is False

    # 查库确认
    rows = b.rows("SELECT notify_enabled FROM schedules WHERE title = ?", "默认通知测试")
    assert rows[0][0] == 0  # SQLite 的 False 是 0


def test_update_and_delete(b):
    """编辑和删除单次日程"""
    # 建
    data = _make_schedule(title="原标题")
    r = b.client.post("/api/schedules", json=data)
    sid = r.json()["id"]
    inst_date = r.json()["instance_date"]

    # 改
    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": inst_date},
        json={"title": "新标题", "scope": "all"},
    )
    assert r.status_code == 200, r.text

    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-18",
    })
    assert r.json()[0]["title"] == "新标题"

    # 删
    r = b.client.delete(
        f"/api/schedules/item/{sid}",
        params={"instance_date": inst_date, "scope": "all"},
    )
    assert r.status_code == 200

    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-18",
    })
    assert len(r.json()) == 0


def test_end_before_start_rejected(b):
    """结束时间早于开始时间 → 400"""
    data = _make_schedule(start="2026-10-12T10:00:00", end="2026-10-12T09:00:00")
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 400


def test_timezone_aware_input_normalized(b):
    """带时区的时间（脚本 / PAT 客户端会这么传）换算成本地墙上时间

    以前一端带时区一端不带时，接口里比 end <= start 直接 TypeError → 500；
    两端都带时区则偏移被静默丢掉。容器里 TZ 是 Asia/Shanghai，所以 00:00Z 就是 08:00。
    """
    # 混着传：以前这一条 500
    data = _make_schedule(start="2026-10-12T08:00:00", end="2026-10-12T09:00:00+08:00")
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 200, r.text

    # 两端都带 Z：按本地墙上时间落库，不再是「偏移直接消失」
    data = _make_schedule(title="UTC时间", start="2026-10-13T00:00:00Z", end="2026-10-13T01:00:00Z")
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 200, r.text
    assert r.json()["start_time"] == "2026-10-13T08:00:00"


# ---------- 全天事件（学生自己建的，不是导入的） ----------

def test_all_day_create_drops_times(b):
    """全天事件不带时刻：传进来的 08:00–09:00 一律归到 00:00–23:59

    表单里全天事件只有日期、没有时间框，这里是直接调接口的那条路径。
    """
    data = _make_schedule(
        title="国庆放假",
        start="2026-10-05T08:00:00", end="2026-10-05T09:00:00",
        all_day=True,
    )
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["all_day"] is True
    assert body["start_time"] == "2026-10-05T00:00:00"
    assert body["end_time"] == "2026-10-05T23:59:00"

    # 查库确认存的也是这一天两头，不是传进来的那个时刻
    rows = b.rows("SELECT start_time, end_time FROM schedules WHERE title = ?", "国庆放假")
    assert rows[0][0].startswith("2026-10-05 00:00")
    assert rows[0][1].startswith("2026-10-05 23:59")


def test_all_day_multi_day_is_one_row_per_day(b):
    """跨天的全天事件是「一天一条」—— 前端按天建、这里逐条落库

    模型里没有跨天日程这个概念：`_to_instance` 只取时刻、日期用 instance_date，
    所以一条记 10-05~10-07 的记录在日历上只会出现在 10-05（iCal 导出倒是能表达
    跨天，但列表路径表达不了）。页面上那根连着的条是前端合出来的（allDayBars）。
    """
    for day in ("2026-10-05", "2026-10-06", "2026-10-07"):
        r = b.client.post("/api/schedules", json=_make_schedule(
            title="国庆放假", start=f"{day}T00:00:00", end=f"{day}T23:59:00", all_day=True,
        ))
        assert r.status_code == 200, r.text

    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-05", "end_date": "2026-10-11",
    })
    items = r.json()
    assert [it["instance_date"] for it in items] == ["2026-10-05", "2026-10-06", "2026-10-07"]
    assert all(it["all_day"] and it["title"] == "国庆放假" for it in items)


def test_all_day_update_also_drops_times(b):
    """编辑成全天事件：时刻同样归到当天两头（前端会把日期一起传上来）"""
    r = b.client.post("/api/schedules", json=_make_schedule(title="大扫除"))
    sid = r.json()["id"]

    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": r.json()["instance_date"]},
        json={
            "all_day": True,
            "start_time": "2026-10-12T08:00:00",
            "end_time": "2026-10-12T09:00:00",
            "scope": "all",
        },
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["all_day"] is True
    assert body["start_time"] == "2026-10-12T00:00:00"
    assert body["end_time"] == "2026-10-12T23:59:00"


# ---------- 重复日程 ----------

def test_weekly_repeat_expands(b):
    """每周重复日程在列表中展开为多条"""
    data = _make_schedule(
        title="每周例会",
        start="2026-10-12T14:00:00",  # 周一
        end="2026-10-12T15:00:00",
        repeat_type="weekly",
        repeat_end_date="2026-11-02",
    )
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 200, r.text

    # 查 10/12-10/25（两周），应该有 2 条
    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-25",
    })
    items = r.json()
    assert len(items) == 2
    assert items[0]["instance_date"] == "2026-10-12"
    assert items[1]["instance_date"] == "2026-10-19"


def test_delete_single_of_repeat(b):
    """删除重复日程的某一次：前后都还在，本次没了"""
    data = _make_schedule(
        title="周课",
        start="2026-10-12T08:00:00",  # 周一
        end="2026-10-12T09:00:00",
        repeat_type="weekly",
        repeat_end_date="2026-10-26",
    )
    r = b.client.post("/api/schedules", json=data)
    sid = r.json()["template_id"]

    # 初始 3 次
    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-26",
    })
    assert len(r.json()) == 3

    # 删中间那次（10/19）
    r = b.client.delete(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-19", "scope": "single"},
    )
    assert r.status_code == 200

    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-26",
    })
    items = r.json()
    assert len(items) == 2
    dates = [it["instance_date"] for it in items]
    assert "2026-10-12" in dates
    assert "2026-10-19" not in dates
    assert "2026-10-26" in dates


def test_update_forward_of_repeat(b):
    """编辑「本次及以后」：前面不变，后面应用新值"""
    data = _make_schedule(
        title="原名称",
        start="2026-10-12T08:00:00",
        end="2026-10-12T09:00:00",
        repeat_type="weekly",
        repeat_end_date="2026-10-26",
    )
    r = b.client.post("/api/schedules", json=data)
    sid = r.json()["template_id"]

    # 从 10/19 起改名为「新名称」
    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-19"},
        json={"title": "新名称", "scope": "forward"},
    )
    assert r.status_code == 200, r.text

    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-26",
    })
    items = r.json()
    assert len(items) == 3

    by_date = {it["instance_date"]: it for it in items}
    assert by_date["2026-10-12"]["title"] == "原名称"
    assert by_date["2026-10-19"]["title"] == "新名称"
    assert by_date["2026-10-26"]["title"] == "新名称"


def test_update_single_of_repeat_keeps_earlier(b):
    """编辑重复日程的「仅本次」：本次之前的那几次必须原样留着

    这个分支要同时表达「前半段」和「后半段」，只靠把模板的起始日挪到下一次的话，
    instance_date 之前的发生日会跟着起点一起消失 —— 改中间一天等于把前面几天
    静默删掉（后半段还在，所以看起来「改成功了」）。
    """
    data = _make_schedule(
        title="原名称",
        start="2026-10-12T08:00:00",  # 周一
        end="2026-10-12T09:00:00",
        repeat_type="weekly",
        repeat_end_date="2026-10-26",
    )
    r = b.client.post("/api/schedules", json=data)
    sid = r.json()["template_id"]

    # 只改中间那次（10/19）
    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-19"},
        json={"title": "临时改名", "scope": "single"},
    )
    assert r.status_code == 200, r.text

    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-26",
    })
    items = r.json()
    dates = sorted(it["instance_date"] for it in items)
    assert dates == ["2026-10-12", "2026-10-19", "2026-10-26"], dates

    by_date = {it["instance_date"]: it for it in items}
    assert by_date["2026-10-12"]["title"] == "原名称"
    assert by_date["2026-10-19"]["title"] == "临时改名"
    # 本次之后仍是重复日程（后半段另立了一条模板）
    assert by_date["2026-10-26"]["title"] == "原名称"
    assert by_date["2026-10-26"]["repeat_type"] == "weekly"


def test_update_single_of_repeat_first_occurrence(b):
    """改「仅本次」且本次就是第一次：后面照常重复，前面本来就没有"""
    data = _make_schedule(
        title="原名称",
        start="2026-10-12T08:00:00",
        end="2026-10-12T09:00:00",
        repeat_type="weekly",
        repeat_end_date="2026-10-26",
    )
    r = b.client.post("/api/schedules", json=data)
    sid = r.json()["template_id"]

    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-12"},
        json={"title": "临时改名", "scope": "single"},
    )
    assert r.status_code == 200, r.text

    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-26",
    })
    items = r.json()
    dates = sorted(it["instance_date"] for it in items)
    assert dates == ["2026-10-12", "2026-10-19", "2026-10-26"], dates

    by_date = {it["instance_date"]: it for it in items}
    assert by_date["2026-10-12"]["title"] == "临时改名"
    assert by_date["2026-10-19"]["title"] == "原名称"
    assert by_date["2026-10-26"]["title"] == "原名称"


def test_update_single_only_occurrence(b):
    """只剩一次发生的重复日程改成「仅本次」：改完就是一条不重复的日程

    这一次既是第一次也是最后一次，原模板没用了 —— 拆完必须被删掉，不能留一条
    展开为空的僵尸行（页面上看不见它，但「第几周」那类按全表算的东西还看得到）。
    顺带守住响应：原行没了之后接口不能去读一个已经删掉的对象。
    """
    data = _make_schedule(
        title="只有一次",
        start="2026-10-12T08:00:00",
        end="2026-10-12T09:00:00",
        repeat_type="weekly",
        repeat_end_date="2026-10-12",
    )
    r = b.client.post("/api/schedules", json=data)
    sid = r.json()["template_id"]

    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-12"},
        json={"title": "改过了", "scope": "single"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["title"] == "改过了"
    assert r.json()["repeat_type"] == "none"

    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-31",
    })
    items = r.json()
    assert [(it["instance_date"], it["title"], it["repeat_type"]) for it in items] == [
        ("2026-10-12", "改过了", "none"),
    ]
    # 原模板没了，只留下这一条
    assert b.rows("SELECT COUNT(*) FROM schedules")[0][0] == 1


def test_update_single_of_repeat_last_occurrence(b):
    """改「仅本次」且本次是最后一次：前面照常重复，后面另立的那条不该出现"""
    data = _make_schedule(
        title="原名称",
        start="2026-10-12T08:00:00",
        end="2026-10-12T09:00:00",
        repeat_type="weekly",
        repeat_end_date="2026-10-26",
    )
    r = b.client.post("/api/schedules", json=data)
    sid = r.json()["template_id"]

    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-26"},
        json={"title": "临时改名", "scope": "single"},
    )
    assert r.status_code == 200, r.text

    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-12-31",
    })
    items = r.json()
    dates = sorted(it["instance_date"] for it in items)
    assert dates == ["2026-10-12", "2026-10-19", "2026-10-26"], dates

    by_date = {it["instance_date"]: it for it in items}
    assert by_date["2026-10-12"]["title"] == "原名称"
    assert by_date["2026-10-19"]["title"] == "原名称"
    assert by_date["2026-10-26"]["title"] == "临时改名"


# ---------- 权限隔离 ----------

def test_student_cannot_see_others(b):
    """学生只能看到自己的日程，看不到别人的"""
    student2 = as_student(b, 1)  # 乙

    # 管理员建一个「仅自己」的日程
    data = _make_schedule(title="管理员私人日程")
    b.client.post("/api/schedules", json=data)

    # 乙看不到
    r = student2.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-18",
    })
    assert len(r.json()) == 0

    # 管理员自己能看到
    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-18",
    })
    assert len(r.json()) == 1


def test_student_cannot_delete_others(b):
    """学生不能删除别人的日程（404）"""
    student2 = as_student(b, 1)

    # 管理员建一个自己的日程
    data = _make_schedule(title="管理员的")
    r = b.client.post("/api/schedules", json=data)
    sid = r.json()["id"]
    inst = r.json()["instance_date"]

    # 乙尝试删 → 404
    r = student2.delete(
        f"/api/schedules/item/{sid}",
        params={"instance_date": inst, "scope": "all"},
    )
    assert r.status_code == 404


def test_committee_batch_create(b):
    """班委批量创建：目标学生都能看到"""
    # 先拿到学生 ID
    r = b.client.get("/api/students", params={"page": 1, "page_size": 10})
    students = {s["student_no"]: s["id"] for s in r.json()["items"]}

    sid_yi = students[ROSTER[1][0]]  # 乙
    sid_bing = students[ROSTER[2][0]]  # 丙

    # 管理员给乙和丙批量建日程
    data = _make_schedule(
        title="班会",
        start="2026-10-13T16:00:00",
        end="2026-10-13T17:00:00",
        target_student_ids=[sid_yi, sid_bing],
    )
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 200, r.text

    # 乙能看到
    student2 = as_student(b, 1)
    r = student2.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-18",
    })
    items = r.json()
    assert len(items) >= 1
    assert any(it["title"] == "班会" for it in items)
    # 来源是 batch
    assert any(it["source_type"] == "batch" for it in items)

    # 丙也能看到
    student3 = as_student(b, 2)
    r = student3.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-18",
    })
    assert any(it["title"] == "班会" for it in r.json())


def test_regular_student_cannot_batch(b):
    """普通学生传了 target_student_ids 也没用，只建自己的"""
    student2 = as_student(b, 1)

    r = b.client.get("/api/students", params={"page": 1, "page_size": 10})
    all_ids = [s["id"] for s in r.json()["items"]]

    data = _make_schedule(
        title="学生想批量发",
        target_all=True,
        target_student_ids=all_ids,
    )
    r = student2.post("/api/schedules", json=data)
    # 应该成功（建自己的），但不会批量
    assert r.status_code == 200

    # 丙那里不会有
    student3 = as_student(b, 2)
    r = student3.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-18",
    })
    assert not any(it["title"] == "学生想批量发" for it in r.json())


# ---------- 导入解析 ----------

def _make_schedule_xlsx(
    first_monday: date = date(2026, 10, 12),
    weeks: int = 2,
    first_week_no: int = 1,
    date_in_header: bool = True,
) -> bytes:
    """造一份简单的周课表 xlsx（默认两周）

    first_monday 是**第 first_week_no 周**那个 sheet 的周一，换掉它就能造出另一份
    课表（另一学期）—— 用来验「新的一份 Excel 重新从第 1 周编号」。
    weeks / first_week_no 用来造补课表那种「只含第 8–10 周」的文件；
    date_in_header=False 造一份表头没写日期的（只剩「周一」），用来验「猜不出锚点，
    得等用户选」。
    """
    wb = openpyxl.Workbook()

    for week_idx in range(weeks):
        sheet_name = f"第{week_idx + first_week_no}周"
        if week_idx == 0:
            ws = wb.active
            ws.title = sheet_name
        else:
            ws = wb.create_sheet(sheet_name)

        monday = first_monday + __import__('datetime').timedelta(weeks=week_idx)
        weekdays = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]

        # 表头
        ws.cell(row=1, column=1, value="时间")
        for i in range(7):
            d = monday + __import__('datetime').timedelta(days=i)
            value = f"{d.month}/{d.day} {weekdays[i]}" if date_in_header else weekdays[i]
            ws.cell(row=1, column=i + 2, value=value)

        # 时间列（07:00 - 21:50，10 分钟一行 = 90 行）
        for row in range(90):
            hour = 7 + row // 6
            minute = (row % 6) * 10
            ws.cell(row=row + 2, column=1,
                    value=f"{hour:02d}:{minute:02d}")

        # 周一 08:00-09:50（第 1 行时间是 07:00 在 row2，08:00 在 row8）
        # 8:00 = (8-7)*6 + 0/10 = 6 → row = 6 + 2 = 8
        # 9:50 = (9-7)*6 + 50/10 = 12+5=17 → row = 17 + 2 = 19
        start_row = 8   # 08:00
        end_row = 19    # 09:50
        ws.merge_cells(start_row=start_row, start_column=2,
                       end_row=end_row, end_column=2)
        ws.cell(row=start_row, column=2,
                value="高等数学\n课程\n教学楼C101")

        # 周三 14:00-15:50（单周才有，模拟单双周 —— 这里第 1 周有、第 2 周没有）
        if week_idx == 0:
            # 14:00 = (14-7)*6 = 42 → row = 42 + 2 = 44
            # 15:50 = (15-7)*6 + 5 = 48+5=53 → row = 53 + 2 = 55
            ws.merge_cells(start_row=44, start_column=4, end_row=55, end_column=4)
            ws.cell(row=44, column=4, value="大学英语\n课程\n语音室")

        # 底部全天事件区：事件写在它覆盖的那一天的列里（B–H = 周一–周日）
        ws.cell(row=92, column=1, value="事件1")
        ws.cell(row=92, column=2, value=f"正式活动 {monday.month}/{monday.day}")
        if week_idx == 0:
            # 跨周六、周日两天：同一段文字在 G、H 两列各写一遍
            ws.cell(row=92, column=7, value="运动会")
            ws.cell(row=92, column=8, value="运动会")
            # 同一天（周日）还有第二条事件
            ws.cell(row=93, column=1, value="事件2")
            ws.cell(row=93, column=8, value="收假返校")
            # 跨周三、四、五三天，写成**横向合并**（D–F 一格）—— 合并之后只有
            # 左上角 D94 有值，逐格取会让这件事件只落在周三一天
            ws.cell(row=94, column=1, value="事件3")
            ws.merge_cells(start_row=94, start_column=4, end_row=94, end_column=6)
            ws.cell(row=94, column=4, value="国庆放假")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_preview_faithful(b):
    """导入预览：忠实展平，不推断重复——每周每块各一条独立日程"""
    xlsx_bytes = _make_schedule_xlsx()

    r = b.client.post(
        "/api/schedules/import/preview",
        files={"file": ("课表.xlsx", xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["week_count"] == 2

    # 高数两周各一块 → 2 条独立日程，日期分别落在两个周一
    math_items = [it for it in data["items"] if "高等数学" in it["title"]]
    assert len(math_items) == 2
    dates = sorted(it["start_date"] for it in math_items)
    assert dates == ["2026-10-12", "2026-10-19"]
    # 全部不重复
    assert all(it["repeat_type"] == "none" for it in math_items)
    # 类型、地点解析正确；类型不再污染备注
    assert math_items[0]["event_type"] == "课程"
    assert "教学楼C" in math_items[0]["location"]
    assert math_items[0]["description"] == ""

    # 大学英语只在第 1 周出现 → 恰好 1 条（周三 10/14）
    eng_items = [it for it in data["items"] if "大学英语" in it["title"]]
    assert len(eng_items) == 1
    assert eng_items[0]["start_date"] == "2026-10-14"


def test_all_day_section_reads_every_column(b):
    """全天事件区逐格取：跨天的事件每天各出一条，同一天可以有好几条

    三个曾经踩过的坑：
    1. 只读 B 列 —— 事件写在它覆盖的那一天的列里，真实课表的「国庆放假」
       横跨 7 列，只读 B 列会让后面 6 天整片消失；
    2. 把全天事件塞进以「星期+开始时间」为键的字典 —— 凌晨 00:00 撞车，
       同一天的第二条会把第一条覆盖掉；
    3. 跨天写成横向合并时只认左上角 —— 合并之后只有最左边那格存着文字，
       逐格读会让 D94:F94 的「国庆放假」只落在周三，后面两天整片消失。

    跨天的两种写法（逐列重复写、横向合并）都摊成**一天一条**，页面上再由
    allDayBars 把相邻的同名条连回一条 —— 模型里没有「跨天日程」这个概念。
    """
    xlsx_bytes = _make_schedule_xlsx()
    r = b.client.post(
        "/api/schedules/import/preview",
        files={"file": ("课表.xlsx", xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200, r.text
    allday = [it for it in r.json()["items"] if it["all_day"]]
    assert [(it["start_date"], it["title"]) for it in allday] == [
        ("2026-10-12", "正式活动 10/12"),   # 第 1 周周一
        ("2026-10-14", "国庆放假"),         # 横向合并 D94:F94 → 周三
        ("2026-10-15", "国庆放假"),         #                   周四
        ("2026-10-16", "国庆放假"),         #                   周五
        ("2026-10-17", "运动会"),           # 周六（同一段文字写了两列）
        ("2026-10-18", "运动会"),           # 周日
        ("2026-10-18", "收假返校"),         # 周日同一天的第二条
        ("2026-10-19", "正式活动 10/19"),   # 第 2 周周一
    ]
    # 全天事件不参与重复推断，类型/地点为空
    assert all(it["repeat_type"] == "none" for it in allday)
    assert all(it["event_type"] == "" and it["location"] == "" for it in allday)


def _typed_cell_xlsx() -> bytes:
    """表头和时间列都是 Excel 的**真日期/真时间**格式（不是文本）

    这种表非常常见：别人把单元格格式设成日期，值就不是字符串了。早先的解析只认
    字符串，整张表会以「无法识别日期表头」「没有识别到时间列」被静默跳过 ——
    预览里一条日程都没有，用户只会以为文件选错了。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "第1周"
    monday = date(2026, 10, 12)
    ws.cell(row=1, column=2, value=datetime(2026, 10, 12))          # 真日期
    ws.cell(row=1, column=3, value="10/13 周二")                    # 同一个表里混字符串
    for row in range(90):
        hour = 7 + row // 6
        minute = (row % 6) * 10
        ws.cell(row=row + 2, column=1, value=time(hour, minute))    # 真时间
    ws.merge_cells(start_row=8, start_column=2, end_row=19, end_column=2)
    ws.cell(row=8, column=2, value="高等数学\n课程\n教学楼C101")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_accepts_real_date_and_time_cells(b):
    """表头 / 时间列是 Excel 的真日期、真时间格式时照样能导入"""
    r = b.client.post(
        "/api/schedules/import/preview",
        files={"file": ("课表.xlsx", _typed_cell_xlsx(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["week_count"] == 1
    assert data["warnings"] == []
    items = [it for it in data["items"] if "高等数学" in it["title"]]
    assert len(items) == 1
    assert items[0]["start_date"] == "2026-10-12"
    # 结束时间 = 最后一行（第 19 行 = 09:50）再往后一格
    assert (items[0]["start_time"], items[0]["end_time"]) == ("08:00", "10:00")


def test_import_header_not_monday(b):
    """表头那天不是周一时按它所在周的周一算，并在警告里写清用了哪一天"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "第1周"
    ws.cell(row=1, column=2, value="10/14 周三")   # 2026-10-14 是周三
    for row in range(90):
        ws.cell(row=row + 2, column=1, value=f"{7 + row // 6:02d}:{(row % 6) * 10:02d}")
    ws.merge_cells(start_row=8, start_column=2, end_row=19, end_column=2)
    ws.cell(row=8, column=2, value="高等数学")
    buf = BytesIO()
    wb.save(buf)

    r = b.client.post(
        "/api/schedules/import/preview",
        files={"file": ("课表.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200, r.text
    data = r.json()
    assert any("2026-10-12" in w for w in data["warnings"]), data["warnings"]
    # 块在 B 列（周一），所以落在那周的周一 10/12；表头那格 10/14 只用来定周
    assert data["items"][0]["start_date"] == "2026-10-12"


def test_import_confirm_writes_to_targets(b):
    """导入确认：目标学生都拿到了日程"""
    xlsx_bytes = _make_schedule_xlsx()

    r = b.client.post(
        "/api/schedules/import/preview",
        files={"file": ("课表.xlsx", xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    preview_items = r.json()["items"]
    token = r.json()["token"]

    # 先拿学生 ID
    r2 = b.client.get("/api/students", params={"page": 1, "page_size": 10})
    all_ids = [s["id"] for s in r2.json()["items"]]

    # 确认导入给全部学生
    r = b.client.post("/api/schedules/import", json={
        "token": token,
        "batch_name": "秋季课表",
        "target_all": True,
        "target_student_ids": [],
        "selected_indices": [],
    })
    assert r.status_code == 200, r.text
    result = r.json()
    # 条数从预览里数，不写死：夹具展平出什么由 test_all_day_section_reads_every_column
    # 逐条钉着，这里只关心「每个目标学生都拿到了同样多的一份」
    expected = len(preview_items)
    assert result["target_count"] == len(all_ids)
    assert result["schedule_per_student"] == expected
    assert result["total_created"] == expected * len(all_ids)

    # 学生乙能看到导入的课
    student2 = as_student(b, 1)
    r = student2.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-18",
    })
    titles = [it["title"] for it in r.json()]
    assert "高等数学" in titles


def test_regular_student_cannot_import(b):
    """普通学生不能调导入接口"""
    student2 = as_student(b, 1)
    xlsx_bytes = _make_schedule_xlsx()

    r = student2.post(
        "/api/schedules/import/preview",
        files={"file": ("课表.xlsx", xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 403  # 权限不足


# ---------- iCal 订阅 ----------

def test_ical_token_and_feed(b):
    """获取订阅 token，用 token 能拿到 ics 文件"""
    # 获取 token
    r = b.client.get("/api/schedules/ical/token")
    assert r.status_code == 200
    data = r.json()
    assert "token" in data
    assert "url" in data
    assert len(data["token"]) > 10

    # 先建一条日程
    b.client.post("/api/schedules", json=_make_schedule(title="iCal测试课"))

    # 用 token 拉 iCal（直接走 HTTP，不通过已登录客户端的 cookies）
    base = str(b.client.base_url).rstrip("/")
    r = httpx.get(f"{base}/api/schedules/ical/{data['token']}.ics")
    assert r.status_code == 200
    assert "VCALENDAR" in r.text
    assert "VEVENT" in r.text
    assert "iCal测试课" in r.text


def _ical_text(b) -> str:
    """拿当前学生的订阅文件正文"""
    token = b.client.get("/api/schedules/ical/token").json()["token"]
    base = str(b.client.base_url).rstrip("/")
    r = httpx.get(f"{base}/api/schedules/ical/{token}.ics")
    assert r.status_code == 200, r.text
    return r.text


def _ical_events(text: str) -> list[dict]:
    """把正文拆成 VEVENT 列表，每个只留测试要看的那几行

    先把续行摊平 —— _fold 会给超 75 字节的行加续行，标题/备注长一点就会被拆开。
    键取「行首到分号为止」：DTSTART;VALUE=DATE:20260131 要能读成 DTSTART
    """
    lines: list[str] = []
    for line in text.split("\r\n"):
        if line.startswith(" ") and lines:
            lines[-1] += line[1:]
        else:
            lines.append(line)

    events: list[dict] = []
    cur: Optional[dict] = None
    for line in lines:
        if line == "BEGIN:VEVENT":
            cur = {}
        elif line == "END:VEVENT":
            if cur is not None:
                events.append(cur)
            cur = None
        elif cur is not None and ":" in line:
            key, value = line.split(":", 1)
            cur[key.split(";")[0]] = value
    return events


def _ical_dates(text: str) -> dict[str, list[str]]:
    """正文里「标题 → 发生日（YYYY-MM-DD）」的映射

    只认每个 VEVENT 的 DTSTART 那天（测试里的日程都不跨天）。带 RRULE 的事件在
    这里只算 DTSTART 一条 —— 要验「和网页完全一样」就不能有 RRULE，见
    test_ical_dates_match_page 里的断言。
    """
    out: dict[str, list[str]] = {}
    for ev in _ical_events(text):
        raw = ev.get("DTSTART", "")
        out.setdefault(ev.get("SUMMARY", ""), []).append(
            f"{raw[0:4]}-{raw[4:6]}-{raw[6:8]}"
        )
    return out


def _page_dates(b, start: str, end: str) -> dict[str, list[str]]:
    """页面上「标题 → 发生日」的映射（走接口，页面拿的就是它）"""
    r = b.client.get("/api/schedules", params={"start_date": start, "end_date": end})
    assert r.status_code == 200, r.text
    out: dict[str, list[str]] = {}
    for it in r.json():
        out.setdefault(it["title"], []).append(it["instance_date"])
    return out


def test_ical_event_shape(b):
    """iCal 正文的几处格式：备注换行不被二次转义、行不超 75 字节、DTSTAMP 是 UTC"""
    b.client.post("/api/schedules", json=_make_schedule(
        title="远足", location="操场", description="带水\n带伞",
        start="2026-10-12T08:00:00", end="2026-10-12T09:00:00",
    ))
    text = _ical_text(b)

    # 换行写成「反斜杠 + n」两个字符。以前拼的是 '\\n'（本身就是两个字符），
    # 再转义一次变成 「\\n」，客户端看到的是字面的 \n，不是换行
    assert "DESCRIPTION:带水\\n带伞" in text
    assert "\\\\n" not in text
    # DTSTAMP 按 RFC5545 必须是 UTC（整行就是 CRLF 包的，别用 ^$ 匹配，
    # 行尾那个 \r 会让 $ 对不上）
    assert re.search(r"\r\nDTSTAMP:\d{8}T\d{6}Z\r\n", text)
    # 每行（含续行）不超过 75 字节 —— 中文一个字 3 字节，长备注很容易超
    for line in text.split("\r\n"):
        assert len(line.encode("utf-8")) <= 75, line


def test_ical_long_line_folded(b):
    """超长的一行要折行，续行以空格开头，且不劈开多字节字符"""
    b.client.post("/api/schedules", json=_make_schedule(
        title="很长的标题" * 20, description="备注" * 60,
    ))
    text = _ical_text(b)
    lines = text.split("\r\n")
    assert any(line.startswith(" ") for line in lines), "应该出现了续行"
    for line in lines:
        assert len(line.encode("utf-8")) <= 75
        assert "�" not in line  # 按字符折，不会出现半截汉字


def test_ical_repeat_dtstart_matches_page(b):
    """单双周：订阅里的日子就是页面上那些日子（逐次展开出来的）

    奇偶按学期周次算（周次号来自导入时确认的锚点，写在这条日程的 week_start 上）。
    用户把开始日填在了相反的一周时，页面从下一次才出现 —— 订阅也得一样：
    不能拿模板开始日当种子往外发。
    """
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    # 先导入一份只含第 8–10 周的补课表（锚点 = 第 8 周周一 10/26）：
    # 学期第 1 周周一被推为 2026-09-07，区间覆盖 10/26–11/15
    r = b.client.post(
        "/api/schedules/import/preview",
        files={"file": ("补课表.xlsx", _make_schedule_xlsx(date(2026, 10, 26), weeks=3, first_week_no=8), mime)},
    )
    assert r.status_code == 200, r.text
    r = b.client.post("/api/schedules/import", json={
        "token": r.json()["token"], "batch_name": "补课表",
        "target_all": True, "target_student_ids": [], "selected_indices": [],
    })
    assert r.status_code == 200, r.text

    r = b.client.post("/api/schedules", json=_make_schedule(
        title="单双周课", start="2026-10-27T08:00:00", end="2026-10-27T09:00:00"))
    sid = r.json()["id"]
    # 10/27 落在第 8 周（双周），而这条是单周课
    b.run(
        "UPDATE schedules SET repeat_type='odd_week', repeat_end_date='2026-12-31', "
        "week_start='2026-09-07' WHERE id=?",
        sid,
    )

    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-26", "end_date": "2026-11-08",
    })
    dates = [it["instance_date"] for it in r.json() if it["title"] == "单双周课"]
    assert dates == ["2026-11-03"]   # 第 9 周（单周）的周二

    text = _ical_text(b)
    # 第 10 周（11/09–11/15）是双周、第 11 周起没有周次号，都只有这一次
    assert _ical_dates(text)["单双周课"] == ["2026-11-03"], text
    assert "DTSTART:20261103T080000" in text
    assert "DTEND:20261103T090000" in text
    assert "RRULE" not in text
    # 只发生一次 → UID 沿用不带日期的那种（不给订阅端制造无谓的「删了再加」）
    assert f"UID:schedule-{sid}@class-oa" in text


def test_ical_monthly_last_day(b):
    """每月重复：页面那边没有那一天的月份取月末，订阅得逐次展开成同样的日子

    BYMONTHDAY 只能写死一个数，表达不了「2 月落到 28」这种逐月变化。老代码只给
    31 号用了 -1，填 30 号时 2 月在订阅里会整个消失（网页上明明是 2/28）。
    """
    def add(title, start, end):
        r = b.client.post("/api/schedules", json=_make_schedule(
            title=title, start=start, end=end,
            repeat_type="monthly", repeat_end_date="2026-04-30"))
        assert r.status_code == 200, r.text

    add("月末总结", "2026-01-31T08:00:00", "2026-01-31T09:00:00")   # 31 号
    add("月中年终", "2026-01-30T08:00:00", "2026-01-30T09:00:00")   # 30 号：2 月该落到 28

    text = _ical_text(b)
    ical = _ical_dates(text)
    page = _page_dates(b, "2026-01-30", "2026-04-30")
    assert sorted(ical["月末总结"]) == sorted(page["月末总结"]) == [
        "2026-01-31", "2026-02-28", "2026-03-31", "2026-04-30"]
    assert sorted(ical["月中年终"]) == sorted(page["月中年终"]) == [
        "2026-01-30", "2026-02-28", "2026-03-30", "2026-04-30"]
    # 展开出来的事件不该再带重复规则（否则客户端会自己长出一堆网页上没有的日子）
    assert "RRULE" not in text
    assert "BYMONTHDAY" not in text

    # 全天事件也一样：DTSTART 是 DATE，每个发生日一条，DTEND 是第二天（半开区间）
    b.client.post("/api/schedules", json=_make_schedule(
        title="月末假", start="2026-01-31T00:00:00", end="2026-01-31T23:59:00",
        all_day=True, repeat_type="monthly", repeat_end_date="2026-04-30"))
    text = _ical_text(b)
    assert sorted(_ical_dates(text)["月末假"]) == [
        "2026-01-31", "2026-02-28", "2026-03-31", "2026-04-30"]
    assert "DTSTART;VALUE=DATE:20260131" in text
    assert "DTEND;VALUE=DATE:20260201" in text


def test_ical_dates_match_page(b):
    """订阅里的每一天都和网页一致（单次 / 单周 / 双周 / 每月 31 / 每月 30）

    日期集合直接比对接口返回的 instance_date。这五条都该是逐次展开出来的独立
    事件，正文里一条 RRULE 都不该有 —— 有 RRULE 的话客户端会自己多长出一堆
    网页上没有的日子，只比 DTSTART 集合是看不出来的。
    """
    def add(title, start, end, **kw):
        r = b.client.post("/api/schedules", json=_make_schedule(
            title=title, start=start, end=end, **kw))
        assert r.status_code == 200, r.text

    add("单次", "2026-01-05T08:00:00", "2026-01-05T09:00:00")
    # 没导过课表：单双周以各自的开始周为第 1 周（见 _periods_for）
    add("单周", "2026-01-05T09:00:00", "2026-01-05T10:00:00",
        repeat_type="odd_week", repeat_end_date="2026-02-15")
    add("双周", "2026-01-05T10:00:00", "2026-01-05T11:00:00",
        repeat_type="even_week", repeat_end_date="2026-02-15")
    add("月卅一", "2026-01-31T11:00:00", "2026-01-31T12:00:00",
        repeat_type="monthly", repeat_end_date="2026-03-31")
    add("月卅", "2026-01-30T12:00:00", "2026-01-30T13:00:00",
        repeat_type="monthly", repeat_end_date="2026-03-31")

    text = _ical_text(b)
    assert "RRULE" not in text, "这五条都不该走 RRULE"
    ical = _ical_dates(text)
    page = _page_dates(b, "2026-01-01", "2026-04-01")

    assert set(ical) == set(page) == {"单次", "单周", "双周", "月卅一", "月卅"}
    for title, days in page.items():
        assert sorted(ical[title]) == sorted(days), title

    # 同一模板展开成多条时 UID 加了日期，正文里不该有重复的 UID
    uids = [ev["UID"] for ev in _ical_events(text)]
    assert len(set(uids)) == len(uids)


def test_ical_weekly_still_uses_rrule(b):
    """每周重复仍是一个 VEVENT + RRULE（它和网页「每 7 天一次」逐日等价，不必展开）"""
    r = b.client.post("/api/schedules", json=_make_schedule(
        title="每周例会", start="2026-10-12T19:00:00", end="2026-10-12T20:00:00",
        repeat_type="weekly", repeat_end_date="2026-11-30"))
    assert r.status_code == 200, r.text

    text = _ical_text(b)
    assert text.count("BEGIN:VEVENT") == 1
    assert "DTSTART:20261012T190000" in text
    # UNTIL 与 DTSTART 同为浮动本地时间（给浮动时间配 Z 的 UNTIL 不合 RFC5545）
    assert "RRULE:FREQ=WEEKLY;INTERVAL=1;BYDAY=MO;UNTIL=20261130T190000" in text, text

    # 全天事件：DTSTART 和 UNTIL 都是 DATE
    b.client.post("/api/schedules", json=_make_schedule(
        title="每周值日", start="2026-10-12T00:00:00", end="2026-10-12T23:59:00",
        all_day=True, repeat_type="weekly", repeat_end_date="2026-11-30"))
    text = _ical_text(b)
    assert "DTSTART;VALUE=DATE:20261012" in text
    assert "UNTIL=20261130" in text, text


def test_ical_weekly_cross_day_until_uses_start_time(b):
    """跨天的每周日程：UNTIL 用开始时刻，否则最后一天那次会被客户端丢掉

    22:00–次日 1:00 这种日程，拿结束时刻（01:00）当 UNTIL 的话，
    repeat_end_date 那天 22:00 的那一次落在 UNTIL 之后 —— 网页上有、手机上没有。

    接口已经不收这种日程了（见 test_cross_day_schedule_rejected），所以这里
    直接把库里的 end_time 改成次日 —— 造一条**历史数据**：老库里有、别的部署
    可能还有（也挡不住有人直接改库），iCal 这边不能因此把最后一天算丢。
    """
    r = b.client.post("/api/schedules", json=_make_schedule(
        title="夜间值守", start="2026-10-12T22:00:00", end="2026-10-12T23:00:00",
        repeat_type="weekly", repeat_end_date="2026-11-30"))
    assert r.status_code == 200, r.text
    b.run("UPDATE schedules SET end_time='2026-10-13 01:00:00' WHERE id=?", r.json()["id"])

    text = _ical_text(b)
    assert "DTSTART:20261012T220000" in text
    assert "RRULE:FREQ=WEEKLY;INTERVAL=1;BYDAY=MO;UNTIL=20261130T220000" in text, text


def test_ical_degenerate_weekly_writes_nothing(b):
    """「删除本次及以后」删掉第一条之后：网页上没了，订阅里也得一条都不写

    删除会把 repeat_end_date 设成开始日的前一天，于是这条日程一次都不发生。
    RRULE 那条路表达不了「一次都没有」—— DTSTART 的那天按 RFC5545 永远算一次
    发生，客户端照样会显示一条已经删掉的日程。
    """
    r = b.client.post("/api/schedules", json=_make_schedule(
        title="待删课", start="2026-10-12T08:00:00", end="2026-10-12T09:00:00",
        repeat_type="weekly", repeat_end_date="2026-11-30"))
    sid = r.json()["id"]
    r = b.client.delete(f"/api/schedules/item/{sid}", params={
        "instance_date": "2026-10-12", "scope": "forward",
    })
    assert r.status_code == 200, r.text

    assert _page_dates(b, "2026-10-01", "2026-12-31").get("待删课") is None
    text = _ical_text(b)
    assert "待删课" not in text
    assert f"schedule-{sid}" not in text


def test_ical_odd_week_outside_coverage_is_absent(b):
    """探针复现：覆盖范围外的单周日程，网页上没有，订阅里也必须没有

    导入两周课表后手动建一条 11/16 起的单周日程 —— 它以**自己的开始周**为第 1 周
    （见 _periods_for），而覆盖范围收在课表区间（10/12–10/25）里，11/16 那周没有
    周次号，网页上一条都不出现。老代码 iCal 那边用「兜底成开始日」硬发了一条
    DTSTART，学生手机上于是有网页上没有的日程。
    """
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    r = b.client.post(
        "/api/schedules/import/preview",
        files={"file": ("课表.xlsx", _make_schedule_xlsx(date(2026, 10, 12)), mime)},
    )
    assert r.status_code == 200, r.text
    r = b.client.post("/api/schedules/import", json={
        "token": r.json()["token"], "batch_name": "课表",
        "target_all": True, "target_student_ids": [], "selected_indices": [],
    })
    assert r.status_code == 200, r.text

    r = b.client.post("/api/schedules", json=_make_schedule(
        title="下学期单周课", start="2026-11-16T19:00:00", end="2026-11-16T20:00:00",
        repeat_type="odd_week", repeat_end_date="2026-12-31"))
    assert r.status_code == 200, r.text
    sid = r.json()["id"]

    assert _page_dates(b, "2026-11-01", "2026-12-31").get("下学期单周课") is None

    text = _ical_text(b)
    assert "下学期单周课" not in text
    assert "DTSTART:20261116T190000" not in text
    # 连 UID 都不该出现：只查某个时刻的话，半对不半对的订阅还能蒙混过去
    assert f"schedule-{sid}@class-oa" not in text
    assert f"schedule-{sid}-" not in text


def test_ical_bad_token_404(b):
    """无效 token 返回 404（不泄露存在性）"""
    base = str(b.client.base_url).rstrip("/")
    r = httpx.get(f"{base}/api/schedules/ical/nonexistent_token_123.ics")
    assert r.status_code == 404


def test_ical_reset_token(b):
    """重置 token 后旧 token 失效"""
    r = b.client.get("/api/schedules/ical/token")
    old_token = r.json()["token"]

    r = b.client.post("/api/schedules/ical/token/reset")
    assert r.status_code == 200
    new_token = r.json()["token"]
    assert new_token != old_token

    base = str(b.client.base_url).rstrip("/")
    # 新 token 能用
    r = httpx.get(f"{base}/api/schedules/ical/{new_token}.ics")
    assert r.status_code == 200
    # 旧 token 失效
    r = httpx.get(f"{base}/api/schedules/ical/{old_token}.ics")
    assert r.status_code == 404

    # 再重置一次：唯一约束是 (student_id, is_active)，上一次那条 is_active=0 的行
    # 留着就会让「把当前这行置为无效」撞约束 → 500。500 的话订阅链接就作废不掉了
    r = b.client.post("/api/schedules/ical/token/reset")
    assert r.status_code == 200, r.text
    third = r.json()["token"]
    assert third not in (old_token, new_token)
    assert httpx.get(f"{base}/api/schedules/ical/{third}.ics").status_code == 200
    # 上上轮那条也不该复活
    assert httpx.get(f"{base}/api/schedules/ical/{new_token}.ics").status_code == 404


def test_ical_alarm_absolute_for_all_day(b):
    """全天日程的闹钟写**绝对时刻**，和 App 推送是同一瞬间

    全天事件存的是当天 00:00–23:59（见 _all_day_times），提前 10 分钟就是
    前一天 23:50（北京时间）→ UTC 15:50Z。写相对的 `-PT10M` 的话，各客户端
    对全天事件里相对闹钟的解释不一致（不少干脆忽略），改成绝对时刻两边才真的
    对得上。普通日程仍写相对的。
    """
    b.client.post("/api/schedules", json=_make_schedule(
        title="校运会", start="2026-11-16T00:00:00", end="2026-11-16T23:59:00",
        all_day=True, notify_enabled=True, notify_minutes_before=10,
    ))
    b.client.post("/api/schedules", json=_make_schedule(
        title="班委会", start="2026-10-12T08:00:00", end="2026-10-12T09:00:00",
        notify_enabled=True, notify_minutes_before=30,
    ))
    text = _ical_text(b)

    events = {ev["SUMMARY"]: ev for ev in _ical_events(text)}
    assert events["校运会"]["DTSTART"] == "20261116"
    assert events["校运会"]["TRIGGER"] == "20261115T155000Z"
    # 普通日程不受影响，还是相对的
    assert events["班委会"]["TRIGGER"] == "-PT30M"


# ---------- 日程提醒的推送正文 ----------

def test_all_day_notification_body_says_all_day(b, monkeypatch):
    """全天日程的推送正文写「全天」，不是哨兵时刻 00:00 - 23:59

    起止在库里就是 00:00 / 23:59，照 strftime 拼出来学生收到会以为日程在半夜。
    这里只验正文怎么拼，所以把收件人查询和发送都短路掉；tpl 从真库里那条读出来，
    免得假对象把「库里存的是什么」这个前提也一起假掉。

    顺带钉住这个函数**能跑到底**：它原来在 `from .email_service import
    email_enabled` 那行 ImportError（没这个名字），而 tick 的 except 会把异常
    吞掉只打一行日志 —— 日程提醒从上线起一条都没发出去过。用例只要调用它就
    会碰到那个 import，改坏了立刻红。
    """
    from datetime import datetime as dt
    from types import SimpleNamespace

    from app.services import schedule_notify_service as sns

    b.client.post("/api/schedules", json=_make_schedule(
        title="校运会", start="2026-11-16T00:00:00", end="2026-11-16T23:59:00",
        all_day=True, notify_enabled=True, notify_minutes_before=10,
        location="操场",
    ))
    title_, all_day, start, end, location, description, source_type = b.rows(
        "SELECT title, all_day, start_time, end_time, location, description, source_type"
        " FROM schedules WHERE title = ?", "校运会",
    )[0]
    tpl = SimpleNamespace(
        title=title_, all_day=bool(all_day),
        start_time=dt.strptime(str(start), "%Y-%m-%d %H:%M:%S.%f"),
        end_time=dt.strptime(str(end), "%Y-%m-%d %H:%M:%S.%f"),
        location=location, description=description, source_type=source_type,
    )

    # 收件人查询全部短路，发送也换成只记正文
    for mod, fn in [
        ("app.services.inbox_service", "recipient_nos"),
        ("app.services.bark_service", "bark_keys"),
        ("app.services.webpush_service", "web_subscriptions"),
        ("app.services.webhook_service", "webhook_targets"),
        ("app.services.email_service", "email_recipients"),
    ]:
        monkeypatch.setattr(f"{mod}.{fn}", lambda db, **kw: [])

    sent = {}
    monkeypatch.setattr(
        sns.ns, "deliver_now",
        lambda db, targets, title, body, **kw: sent.update(title=title, body=body),
    )

    sns._send_notification(None, tpl, SimpleNamespace(student_no="admin"),
                           dt(2026, 11, 15, 23, 50))

    assert sent["title"] == "日程提醒：校运会"
    assert sent["body"].splitlines()[0] == "时间：全天"
    assert "00:00" not in sent["body"] and "23:59" not in sent["body"]
    assert "地点：操场" in sent["body"]


# ---------- 批次管理 ----------

def test_batch_list_and_delete(b):
    """班委能看到自己创建的批次，并能撤销"""
    # 先批量建一个
    r = b.client.get("/api/students", params={"page": 1, "page_size": 10})
    all_ids = [s["id"] for s in r.json()["items"]]

    data = _make_schedule(
        title="临时活动",
        target_student_ids=all_ids,
    )
    # 用批量创建（走 batch 路径）
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 200

    # 看批次列表。上面那次批量创建其实也产生了批次（班委分支走的就是 create_batch），
    # 这里另外导入一个是为了按名字认得出「撤销的是哪一个」
    xlsx_bytes = _make_schedule_xlsx()
    r = b.client.post(
        "/api/schedules/import/preview",
        files={"file": ("x.xlsx", xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    token = r.json()["token"]
    r = b.client.post("/api/schedules/import", json={
        "token": token, "batch_name": "测试批次",
        "target_all": True, "target_student_ids": [],
        "selected_indices": [],
    })
    batch_id = r.json()["batch_id"]

    # 批次列表有它
    r = b.client.get("/api/schedules/batches")
    assert r.status_code == 200
    batches = r.json()
    assert any(b["id"] == batch_id for b in batches)

    # 删批次
    r = b.client.delete(f"/api/schedules/batches/{batch_id}")
    assert r.status_code == 200
    assert r.json()["deleted"] > 0

    # 学生那边没了
    student2 = as_student(b, 1)
    r = student2.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-18",
    })
    assert not any(it["title"] == "高等数学" for it in r.json())


# ---------- 跨学生下发 ----------
#
# 班委建的多条日程，实现上是给每个学生各写一条（create_batch），同一条原始日程的
# 副本共用 source_series_id。默认改/删只动自己那条；带上 share=true 才下发到其他人
# 的表里 —— 这个能力只属于「当初把这条日程分配出去的人」。

def _list(client, start="2026-10-12", end="2026-11-08"):
    """某个客户端（某个学生）看到的一段日程"""
    r = client.get("/api/schedules", params={"start_date": start, "end_date": end})
    assert r.status_code == 200, r.text
    return r.json()


def _items_of(items, title):
    """按标题筛出实例，按日期排好——断言展开结果时用它比对着看"""
    return sorted(
        [(it["instance_date"], it["title"], it["share_count"]) for it in items
         if it["title"].startswith(title)],
    )


def _student_ids(b):
    r = b.client.get("/api/students", params={"page": 1, "page_size": 10})
    return {s["student_no"]: s["id"] for s in r.json()["items"]}


def test_share_count_in_list(b):
    """列表里能看出这条日程发给了几个人；收件人看到的是 1"""
    ids = _student_ids(b)
    b.client.post("/api/schedules", json=_make_schedule(
        title="班会", target_student_ids=[ids[ROSTER[1][0]], ids[ROSTER[2][0]]]))
    b.client.post("/api/schedules", json=_make_schedule(
        title="自己的", start="2026-10-14T08:00:00", end="2026-10-14T09:00:00"))

    by = {it["title"]: it for it in _list(b.client)}
    # 自己 + 乙 + 丙 = 3
    assert by["班会"]["share_count"] == 3
    assert by["自己的"]["share_count"] == 1

    # 乙是收件人：没有资格再往下发，所以只算他自己那份
    student2 = as_student(b, 1)
    assert [it["share_count"] for it in _list(student2) if it["title"] == "班会"] == [1]


def test_share_update_fans_out(b):
    """建这条日程的人改标题，勾了「所有人的」就下发到乙丙的表里"""
    ids = _student_ids(b)
    b.client.post("/api/schedules", json=_make_schedule(
        title="班会", target_student_ids=[ids[ROSTER[1][0]], ids[ROSTER[2][0]]]))
    sid = [it for it in _list(b.client) if it["title"] == "班会"][0]["id"]

    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-12", "share": True},
        json={"title": "班会（改）", "scope": "all"},
    )
    assert r.status_code == 200, r.text

    assert _items_of(_list(b.client), "班会") == [("2026-10-12", "班会（改）", 3)]
    for idx in (1, 2):
        assert _items_of(_list(as_student(b, idx)), "班会") == [
            ("2026-10-12", "班会（改）", 1),
        ]
    # 三条都在库里（每人一条），不是只改了自己那条
    assert b.rows("SELECT COUNT(*) FROM schedules WHERE title='班会（改）'")[0][0] == 3


def test_share_update_single_fans_out(b):
    """「仅本次」下发：三个人表里都只剩本次改名，前后几次照样在

    这是分配出去的重复日程最常见的改法。它在每个人表里都要把模板切成
    前半段/后半段（见 update_schedule 的 SCOPE_SINGLE），任何一段没落好，
    下发之后别人的页面上就会多一次或者少一次。
    """
    ids = _student_ids(b)
    data = _make_schedule(
        title="班会", start="2026-10-12T08:00:00", end="2026-10-12T09:00:00",
        repeat_type="weekly", repeat_end_date="2026-10-26",
        target_student_ids=[ids[ROSTER[1][0]], ids[ROSTER[2][0]]],
    )
    assert b.client.post("/api/schedules", json=data).status_code == 200

    sid = [it for it in _list(b.client)
           if it["title"] == "班会" and it["instance_date"] == "2026-10-19"][0]["id"]
    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-19", "share": True},
        json={"title": "班会（本次改的）", "scope": "single"},
    )
    assert r.status_code == 200, r.text

    titles = [("2026-10-12", "班会"), ("2026-10-19", "班会（本次改的）"),
              ("2026-10-26", "班会")]
    # 甲自己看到 share_count=3，收件人看到 1；日期和标题三个人必须一模一样
    assert [(d, t) for d, t, _ in _items_of(
        _list(b.client, "2026-10-12", "2026-10-31"), "班会")] == titles
    for idx in (1, 2):
        got = _items_of(_list(as_student(b, idx), "2026-10-12", "2026-10-31"), "班会")
        assert got == [(d, t, 1) for d, t in titles], got
    # 每人被拆成三条：前半段、本次、后半段
    assert b.rows("SELECT COUNT(*) FROM schedules WHERE title LIKE '班会%'")[0][0] == 9


def test_update_without_share_keeps_others(b):
    """不带 share 就只改自己那份——「也允许只改自己的」"""
    ids = _student_ids(b)
    b.client.post("/api/schedules", json=_make_schedule(
        title="班会", target_student_ids=[ids[ROSTER[1][0]], ids[ROSTER[2][0]]]))
    sid = [it for it in _list(b.client) if it["title"] == "班会"][0]["id"]

    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-12"},
        json={"title": "班会（只改我的）", "scope": "all"},
    )
    assert r.status_code == 200, r.text

    assert _items_of(_list(b.client), "班会") == [("2026-10-12", "班会（只改我的）", 3)]
    for idx in (1, 2):
        assert _items_of(_list(as_student(b, idx)), "班会") == [("2026-10-12", "班会", 1)]


def test_share_delete_fans_out(b):
    """勾了「所有人的」删除，乙丙表里那条也一并没掉"""
    ids = _student_ids(b)
    b.client.post("/api/schedules", json=_make_schedule(
        title="班会", target_student_ids=[ids[ROSTER[1][0]], ids[ROSTER[2][0]]]))
    sid = [it for it in _list(b.client) if it["title"] == "班会"][0]["id"]

    r = b.client.delete(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-12", "scope": "all", "share": True},
    )
    assert r.status_code == 200, r.text
    assert r.json()["deleted"] == 3

    for client in (b.client, as_student(b, 1), as_student(b, 2)):
        assert not [it for it in _list(client) if it["title"] == "班会"]

    # 对照组：不勾就只删自己那条
    b.client.post("/api/schedules", json=_make_schedule(
        title="春游", target_student_ids=[ids[ROSTER[1][0]], ids[ROSTER[2][0]]]))
    sid = [it for it in _list(b.client) if it["title"] == "春游"][0]["id"]
    r = b.client.delete(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-12", "scope": "all"},
    )
    assert r.json()["deleted"] == 1
    assert not [it for it in _list(b.client) if it["title"] == "春游"]
    assert [it for it in _list(as_student(b, 1)) if it["title"] == "春游"]


def test_share_only_by_creator(b):
    """收到的日程只能改自己那份：勾「所有人的」返回 403，不是静默降级"""
    ids = _student_ids(b)
    b.client.post("/api/schedules", json=_make_schedule(
        title="班会", target_student_ids=[ids[ROSTER[1][0]]]))
    sid = [it for it in _list(b.client) if it["title"] == "班会"][0]["id"]

    # 乙点开的是**自己那条**（库里每人一行，ID 各不同）
    student2 = as_student(b, 1)
    own = [it for it in _list(student2) if it["title"] == "班会"][0]["id"]

    r = student2.put(
        f"/api/schedules/item/{own}",
        params={"instance_date": "2026-10-12", "share": True},
        json={"title": "乙偷改全班", "scope": "all"},
    )
    assert r.status_code == 403, r.text
    # 报错之前不能已经动过库（哪怕只动了他自己那条）
    assert b.rows("SELECT COUNT(*) FROM schedules WHERE title='乙偷改全班'")[0][0] == 0
    assert _items_of(_list(b.client), "班会") == [("2026-10-12", "班会", 2)]

    # 不勾就随便改，只动自己那条
    r = student2.put(
        f"/api/schedules/item/{own}",
        params={"instance_date": "2026-10-12"},
        json={"title": "班会（乙改的）", "scope": "all"},
    )
    assert r.status_code == 200, r.text
    assert _items_of(_list(b.client), "班会") == [("2026-10-12", "班会", 2)]
    assert _items_of(_list(student2), "班会") == [("2026-10-12", "班会（乙改的）", 1)]


def test_import_share_fans_out(b):
    """导入的课表也能下发，且只命中同一条原始日程，不波及同批次别的课"""
    xlsx_bytes = _make_schedule_xlsx()
    r = b.client.post(
        "/api/schedules/import/preview",
        files={"file": ("课表.xlsx", xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    r = b.client.post("/api/schedules/import", json={
        "token": r.json()["token"], "batch_name": "秋季课表",
        "target_all": True, "target_student_ids": [], "selected_indices": [],
    })
    assert r.status_code == 200, r.text

    # 高数两周两块，是**两条**互不相干的日程（同批次、同标题、不同日期）
    mine = _items_of(_list(b.client), "高等数学")
    assert mine == [("2026-10-12", "高等数学", 3), ("2026-10-19", "高等数学", 3)]
    sid = [it for it in _list(b.client)
           if it["title"] == "高等数学" and it["instance_date"] == "2026-10-12"][0]["id"]

    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-12", "share": True},
        json={"title": "高等数学（改）", "scope": "all"},
    )
    assert r.status_code == 200, r.text

    for idx in (1, 2):
        # 10-12 那天跟着变了；10-19 那条纹丝不动
        assert _items_of(_list(as_student(b, idx)), "高等数学") == [
            ("2026-10-12", "高等数学（改）", 1),
            ("2026-10-19", "高等数学", 1),
        ]
    # 三人各两条高数：10-12 那三条都改了名，剩下三条是 10-19 的
    assert b.rows("SELECT COUNT(*) FROM schedules WHERE title='高等数学'")[0][0] == 3
    assert b.rows("SELECT COUNT(*) FROM schedules WHERE title='高等数学（改）'")[0][0] == 3


def test_share_skips_drifted_copy(b):
    """别人自己把某一次删掉了，下发不能把他的日程静默改没

    乙先自己删掉 10-26 那一次 —— 他那条模板当场被拆成两段。甲再用
    「本次及以后」下发时，乙那两段其实都不包含 10-26，必须整个跳过：
    硬套的话，后半段（11-02）会被截到 10-25 截止，展开成空区间，
    在乙的页面上等于那次日程凭空消失。
    """
    ids = _student_ids(b)
    data = _make_schedule(
        title="班会", start="2026-10-12T08:00:00", end="2026-10-12T09:00:00",
        repeat_type="weekly", repeat_end_date="2026-11-02",
        target_student_ids=[ids[ROSTER[1][0]], ids[ROSTER[2][0]]],
    )
    assert b.client.post("/api/schedules", json=data).status_code == 200

    # 乙自己删掉中间那次（10-26）
    student2 = as_student(b, 1)
    sid = [it for it in _list(student2) if it["title"] == "班会"][0]["id"]
    r = student2.delete(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-26", "scope": "single"},
    )
    assert r.status_code == 200, r.text
    assert [d for d, _, _ in _items_of(_list(student2), "班会")] == [
        "2026-10-12", "2026-10-19", "2026-11-02",
    ]
    assert b.rows("SELECT COUNT(*) FROM schedules WHERE title='班会' AND student_id=?",
                  ids[ROSTER[1][0]])[0][0] == 2  # 被拆成两段

    # 甲用「本次及以后」下发 10-26 起的那一段
    sid = [it for it in _list(b.client)
           if it["title"] == "班会" and it["instance_date"] == "2026-10-26"][0]["id"]
    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-26", "share": True},
        json={"title": "班会（改）", "scope": "forward"},
    )
    assert r.status_code == 200, r.text

    # 丙正常收到：10-26 和 11-02 变成新标题
    assert _items_of(_list(as_student(b, 2)), "班会") == [
        ("2026-10-12", "班会", 1),
        ("2026-10-19", "班会", 1),
        ("2026-10-26", "班会（改）", 1),
        ("2026-11-02", "班会（改）", 1),
    ]
    # 乙那两段都没被动过：还是两行、11-02 那次还在、标题还是旧的
    assert b.rows("SELECT COUNT(*) FROM schedules WHERE title LIKE '班会%' AND student_id=?",
                  ids[ROSTER[1][0]])[0][0] == 2
    assert _items_of(_list(student2), "班会") == [
        ("2026-10-12", "班会", 1),
        ("2026-10-19", "班会", 1),
        ("2026-11-02", "班会", 1),
    ]


def test_share_multi_day_all_day(b):
    """跨天的全天事件是「一天一条」，每天各自能下发；新增的天沿用原来那批人"""
    ids = _student_ids(b)
    targets = [ids[ROSTER[1][0]], ids[ROSTER[2][0]]]
    # 前端的 saveAllDay 就是一天一次请求，这里照做
    for day in ("2026-10-12", "2026-10-13", "2026-10-14"):
        r = b.client.post("/api/schedules", json=_make_schedule(
            title="运动会", start=f"{day}T00:00:00", end=f"{day}T23:59:00",
            all_day=True, target_student_ids=targets,
        ))
        assert r.status_code == 200, r.text

    student2 = as_student(b, 1)
    assert [d for d, _, _ in _items_of(_list(student2), "运动会")] == [
        "2026-10-12", "2026-10-13", "2026-10-14",
    ]

    # 改中间那天：只有那天跟着变，首尾两天不受影响
    sid = [it for it in _list(b.client)
           if it["title"] == "运动会" and it["instance_date"] == "2026-10-13"][0]["id"]
    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-13", "share": True},
        json={"title": "运动会（改）", "scope": "all"},
    )
    assert r.status_code == 200, r.text
    assert [t for _, t, _ in _items_of(_list(student2), "运动会")] == [
        "运动会", "运动会（改）", "运动会",
    ]

    # 往后拖一天：新加的那天得发给原来那批人（share_from 沿用分组），
    # 不然「3 天改成 4 天」之后，第 4 天只有甲自己有
    r = b.client.post("/api/schedules", json=_make_schedule(
        title="运动会", start="2026-10-15T00:00:00", end="2026-10-15T23:59:00",
        all_day=True, share_from=sid,
    ))
    assert r.status_code == 200, r.text
    assert [d for d, _, _ in _items_of(_list(student2), "运动会")] == [
        "2026-10-12", "2026-10-13", "2026-10-14", "2026-10-15",
    ]

    # share_from 指向自己收来的（不是自己分配的）日程时不认账，只建给自己 ——
    # 不然班委之间互相传个日程 ID 就能往别人表里塞东西。
    # 得先把乙提成班委，否则被「不是班委」这一层挡住，测不到这个判断
    r = b.client.put(f"/api/admin/users/{ROSTER[1][0]}/role", json={"role": "committee"})
    assert r.status_code == 200, r.text
    committee2 = as_student(b, 1)
    received = [it for it in _list(committee2) if it["title"] == "运动会"][0]
    r = committee2.post("/api/schedules", json=_make_schedule(
        title="乙私自加的", start="2026-10-16T00:00:00", end="2026-10-16T23:59:00",
        all_day=True, share_from=received["id"],
    ))
    assert r.status_code == 200, r.text
    assert [it["title"] for it in _list(committee2) if it["title"] == "乙私自加的"]
    assert not [it for it in _list(b.client) if it["title"] == "乙私自加的"]


# ---------- 查询范围限制 ----------

def test_range_limit(b):
    """查询超过 92 天会被拒绝"""
    r = b.client.get("/api/schedules", params={
        "start_date": "2026-01-01", "end_date": "2026-12-31",
    })
    assert r.status_code == 400
    assert "92" in r.json()["detail"]


# ---------- 回归：重复必设结束日、单双周基准 ----------

def test_repeat_without_end_rejected(b):
    """后端兜底：重复日程必须设结束日期，直接调接口也不能漏"""
    data = _make_schedule(
        title="缺结束日的重复",
        start="2026-10-12T08:00:00",
        end="2026-10-12T09:00:00",
        repeat_type="weekly",
        repeat_end_date=None,
    )
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 400
    assert "结束日期" in r.json()["detail"]


def test_cross_day_schedule_rejected(b):
    """跨天日程（22:00 → 次日 01:00）一律拒收

    模型里没有「跨天日程」这个概念：跨天的全天事件是摊成一天一条再由前端合回
    一条，前端表单也拦着。只有后端这个校验器以前比的是完整 datetime，于是
    「次日 01:00」算「晚于」被放行 —— 而 _to_instance 展开实例时只取 .time()，
    页面上就成了结束(01:00)早于开始(22:00)，iCal 那边按时长算却是次日 01:00，
    两边对不上。关掉之后这个状态压根不存在。
    """
    r = b.client.post("/api/schedules", json=_make_schedule(
        title="跨天日程", start="2026-10-12T22:00:00", end="2026-10-13T01:00:00",
    ))
    assert r.status_code == 400, r.text
    assert "当天结束" in r.json()["detail"]
    assert b.rows("SELECT id FROM schedules WHERE title = ?", "跨天日程") == []


def test_cross_day_update_rejected(b):
    """编辑时只把结束时刻挪到次日也要拦住

    以前要 start_time / end_time 两个都传才校验，光改 end_time 就能绕过去。
    现在校验的是「改完之后」的那一对时刻（没传的那个取库里的现值）。
    """
    r = b.client.post("/api/schedules", json=_make_schedule(
        title="只改结束时刻", start="2026-10-12T22:00:00", end="2026-10-12T23:00:00",
    ))
    assert r.status_code == 200, r.text
    sid = r.json()["id"]

    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-12"},
        json={"scope": "all", "end_time": "2026-10-13T01:00:00"},
    )
    assert r.status_code == 400, r.text
    assert "当天结束" in r.json()["detail"]

    # 库里那条没被动过（end_time 还是当天 23:00）
    end = b.rows("SELECT end_time FROM schedules WHERE id = ?", sid)[0][0]
    assert str(end).startswith("2026-10-12 23:00")


def test_odd_week_expansion_honors_week_start(b):
    """带周次基准的日程（课表导入、带锚点下发）：以基准的周次号为准

    直接构造模板和区间：基准 2026-08-31（第 1 周·单），
    日程 start=2026-09-08（第 2 周·双），repeat_type=even_week。
    基准生效的话 09-08 本身就该是「双周」并出现；
    若退化到「以 start 所在周为第1周」，09-08 会被当成第1周·单而被跳过。
    """
    from app.services.schedule_service import Period, _next_occurrence
    from app.models.schedule import REPEAT_EVEN_WEEK
    from app.models import Schedule

    periods = [Period(
        week_start=date(2026, 8, 31),       # 第 1 周
        first_day=date(2026, 8, 31),
        last_day=date(2026, 9, 27),         # 覆盖到第 4 周周日
    )]
    tpl = Schedule(
        student_id=1, title="双周课", location="", description="",
        start_time=datetime(2026, 9, 8, 9, 0),
        end_time=datetime(2026, 9, 8, 10, 0),
        repeat_type=REPEAT_EVEN_WEEK,
        repeat_end_date=date(2026, 9, 30),
        notify_enabled=False, notify_minutes_before=10,
        source_type="personal",
        # 导入课表时写下的编号基准（有它才按课表编号）
        week_start=date(2026, 8, 31),
    )
    assert _next_occurrence(tpl, date(2026, 9, 8), periods) == date(2026, 9, 8)
    # 下一个偶数周（第4周）
    assert _next_occurrence(tpl, date(2026, 9, 9), periods) == date(2026, 9, 22)


def test_odd_week_skips_weeks_without_a_number(b):
    """没有周次号的周：单双周日程一律不出现（严格按周次号的奇偶）

    课表只覆盖到第 2 周，而日程的重复结束日远在那之后 —— 第 3 周及以后没有周次号，
    既不是单周也不是双周，整周跳过，不会一路按奇偶排下去。
    """
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    r = b.client.post(
        "/api/schedules/import/preview",
        files={"file": ("课表.xlsx", _make_schedule_xlsx(date(2026, 10, 12)), mime)},
    )
    assert r.status_code == 200, r.text
    r = b.client.post("/api/schedules/import", json={
        "token": r.json()["token"], "batch_name": "课表",
        "target_all": True, "target_student_ids": [], "selected_indices": [],
    })
    assert r.status_code == 200, r.text

    # 周二的单周课：第 1 周（单）出现；第 3 周之后没有周次号，一条都不该有
    r = b.client.post("/api/schedules", json=_make_schedule(
        title="单周选修", start="2026-10-13T08:00:00", end="2026-10-13T09:00:00"))
    b.run(
        "UPDATE schedules SET repeat_type='odd_week', repeat_end_date='2026-11-30' WHERE id=?",
        r.json()["id"],
    )

    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-11-30",
    })
    dates = [it["instance_date"] for it in r.json() if it["title"] == "单周选修"]
    assert dates == ["2026-10-13"]   # 只有第 1 周有周次号

    # 订阅那一半也得一样：只有这一条，且不带 RRULE（RRULE 没法表达「中间整周空掉」）
    text = _ical_text(b)
    assert _ical_dates(text)["单周选修"] == ["2026-10-13"], text
    assert "RRULE" not in text


def test_ical_single_schedule_ignores_stale_repeat_end_date(b):
    """单次日程带着一个早于开始日的 repeat_end_date：两边都还得有这一条

    接口不拦这种脏数据（repeat_type=none 时 repeat_end_date 没人管）。页面一直
    是「单次日程不看它」的口径，订阅要是照 _next_occurrence 的写法（把它当结束日）
    去展开，这条就会在订阅里凭空消失。
    """
    r = b.client.post("/api/schedules", json=_make_schedule(
        title="单次脏数据", start="2026-10-12T08:00:00", end="2026-10-12T09:00:00"))
    sid = r.json()["id"]
    b.run("UPDATE schedules SET repeat_end_date='2026-09-01' WHERE id=?", sid)

    assert _page_dates(b, "2026-10-01", "2026-10-31")["单次脏数据"] == ["2026-10-12"]
    text = _ical_text(b)
    assert _ical_dates(text)["单次脏数据"] == ["2026-10-12"], text


def test_manual_odd_week_without_week_start_counts_from_its_own_first_week(b):
    """手动建的单双周（没走课表导入）：以它自己的开始周为第 1 周

    课表基准是 2026-08-31（第 1 周），于是 09-07 在课表里是第 2 周（双）；
    这条日程却是从 09-07 起按单周重复的。留空 week_start 就得数它自己的第 1 周，
    写上周次基准（导入/带锚点下发）才按课表编号 —— 两套口径各管各的，
    否则同一条共享出去的单双周日程会在各人的表里落到相反的周（见 _periods_for）。
    """
    from app.services.schedule_service import Period, _next_occurrence
    from app.models.schedule import REPEAT_ODD_WEEK
    from app.models import Schedule

    periods = [Period(
        week_start=date(2026, 8, 31),   # 课表的第 1 周
        first_day=date(2026, 8, 31),
        last_day=date(2026, 10, 11),    # 覆盖到 10-11 那周的周日
    )]

    def make(**kw):
        return Schedule(
            student_id=1, title="值班", location="", description="",
            start_time=datetime(2026, 9, 7, 9, 0),
            end_time=datetime(2026, 9, 7, 10, 0),
            repeat_type=REPEAT_ODD_WEEK,
            repeat_end_date=date(2026, 10, 11),
            notify_enabled=False, notify_minutes_before=10,
            source_type="personal",
            **kw,
        )

    tpl = make()   # week_start 留空 = 学生自己建的
    # 09-07 是它自己的第 1 周（单）→ 出现（按课表编号是第 2 周·双，会被跳过）
    assert _next_occurrence(tpl, date(2026, 9, 7), periods) == date(2026, 9, 7)
    # 09-14 是它自己的第 2 周，跳过；下一个单周是 09-21
    assert _next_occurrence(tpl, date(2026, 9, 8), periods) == date(2026, 9, 21)

    # 写上课表给的基准，就按课表编号：09-07 是第 2 周（双），跳过，第一次是 09-14
    tpl.week_start = date(2026, 8, 31)
    assert _next_occurrence(tpl, date(2026, 9, 7), periods) == date(2026, 9, 14)


def test_manual_odd_week_fallback_consistent(b):
    """没导入过课表时，手动建的单双周覆盖到重复结束日为止

    一条 week_start 都没有 → 学期区间退回「按现有日程的跨度兜一个」。这个跨度必须
    算到 repeat_end_date，否则区间只有一周宽，第二次以后的重复全落在区间外。
    （编号基准不看这个跨度，见 _periods_for。）
    """
    data = _make_schedule(
        title="双周值班",
        start="2026-10-12T08:00:00",  # 被当作第1周（单）
        end="2026-10-12T09:00:00",
        repeat_type="even_week",
        repeat_end_date="2026-10-26",
    )
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 200, r.text

    # 第1周（单）当天、第3周都不该出现；第2周（双）出现
    r = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-25",
    })
    items = r.json()
    dates = [it["instance_date"] for it in items]
    assert "2026-10-12" not in dates
    assert "2026-10-19" in dates


def test_shared_odd_week_agrees_across_students(b):
    """同一条共享出去的单双周日程，在每个人表里落到同样的日子

    回归：兜底区间一度是「该学生全部日程的跨度」，谁的表里多一条更早的日程，
    编号基准就整体往前挪，奇偶跟着翻 —— 乙的表里整条日程会平移一周，而甲丙正常。
    现在手动日程一律数它自己的第 1 周，与其他日程无关。
    """
    ids = _student_ids(b)
    # 乙比甲多一条更早的日程：曾经正是它把乙的兜底基准往前拉了一周
    as_student(b, 1).post("/api/schedules", json=_make_schedule(
        title="乙的更早日程", start="2026-10-07T08:00:00", end="2026-10-07T09:00:00"))

    r = b.client.post("/api/schedules", json=_make_schedule(
        title="单周值班",
        start="2026-10-12T08:00:00", end="2026-10-12T09:00:00",
        repeat_type="odd_week", repeat_end_date="2026-11-30",
        target_student_ids=[ids[ROSTER[1][0]], ids[ROSTER[2][0]]],
    ))
    assert r.status_code == 200, r.text

    # 甲：第 1 周（10-12 那周）起，单周各一次
    mine = [d for d, _, _ in _items_of(_list(b.client, "2026-10-12", "2026-11-30"), "单周值班")]
    assert mine == ["2026-10-12", "2026-10-26", "2026-11-09", "2026-11-23"]

    # 乙、丙看到的必须一模一样（乙那条 10-07 的日程不该影响任何人的周次）
    for idx in (1, 2):
        theirs = [
            d for d, _, _ in _items_of(
                _list(as_student(b, idx), "2026-10-12", "2026-11-30"), "单周值班")
        ]
        assert theirs == mine


def test_split_keeps_week_start(b):
    """「只改本次 / 本次及以后」拆出来的新条目，得跟着留上周次基准

    _to_instance 是拿模板的字段去 create_schedule 的，漏一项就丢一项：早期版本
    漏了 week_start，拆完的条目没有编号基准，单双周改用它自己的开始周去数。
    """
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    r = b.client.post(
        "/api/schedules/import/preview",
        files={"file": ("课表.xlsx", _make_schedule_xlsx(weeks=6), mime)},
    )
    assert r.status_code == 200, r.text
    r = b.client.post("/api/schedules/import", json={
        "token": r.json()["token"], "batch_name": "课表",
        "target_all": True, "target_student_ids": [], "selected_indices": [],
    })
    assert r.status_code == 200, r.text

    # 高数（周一那块）改成双周、英语（周三那块）改成单周，都按课表写下的编号
    b.run("UPDATE schedules SET repeat_type='even_week', repeat_end_date='2026-11-30'"
          " WHERE title LIKE '高等数学%' AND start_time LIKE '2026-10-12%'")
    b.run("UPDATE schedules SET repeat_type='odd_week', repeat_end_date='2026-11-30'"
          " WHERE title LIKE '大学英语%'")
    math_id = b.rows("SELECT id FROM schedules WHERE title LIKE '高等数学%'"
                     " AND start_time LIKE '2026-10-12%'")[0][0]
    eng_id = b.rows("SELECT id FROM schedules WHERE title LIKE '大学英语%'")[0][0]

    # 双周是 10-19 / 11-02 / 11-16：改掉 10-19 这一次 → 单次条目 + 后半段模板
    r = b.client.put(f"/api/schedules/item/{math_id}",
                     params={"instance_date": "2026-10-19", "share": False},
                     json={"title": "高数（调课）", "scope": "single"})
    assert r.status_code == 200, r.text
    # 单周是 10-14 / 10-28 / 11-11：从 10-28 起改 → 后半段模板
    r = b.client.put(f"/api/schedules/item/{eng_id}",
                     params={"instance_date": "2026-10-28", "share": False},
                     json={"title": "英语（改）", "scope": "forward"})
    assert r.status_code == 200, r.text

    assert b.rows("SELECT week_start FROM schedules WHERE title='高数（调课）'") == [
        ("2026-10-12",)]
    assert b.rows("SELECT week_start FROM schedules WHERE title='英语（改）'") == [
        ("2026-10-12",)]


# ---------- 周次（跟表名，不靠前端估算） ----------

def test_semester_week_follows_sheet_name(b):
    """周次跟表名走：导入时记下「学期第 1 周的周一」，页面按它算第几周

    夹具第 1 周周一是 10/12、第 2 周是 10/19，由表名折算出的基准两周边有同一个。
    前端曾经按「9 月 1 日所在周」自己估第 1 周，比表名早两周 —— 表名「第4周」
    的课在页面上会显示成第 6 周（用户报的就是这个）。
    """
    r = b.client.post(
        "/api/schedules/import/preview",
        files={"file": ("课表.xlsx", _make_schedule_xlsx(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200, r.text
    r = b.client.post("/api/schedules/import", json={
        "token": r.json()["token"],
        "batch_name": "秋季课表",
        "target_all": True,
        "target_student_ids": [],
        "selected_indices": [],
    })
    assert r.status_code == 200, r.text

    # 两个 sheet 折算出的基准必须一致，并且都落到了日程上
    assert b.rows("SELECT DISTINCT week_start FROM schedules") == [("2026-10-12",)]

    # 接口把周次范围交回前端（页面上的「第几周」按它算）。
    # first_day/last_day = 这份课表实际覆盖到的那几周，上界 = 第 2 周周日 10/25 ——
    # 再往后的周就不编号了，不会一路算到「第 35 周」
    info = b.client.get("/api/schedules/semester")
    assert info.status_code == 200, info.text
    assert info.json()["periods"] == [
        {"week_start": "2026-10-12", "first_day": "2026-10-12", "last_day": "2026-10-25"},
    ]


def test_semester_periods_allow_a_new_term_to_restart(b):
    """一份课表一个区间：下学期再导入一份新课表，又是一个新的「第 1 周」

    区间是「按 week_start 分组、取这组课表覆盖到的头尾」推出来的，所以两份课表
    各是各的，页面按当前这周落在哪个区间挑周次，区间之外（寒暑假）不编号。
    """
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    for name, first_monday in (("上学期", date(2026, 10, 12)), ("下学期", date(2027, 3, 1))):
        r = b.client.post(
            "/api/schedules/import/preview",
            files={"file": (f"{name}.xlsx", _make_schedule_xlsx(first_monday), mime)},
        )
        assert r.status_code == 200, r.text
        r = b.client.post("/api/schedules/import", json={
            "token": r.json()["token"],
            "batch_name": name,
            "target_all": True,
            "target_student_ids": [],
            "selected_indices": [],
        })
        assert r.status_code == 200, r.text

    info = b.client.get("/api/schedules/semester")
    assert info.status_code == 200, info.text
    assert info.json()["periods"] == [
        {"week_start": "2026-10-12", "first_day": "2026-10-12", "last_day": "2026-10-25"},
        {"week_start": "2027-03-01", "first_day": "2027-03-01", "last_day": "2027-03-14"},
    ]


def test_semester_info_falls_back_to_earliest_schedule(b):
    """没写过 week_start 的老数据：退回按现有日程的跨度给一个区间

    升级前导入的课表没有这一列，不能让周次因此整个消失。
    """
    data = _make_schedule(title="老日程", start="2026-10-15T08:00:00", end="2026-10-15T09:40:00")
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 200, r.text
    # 手动创建的日程不带基准（只有课表导入才写）
    assert b.rows("SELECT DISTINCT week_start FROM schedules") == [(None,)]

    info = b.client.get("/api/schedules/semester")
    assert info.status_code == 200, info.text
    assert info.json()["periods"] == [
        # 10/15 是周四 → 所属周周一是 10/12、周日是 10/18
        {"week_start": "2026-10-12", "first_day": "2026-10-12", "last_day": "2026-10-18"},
    ]


def test_semester_info_empty_without_schedules(b):
    """一条日程都没有时返回空 —— 前端就不显示周次，而不是硬估一个错的"""
    info = b.client.get("/api/schedules/semester")
    assert info.status_code == 200, info.text
    assert info.json()["periods"] == []


# ---------- 导入的周次锚点 ----------

_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _preview(b, filename: str, xlsx: bytes, anchor: Optional[str] = None):
    """调一次导入预览（anchor 传字符串日期就是用户指定锚点）"""
    files = {"file": (filename, xlsx, _MIME)}
    data = {"anchor_monday": anchor} if anchor else None
    return b.client.post("/api/schedules/import/preview", files=files, data=data)


def test_import_anchor_overrides_header(b):
    """用户指定的锚点优先 —— 表头日期只是建议值

    文件是要求「第 8–10 周怎么办」时那张补课表：表头写的是 10/26 起，
    但用户指认第 8 周的周一其实是 11/02，整份表就该整体往后挪一周，
    编号基准（week_start）也跟着变。
    """
    xlsx = _make_schedule_xlsx(date(2026, 10, 26), weeks=3, first_week_no=8)

    # 不给锚点：按表头猜，第 8 周 = 10/26
    r = _preview(b, "补课表.xlsx", xlsx)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["week_first"] == 8 and data["week_last"] == 10
    assert data["anchor_monday"] == "2026-10-26"
    assert [it["start_date"] for it in data["items"] if "高等数学" in it["title"]] == [
        "2026-10-26", "2026-11-02", "2026-11-09",
    ]

    # 用户改成 11/02：所有日期整体挪一周，并提醒表头和锚点对不上
    r = _preview(b, "补课表.xlsx", xlsx, anchor="2026-11-02")
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["anchor_monday"] == "2026-11-02"
    assert [it["start_date"] for it in data["items"] if "高等数学" in it["title"]] == [
        "2026-11-02", "2026-11-09", "2026-11-16",
    ]
    assert any("表头写的是 2026-10-26" in w for w in data["warnings"]), data["warnings"]


def test_import_anchor_shifts_week_start(b):
    """锚点决定编号基准：补课表只覆盖第 8–10 周，第 1 周由锚点往回推

    week_start（页面上的「第几周」靠它算）要用 11/02 倒推 7 周 = 09/14，
    而不是表头那套 09/07 —— 锚点说了算。
    """
    xlsx = _make_schedule_xlsx(date(2026, 10, 26), weeks=3, first_week_no=8)
    r = _preview(b, "补课表.xlsx", xlsx, anchor="2026-11-02")
    assert r.status_code == 200, r.text
    token = r.json()["token"]
    r = b.client.post("/api/schedules/import", json={
        "token": token, "batch_name": "补课表",
        "target_all": True, "target_student_ids": [], "selected_indices": [],
    })
    assert r.status_code == 200, r.text

    assert b.rows("SELECT DISTINCT week_start FROM schedules") == [("2026-09-14",)]
    # 区间只收在这份表覆盖到的那几周里：11/02（第 8 周）~ 11/22（第 10 周周日）
    info = b.client.get("/api/schedules/semester").json()["periods"]
    assert info == [
        {"week_start": "2026-09-14", "first_day": "2026-11-02", "last_day": "2026-11-22"},
    ]


def test_import_headerless_needs_an_anchor(b):
    """表头没写日期时猜不出锚点：先一条都不解析，等用户指认后再解析

    猜不出硬按「今年」算会把整份课表落到错的日子上，宁可不解析 ——
    响应里带 week_first，页面照着问「第 8 周的周一是哪天」。
    """
    xlsx = _make_schedule_xlsx(date(2026, 10, 26), weeks=3, first_week_no=8,
                               date_in_header=False)

    r = _preview(b, "补课表.xlsx", xlsx)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["week_first"] == 8
    assert data["anchor_monday"] is None
    assert data["items"] == []
    assert any("请先指定第 8 周的周一" in w for w in data["warnings"]), data["warnings"]

    # 用户指认第 8 周的周一 = 10/26，日期就出来了
    r = _preview(b, "补课表.xlsx", xlsx, anchor="2026-10-26")
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["anchor_monday"] == "2026-10-26"
    assert [it["start_date"] for it in data["items"] if "高等数学" in it["title"]] == [
        "2026-10-26", "2026-11-02", "2026-11-09",
    ]


def test_import_anchor_not_monday_snaps_back(b):
    """锚点不是周一：归到它所在周的周一，并说一声（页面上只让选周一，接口不管）"""
    r = _preview(b, "课表.xlsx", _make_schedule_xlsx(), anchor="2026-10-14")  # 周三
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["anchor_monday"] == "2026-10-12"
    assert any("不是周一" in w for w in data["warnings"]), data["warnings"]
    assert any(it["start_date"] == "2026-10-12" for it in data["items"])


def test_calendar_covers_the_whole_day():
    """全天 00:00–24:00；导入至少接受 5 分钟粒度"""
    from app.models.schedule import DAY_END_HOUR, DAY_START_HOUR, IMPORT_MINUTE_STEP

    assert (DAY_START_HOUR, DAY_END_HOUR) == (0, 24)
    assert IMPORT_MINUTE_STEP == 5


def test_import_accepts_five_minute_rows():
    """5 分钟一行的课表：08:00–09:35 的块要落成 08:00–09:40

    结束时间按合并块下一行的真实时间算。以前写死 +10 分钟，
    这种 5 分钟表格会整体少算一半。
    """
    from app.services.schedule_import_service import parse_schedule_excel

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "第1周"
    ws.cell(row=1, column=1, value="时间")
    ws.cell(row=1, column=2, value="周一")
    # 08:00 起每 5 分钟一行，共 25 行到 10:00
    for i in range(25):
        total = 8 * 60 + i * 5
        ws.cell(row=i + 2, column=1, value=f"{total // 60:02d}:{total % 60:02d}")
    # 08:00（第 2 行）到 09:35（第 21 行）
    ws.merge_cells(start_row=2, start_column=2, end_row=21, end_column=2)
    ws.cell(row=2, column=2, value="高等数学\n课程\n教学楼C101")
    buf = BytesIO()
    wb.save(buf)

    items, warnings, _ = parse_schedule_excel(buf.getvalue(), date(2026, 10, 12))
    assert not warnings
    assert len(items) == 1, items
    assert items[0]["start_time"] == "08:00"
    assert items[0]["end_time"] == "09:40"


def test_import_clamps_a_block_ending_at_midnight():
    """时间列最后一格 23:55 的块，结束时间停在 23:59，不会算出 24:00"""
    from app.services.schedule_import_service import _add_minutes

    assert _add_minutes("23:55", 5) == "23:59"
    assert _add_minutes("23:50", 10) == "23:59"


# ---------- 地点坐标（GCJ-02） ----------

def test_create_schedule_with_geo(b):
    """坐标能传进来、存下来、再返回（创建响应 + 列表接口 + 库里三处一致）"""
    data = _make_schedule(title="图书馆自习", lat=39.907501, lng=116.397232)
    r = b.client.post("/api/schedules", json=data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert (body["lat"], body["lng"]) == (39.907501, 116.397232)

    items = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-18",
    }).json()
    assert [(it["lat"], it["lng"]) for it in items] == [(39.907501, 116.397232)]

    assert b.rows("SELECT lat, lng FROM schedules WHERE title = ?", "图书馆自习") == [
        (39.907501, 116.397232)
    ]


def test_schedule_without_geo_returns_null(b):
    """没传坐标就是 null —— 老日程（补列后是 NULL）和手填日程都走这一条，不是 0"""
    r = b.client.post("/api/schedules", json=_make_schedule(title="没坐标"))
    assert r.status_code == 200, r.text
    assert r.json()["lat"] is None and r.json()["lng"] is None

    items = b.client.get("/api/schedules", params={
        "start_date": "2026-10-12", "end_date": "2026-10-18",
    }).json()
    assert all(it["lat"] is None and it["lng"] is None for it in items)


def test_geo_must_be_paired_and_in_range(b):
    """只给一个坐标、或超出合法范围 → 400（不能让半对坐标落库）"""
    r = b.client.post("/api/schedules", json=_make_schedule(title="只有纬度", lat=39.9))
    assert r.status_code == 400, r.text
    assert "经纬度" in r.json()["detail"]

    r = b.client.post("/api/schedules", json=_make_schedule(title="只有经度", lng=116.4))
    assert r.status_code == 400, r.text

    r = b.client.post("/api/schedules", json=_make_schedule(title="越界", lat=91, lng=116.4))
    assert r.status_code == 400, r.text
    r = b.client.post("/api/schedules", json=_make_schedule(title="越界2", lat=39.9, lng=181))
    assert r.status_code == 400, r.text

    assert b.rows("SELECT COUNT(*) FROM schedules") == [(0,)]


def test_update_geo_and_repeat_split_keeps_it(b):
    """改坐标能生效；重复日程「仅本次」拆出来的那条也要带着坐标

    拆分是拿 _to_instance 的 dict 去 create_schedule 的，漏了 lat/lng 的话新记录
    会静默丢坐标（week_start 踩过同一个坑）。
    """
    r = b.client.post("/api/schedules", json=_make_schedule(
        title="周会", repeat_type="weekly", repeat_end_date="2026-11-02",
    ))
    assert r.status_code == 200, r.text
    sid = r.json()["id"]

    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-12"},
        json={"lat": 31.230416, "lng": 121.473701, "scope": "all"},
    )
    assert r.status_code == 200, r.text
    assert (r.json()["lat"], r.json()["lng"]) == (31.230416, 121.473701)

    # 只改中间那一次：拆出来的独立日程要继承坐标
    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-19"},
        json={"title": "周会（改期）", "scope": "single"},
    )
    assert r.status_code == 200, r.text

    rows = b.rows("SELECT title, lat, lng FROM schedules ORDER BY start_time")
    assert [(t, lat, lng) for t, lat, lng in rows] == [
        ("周会", 31.230416, 121.473701),
        ("周会（改期）", 31.230416, 121.473701),
        ("周会", 31.230416, 121.473701),
    ]

    # 只改一个坐标时，另一个取库里现有的：成对成立才放行
    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-12"},
        json={"lat": 30.5, "scope": "all"},
    )
    assert r.status_code == 200, r.text
    assert (r.json()["lat"], r.json()["lng"]) == (30.5, 121.473701)

    # 整份表单提交时会把没填的坐标一起带成 null：那是「不改」，不是「清空」，
    # 更不能被当成「只填了一个坐标」拒掉
    r = b.client.put(
        f"/api/schedules/item/{sid}",
        params={"instance_date": "2026-10-12"},
        json={"title": "周会（改名）", "lat": None, "lng": None, "scope": "all"},
    )
    assert r.status_code == 200, r.text
    assert (r.json()["lat"], r.json()["lng"]) == (30.5, 121.473701)


def test_ical_geo_property(b):
    """订阅里带标准 GEO + Apple 结构化地点；没坐标的日程两条都不写

    Apple 日历的交通时间/出发提醒认的是 X-APPLE-STRUCTURED-LOCATION，
    单有标准 GEO 往往被忽略、仍按 LOCATION 短语重新地理编码。
    """
    b.client.post("/api/schedules", json=_make_schedule(
        title="有坐标", location="图书馆", lat=39.907501, lng=116.397232,
    ))
    b.client.post("/api/schedules", json=_make_schedule(
        title="校园坐标", location="多地点教室",
        start="2026-10-14T08:00:00", end="2026-10-14T09:00:00",
        lat=12.34, lng=56.78,
    ))
    b.client.post("/api/schedules", json=_make_schedule(
        title="带引号的地点", location='甲"乙,丙;丁',
        start="2026-10-15T08:00:00", end="2026-10-15T09:00:00",
        lat=1.0, lng=2.0,
    ))
    b.client.post("/api/schedules", json=_make_schedule(
        title="没坐标", start="2026-10-13T08:00:00", end="2026-10-13T09:00:00",
    ))

    text = _ical_text(b)
    # 折行还原后再比对整条属性（X-APPLE 这行肯定超过 75 字节）
    flat = text.replace("\r\n ", "")
    events = {ev.get("SUMMARY"): ev for ev in _ical_events(text)}
    assert events["有坐标"]["GEO"] == "39.907501;116.397232"
    assert events["有坐标"]["X-APPLE-STRUCTURED-LOCATION"] == "geo:39.907501,116.397232"
    # 不认识的坐标退回 location 文本
    assert 'X-TITLE="图书馆":geo:39.907501,116.397232' in flat
    # 校园坐标按主地点名，不把「多地点教室」原样塞进去
    assert 'X-TITLE="教学楼 A":geo:12.34,56.78' in flat
    # 参数里的双引号按 RFC 6868 编码，逗号/分号在引号内不拆参数
    assert 'X-TITLE="甲^\'乙,丙;丁":geo:1.0,2.0' in flat
    assert "X-APPLE-STRUCTURED-LOCATION" not in events["没坐标"]
    assert "GEO" not in events["没坐标"]
