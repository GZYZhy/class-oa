"""投票收集的回归：六种题型、每人多份、竞争名额、匿名、删除、导出、开始提醒

    cd backend && python -m pytest tests/test_collections.py -q

重点盯：
1. 一人多份与全局竞争名额（删除后释放）；
2. 必答/选答和每种题型的服务端校验；
3. 匿名只隐藏身份，不隐藏「谁还没填」；
4. 管理员才能删答卷；
5. 开始提醒只在到点发一次，改开始时间跟着新时间走（通过直接调调度器 tick 验证）。
"""

from datetime import datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from conftest import png_bytes

PW = "initpass123"
ADMIN = "2024000001"
ADMIN_PW = "010001"
CREATOR = "2024000002"
CREATOR_PW = "010002"
VOTER = "2024000003"
VOTER_PW = "010003"
OTHER = "2024000004"
OTHER_PW = "010004"
THIRD = "2024000005"
THIRD_PW = "010005"

PEOPLE = [
    (CREATOR, "学生02", CREATOR_PW),
    (VOTER, "学生03", VOTER_PW),
    (OTHER, "学生04", OTHER_PW),
    (THIRD, "学生05", THIRD_PW),
]
PNG = png_bytes()


def _world(backend, db_name):
    # 关掉后台开始提醒线程，提醒由用例自己调 tick，结果才确定
    b = backend(
        db=db_name, BOOTSTRAP_PASSWORD=PW, START_NOTIFY_SCHEDULER="0"
    )
    assert b.login(password=PW).status_code == 200

    r = b.client.post("/api/students", json={
        "student_no": ADMIN, "name": "甲", "id_card": "110101200501010001",
    })
    assert r.status_code < 300, r.text
    for no, name, pw in PEOPLE:
        r = b.client.post("/api/students", json={
            "student_no": no, "name": name,
            "id_card": f"11010120050101{no[-4:]}",
        })
        assert r.status_code < 300, r.text

    assert b.client.post("/api/setup/finish", json={"admin_student_no": ADMIN}).status_code == 200
    # 初始化完成后先以管理员登录，才能给发起人提权
    b.logout()
    assert b.login(student_no=ADMIN, password=ADMIN_PW).status_code == 200
    # 出卷人是班委，管理员是另一个人
    assert b.client.put(f"/api/admin/users/{CREATOR}/role", json={"role": "committee"}).status_code < 300
    b.logout()
    assert b.login(student_no=CREATOR, password=CREATOR_PW).status_code == 200
    return b


@pytest.fixture
def cb(backend):
    return _world(backend, "collections.db")


def _login(b, student_no, password):
    b.logout()
    assert b.login(student_no=student_no, password=password).status_code == 200


def _stamp(hours):
    return (datetime.now() + timedelta(hours=hours)).strftime("%Y-%m-%d %H:%M")


def _q(qtype, title="题目", *, required=True, options=None, max_choices=None):
    return {
        "title": title,
        "qtype": qtype,
        "required": required,
        "options": options or [],
        "max_choices": max_choices,
    }


def _body(**kw):
    payload = {
        "title": "收集测试",
        "description": "说明",
        "anonymous": False,
        "start_at": _stamp(-1),
        "end_at": _stamp(3),
        "notify_on_start": True,
        "submit_quota": 1,
        "max_responses": 0,
        "targets": [VOTER],
        "questions": [_q("short_text", "姓名")],
    }
    payload.update(kw)
    return payload


def _create(b, **kw):
    r = b.client.post("/api/collections", json=_body(**kw))
    assert r.status_code < 300, r.text
    return r.json()


def _submit(b, collection_id, answers):
    return b.client.post(f"/api/collections/{collection_id}/responses", json={"answers": answers})


def _poll_body(**kw):
    payload = {
        "title": "投票测试",
        "description": "",
        "anonymous": False,
        "start_at": _stamp(1),
        "end_at": _stamp(3),
        "notify_on_start": True,
        "targets": [VOTER],
        "questions": [{
            "title": "选一个", "multiple": False, "max_choices": 1,
            "options": ["甲", "乙"],
        }],
    }
    payload.update(kw)
    return payload


def _tick(b):
    """在测试进程里对测试库跑一轮开始提醒（后台线程被关掉了）"""
    from app.services import start_notify_service as sns

    engine = create_engine(f"sqlite:///{b.db}")
    Session = sessionmaker(bind=engine)
    db = Session()
    try:
        return sns.tick(now=datetime.now(), db=db)
    finally:
        db.close()


def _inbox_count(b, student_no):
    return b.rows(
        "SELECT COUNT(*) FROM inbox_messages WHERE student_no = ?", student_no
    )[0][0]


# ---------- 六种题型 ----------

def test_create_submit_all_question_types_and_export(cb):
    collection = _create(
        cb,
        questions=[
            _q("single", "单选", options=["甲", "乙"]),
            _q("multiple", "多选", options=["甲", "乙", "丙"], max_choices=2),
            _q("short_text", "短文本"),
            _q("long_text", "长文本", required=False),
            _q("image", "图片", required=False),
            _q("rating", "评分"),
        ],
    )
    cid = collection["id"]
    assert collection["share_token"].startswith(f"{cid}.")
    assert collection["scope_label"] == "定向 1 人"
    assert collection["submit_quota"] == 1

    # 游客看公开信息：有标题/范围/份数，没有题目
    cb.logout()
    info = cb.client.get(f"/api/collections/share/{collection['share_token']}/info")
    assert info.status_code == 200, info.text
    data = info.json()
    assert data["title"] == "收集测试" and data["in_scope"] is False
    assert data["questions"] == []
    assert data["submit_quota"] == 1
    assert data["start_at"] and data["end_at"]

    # 范围内的人上传图片、提交六种答案
    _login(cb, VOTER, VOTER_PW)
    upload = cb.client.post(
        f"/api/collections/{cid}/files",
        files={"file": ("证明.png", PNG, "image/png")},
    )
    assert upload.status_code == 200, upload.text
    token = upload.json()["token"]

    r = _submit(cb, cid, [
        {"question_id": collection["questions"][0]["id"], "choices": [1]},
        {"question_id": collection["questions"][1]["id"], "choices": [0, 2]},
        {"question_id": collection["questions"][2]["id"], "text": "张三"},
        {"question_id": collection["questions"][3]["id"], "text": "一段长文本"},
        {"question_id": collection["questions"][4]["id"], "files": [token]},
        {"question_id": collection["questions"][5]["id"], "rating": 4},
    ])
    assert r.status_code == 200, r.text
    assert r.json()["my_submitted_count"] == 1
    assert r.json()["can_submit"] is False

    # 提交后公开信息里我的进度正确，题目仍只在范围内下发
    data = cb.client.get(f"/api/collections/share/{collection['share_token']}/info").json()
    assert data["my_submitted_count"] == 1 and data["quota_reached"] is True
    assert len(data["questions"]) == 6

    # 发起人看答卷：按提交时间排序、姓名可读、答案齐全
    _login(cb, CREATOR, CREATOR_PW)
    responses = cb.client.get(f"/api/collections/{cid}/responses").json()
    assert responses["total"] == 1
    item = responses["items"][0]
    assert item["student_no"] == VOTER
    assert len(item["answers"]) == 6
    image_answer = next(a for a in item["answers"] if a["qtype"] == "image")
    assert image_answer["files"][0]["original"] == "证明.png"
    file = cb.client.get(
        f"/api/collections/{cid}/responses/{item['id']}/files/0"
    )
    assert file.status_code == 200
    assert file.headers["content-type"].startswith("image/")

    # 导出：一份答卷一行，表头有姓名/学号和各题
    exported = cb.client.get(f"/api/collections/{cid}/responses/export")
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(BytesIO(exported.content))
    sheet = workbook["答卷"]
    headers = [cell.value for cell in sheet[1]]
    assert "姓名" in headers and "学号" in headers
    assert headers[4:] == ["单选", "多选", "短文本", "长文本", "图片", "评分"]
    values = [cell.value for cell in sheet[2]]
    assert values[4] == "B. 乙"
    assert values[5] == "A. 甲；C. 丙"
    assert values[6] == "张三" and values[7] == "一段长文本"
    assert values[9] == 4


# ---------- 配额 / 竞争 / 删除权限 ----------

def test_per_person_quota(cb):
    collection = _create(cb, submit_quota=2)
    cid = collection["id"]
    qid = collection["questions"][0]["id"]
    _login(cb, VOTER, VOTER_PW)
    assert _submit(cb, cid, [{"question_id": qid, "text": "第一份"}]).status_code == 200
    assert _submit(cb, cid, [{"question_id": qid, "text": "第二份"}]).status_code == 200
    r = _submit(cb, cid, [{"question_id": qid, "text": "第三份"}])
    assert r.status_code == 400 and "最多填写 2 份" in r.json()["detail"]

    _login(cb, CREATOR, CREATOR_PW)
    row = cb.client.get("/api/collections").json()["items"][0]
    assert row["submitted_count"] == 2
    assert row["completed_count"] == 1
    assert row["my_submitted_count"] == 0


def test_pending_includes_self_but_remind_skips_self(cb):
    """发起人自己也在范围里时，未填名单要算上他；提醒仍然跳过本人"""
    collection = _create(cb, targets=[CREATOR, VOTER])
    cid = collection["id"]

    pending = cb.client.get(f"/api/collections/{cid}/pending").json()
    assert pending["total"] == 2
    by_no = {row["student_no"]: row for row in pending["items"]}
    assert by_no[CREATOR]["is_self"] is True
    assert by_no[VOTER]["is_self"] is False

    reminded = cb.client.post(f"/api/collections/{cid}/remind")
    assert reminded.status_code == 200, reminded.text
    assert reminded.json()["reminded"] == 1

    # 只有自己一个人在范围里：未填显示 1 人，但催填不会给自己发
    solo = _create(cb, targets=[CREATOR], title="只给自己填")
    solo_pending = cb.client.get(
        f"/api/collections/{solo['id']}/pending"
    ).json()
    assert solo_pending["total"] == 1
    assert solo_pending["items"][0]["is_self"] is True
    assert cb.client.post(
        f"/api/collections/{solo['id']}/remind"
    ).json()["reminded"] == 0


def test_remind_unfilled_and_partial_are_separate(cb):
    """「提醒未填」只催一份没交的；「提醒未填满」只催交过但没填满的"""
    collection = _create(cb, submit_quota=2, targets=[VOTER, OTHER])
    cid = collection["id"]
    qid = collection["questions"][0]["id"]

    _login(cb, VOTER, VOTER_PW)
    assert _submit(cb, cid, [{"question_id": qid, "text": "交了一份"}]).status_code == 200

    _login(cb, CREATOR, CREATOR_PW)
    pending = cb.client.get(f"/api/collections/{cid}/pending").json()
    assert pending["total"] == 2
    assert pending["unfilled_count"] == 1  # OTHER
    assert pending["partial_count"] == 1  # VOTER

    r = cb.client.post(f"/api/collections/{cid}/remind", params={"scope": "unfilled"})
    assert r.status_code == 200, r.text
    assert r.json()["reminded"] == 1
    assert _inbox_count(cb, VOTER) == 0
    assert _inbox_count(cb, OTHER) == 1

    r = cb.client.post(f"/api/collections/{cid}/remind", params={"scope": "partial"})
    assert r.status_code == 200, r.text
    assert r.json()["reminded"] == 1
    assert _inbox_count(cb, VOTER) == 1
    assert _inbox_count(cb, OTHER) == 1

    # 匿名时行里不显示份数，但聚合计数和分组仍然按库里的真实份数算
    anonymous = _create(
        cb,
        anonymous=True,
        submit_quota=2,
        targets=[VOTER, OTHER],
        title="匿名分开提醒",
    )
    _login(cb, VOTER, VOTER_PW)
    assert _submit(
        cb, anonymous["id"], [{"question_id": anonymous["questions"][0]["id"], "text": "x"}]
    ).status_code == 200
    _login(cb, CREATOR, CREATOR_PW)
    rows = cb.client.get(f"/api/collections/{anonymous['id']}/pending").json()
    assert rows["unfilled_count"] == 1 and rows["partial_count"] == 1
    assert all(row["submitted_count"] is None for row in rows["items"])


def test_collection_targets_only_for_managers(cb):
    """复用范围用的名单接口只给发起人/管理员"""
    collection = _create(cb)
    cid = collection["id"]

    own = cb.client.get(f"/api/collections/{cid}/targets")
    assert own.status_code == 200, own.text
    assert [row["student_no"] for row in own.json()["items"]] == [VOTER]

    _login(cb, VOTER, VOTER_PW)
    assert cb.client.get(f"/api/collections/{cid}/targets").status_code == 404

    _login(cb, ADMIN, ADMIN_PW)
    assert cb.client.get(f"/api/collections/{cid}/targets").status_code == 200


def test_competitive_cap_and_delete_releases_slot(cb):
    collection = _create(
        cb,
        max_responses=2, targets=[VOTER, OTHER, THIRD]
    )
    cid = collection["id"]
    qid = collection["questions"][0]["id"]

    _login(cb, VOTER, VOTER_PW)
    assert _submit(cb, cid, [{"question_id": qid, "text": "a"}]).status_code == 200
    _login(cb, OTHER, OTHER_PW)
    assert _submit(cb, cid, [{"question_id": qid, "text": "b"}]).status_code == 200
    _login(cb, THIRD, THIRD_PW)
    r = _submit(cb, cid, [{"question_id": qid, "text": "c"}])
    assert r.status_code == 400 and "名额已满" in r.json()["detail"]

    # 发起人不能删答卷
    _login(cb, CREATOR, CREATOR_PW)
    responses = cb.client.get(f"/api/collections/{cid}/responses").json()
    first_id = responses["items"][0]["id"]
    assert cb.client.delete(f"/api/collections/{cid}/responses/{first_id}").status_code == 403

    # 管理员可以删；删完名额释放，第三个人能补上
    _login(cb, ADMIN, ADMIN_PW)
    deleted = cb.client.delete(f"/api/collections/{cid}/responses/{first_id}")
    assert deleted.status_code == 200, deleted.text
    assert deleted.json()["response_count"] == 1

    _login(cb, THIRD, THIRD_PW)
    assert _submit(cb, cid, [{"question_id": qid, "text": "c"}]).status_code == 200

    # 管理员批量删除
    _login(cb, ADMIN, ADMIN_PW)
    responses = cb.client.get(f"/api/collections/{cid}/responses").json()["items"]
    ids = [item["id"] for item in responses]
    r = cb.client.post(
        f"/api/collections/{cid}/responses/batch-delete",
        json={"response_ids": ids},
    )
    assert r.status_code == 200, r.text
    assert r.json()["deleted"] == len(ids)


def test_required_optional_and_type_validation(cb):
    collection = _create(
        cb,
        questions=[
            _q("single", "单选", options=["甲", "乙"]),
            _q("short_text", "必答文本"),
            _q("rating", "选答评分", required=False),
        ]
    )
    cid = collection["id"]
    q1, q2, q3 = [q["id"] for q in collection["questions"]]

    _login(cb, VOTER, VOTER_PW)
    # 必答题缺失
    r = _submit(cb, cid, [{"question_id": q1, "choices": [0]}])
    assert r.status_code == 400 and "必答题" in r.json()["detail"]
    # 单选只能一项
    r = _submit(cb, cid, [
        {"question_id": q1, "choices": [0, 1]},
        {"question_id": q2, "text": "ok"},
    ])
    assert r.status_code == 400 and "单选题" in r.json()["detail"]
    # 评分越界
    r = _submit(cb, cid, [
        {"question_id": q1, "choices": [0]},
        {"question_id": q2, "text": "ok"},
        {"question_id": q3, "rating": 6},
    ])
    assert r.status_code == 400 and "评分" in r.json()["detail"]
    # 必答齐全、选答不填 → 成功
    r = _submit(cb, cid, [
        {"question_id": q1, "choices": [0]},
        {"question_id": q2, "text": "ok"},
    ])
    assert r.status_code == 200, r.text


def test_list_search_is_independent(cb):
    """收集列表的标题搜索只命中自己这个 Tab 的数据"""
    _create(cb, title="入党推优材料收集")
    _create(cb, title="国庆去向登记")
    r = cb.client.get("/api/collections", params={"keyword": "推优"})
    assert r.status_code == 200, r.text
    titles = [item["title"] for item in r.json()["items"]]
    assert titles == ["入党推优材料收集"]


# ---------- 匿名 ----------

def test_anonymous_hides_identity_but_keeps_pending_names(cb):
    collection = _create(
        cb,
        anonymous=True,
        targets=[VOTER, OTHER],
        questions=[_q("short_text", "备注")],
    )
    cid = collection["id"]
    qid = collection["questions"][0]["id"]
    _login(cb, VOTER, VOTER_PW)
    assert _submit(cb, cid, [{"question_id": qid, "text": "悄悄说"}]).status_code == 200

    _login(cb, CREATOR, CREATOR_PW)
    item = cb.client.get(f"/api/collections/{cid}/responses").json()["items"][0]
    assert item.get("student_no") is None and item.get("student_name") is None
    assert item["answers"][0]["text"] == "悄悄说"

    # 未填名单：匿名也允许知道「谁还没填」，但不给份数
    pending = cb.client.get(f"/api/collections/{cid}/pending").json()
    assert pending["total"] == 1
    assert pending["items"][0]["name"] == "学生04"
    assert pending["items"][0]["submitted_count"] is None
    assert pending["items"][0]["remaining"] is None

    exported = cb.client.get(f"/api/collections/{cid}/responses/export")
    assert exported.status_code == 200
    headers = [cell.value for cell in load_workbook(BytesIO(exported.content))["答卷"][1]]
    assert "姓名" not in headers and "学号" not in headers


# ---------- 开始提醒 ----------

def test_start_notify_follows_time_and_does_not_repeat(cb):
    """没到点不发；改开始时间跟着走；到点只发一次；错过整场不补发；投票也同一套"""
    # 进程停过一整场（开始和截止都过去了）：回来时不补发
    expired = _create(
        cb, title="已经结束的收集", start_at=_stamp(1), end_at=_stamp(3)
    )
    cb.run(
        "UPDATE collections SET start_at = ?, end_at = ? WHERE id = ?",
        (datetime.now() - timedelta(hours=3)).strftime("%Y-%m-%d %H:%M:%S.%f"),
        (datetime.now() - timedelta(hours=2)).strftime("%Y-%m-%d %H:%M:%S.%f"),
        expired["id"],
    )
    _tick(cb)
    assert _inbox_count(cb, VOTER) == 0
    assert cb.rows(
        "SELECT start_notified_at FROM collections WHERE id = ?", expired["id"]
    )[0][0] is None

    collection = _create(
        cb,
        title="需要提醒的收集", start_at=_stamp(1), end_at=_stamp(3)
    )
    cid = collection["id"]
    assert _inbox_count(cb, VOTER) == 0

    # 改开始时间（未开始、未提醒）：提醒跟着新时间
    updated = cb.client.patch(
        f"/api/collections/{cid}",
        json={
            "title": "需要提醒的收集",
            "description": "说明",
            "start_at": _stamp(2),
        },
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["start_notified_at"] is None
    assert _inbox_count(cb, VOTER) == 0

    # 把开始时间拨到过去，手动跑一轮巡检 → 发一次
    cb.run(
        "UPDATE collections SET start_at = ? WHERE id = ?",
        (datetime.now() - timedelta(hours=1)).strftime("%Y-%m-%d %H:%M:%S.%f"),
        cid,
    )
    assert _tick(cb) >= 1
    assert _inbox_count(cb, VOTER) == 1
    _tick(cb)
    assert _inbox_count(cb, VOTER) == 1, "同一个活动不能重复提醒"

    # 投票同样：创建时不发，到点发一次
    r = cb.client.post("/api/polls", json=_poll_body())
    assert r.status_code < 300, r.text
    pid = r.json()["id"]
    assert _inbox_count(cb, VOTER) == 1
    cb.run(
        "UPDATE polls SET start_at = ? WHERE id = ?",
        (datetime.now() - timedelta(hours=1)).strftime("%Y-%m-%d %H:%M:%S.%f"),
        pid,
    )
    _tick(cb)
    assert _inbox_count(cb, VOTER) == 2
    _tick(cb)
    assert _inbox_count(cb, VOTER) == 2


def test_pause_and_resume_blocks_submission(cb):
    """暂停不是结束：暂停期间不能填，恢复后照常；重复暂停/恢复幂等"""
    collection = _create(cb)
    cid = collection["id"]
    qid = collection["questions"][0]["id"]

    assert cb.client.post(f"/api/collections/{cid}/pause").status_code == 200
    assert cb.client.post(f"/api/collections/{cid}/pause").status_code == 200

    _login(cb, VOTER, VOTER_PW)
    info = cb.client.get(
        f"/api/collections/share/{collection['share_token']}/info"
    ).json()
    assert info["status"] == "paused" and info["can_submit"] is False
    r = _submit(cb, cid, [{"question_id": qid, "text": "暂停时提交"}])
    assert r.status_code == 400 and "暂停" in r.json()["detail"]

    _login(cb, CREATOR, CREATOR_PW)
    assert cb.client.post(f"/api/collections/{cid}/resume").status_code == 200
    assert cb.client.post(f"/api/collections/{cid}/resume").status_code == 200

    _login(cb, VOTER, VOTER_PW)
    assert _submit(cb, cid, [{"question_id": qid, "text": "恢复后提交"}]).status_code == 200


# ---------- 每人一份时的回显 / 本人撤回 ----------

def test_quota_one_shows_my_response_and_withdraw(cb):
    """配额 1：答过打开回显自己的答卷；允许撤回后撤回，再打开回到作答表单

    同时验证「本人能看自己答卷里的图片」，以及非本人/非管理者仍是 404。
    """
    collection = _create(
        cb,
        questions=[
            _q("short_text", "备注"),
            _q("image", "证明", required=False),
        ],
    )
    cid = collection["id"]
    text_q, image_q = [question["id"] for question in collection["questions"]]
    assert collection["allow_withdraw"] is False

    _login(cb, VOTER, VOTER_PW)
    upload = cb.client.post(
        f"/api/collections/{cid}/files",
        files={"file": ("证明.png", PNG, "image/png")},
    )
    assert upload.status_code == 200, upload.text
    assert _submit(
        cb,
        cid,
        [
            {"question_id": text_q, "text": "我的答案"},
            {"question_id": image_q, "files": [upload.json()["token"]]},
        ],
    ).status_code == 200

    info = cb.client.get(
        f"/api/collections/share/{collection['share_token']}/info"
    ).json()
    assert info["quota_reached"] is True and info["can_submit"] is False
    my = info["my_responses"][0]
    assert my and my["seq_no"] == 1
    assert my["answers"][0]["text"] == "我的答案"
    assert len(my["answers"][1]["files"]) == 1

    # 本人能读自己答卷的图片；范围外的人读不到
    own = cb.client.get(f"/api/collections/{cid}/responses/{my['id']}/files/0")
    assert own.status_code == 200, own.text
    _login(cb, OTHER, OTHER_PW)
    assert cb.client.get(
        f"/api/collections/{cid}/responses/{my['id']}/files/0"
    ).status_code == 404

    # 不是本人的答卷统一 404（不暴露收集的设置）
    assert cb.client.post(
        f"/api/collections/{cid}/responses/{my['id']}/withdraw"
    ).status_code == 404

    # 没开启「允许撤回」时，本人也撤不了
    _login(cb, VOTER, VOTER_PW)
    assert cb.client.post(
        f"/api/collections/{cid}/responses/{my['id']}/withdraw"
    ).status_code == 400

    # 出卷人打开开关
    _login(cb, CREATOR, CREATOR_PW)
    patched = cb.client.patch(
        f"/api/collections/{cid}",
        json={"title": "收集测试", "description": "说明", "allow_withdraw": True},
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["allow_withdraw"] is True

    _login(cb, VOTER, VOTER_PW)
    withdrawn = cb.client.post(f"/api/collections/{cid}/responses/{my['id']}/withdraw")
    assert withdrawn.status_code == 200, withdrawn.text
    assert withdrawn.json()["response_count"] == 0
    assert withdrawn.json()["my_submitted_count"] == 0
    assert withdrawn.json()["can_submit"] is True

    info = cb.client.get(
        f"/api/collections/share/{collection['share_token']}/info"
    ).json()
    assert info["my_responses"] == [] and info["can_submit"] is True


def test_withdraw_earlier_copy_and_resubmit(cb):
    """撤回较早的一份会留下空号：再交一份不能和现有最大序号撞车"""
    collection = _create(cb, submit_quota=2, allow_withdraw=True)
    cid = collection["id"]
    qid = collection["questions"][0]["id"]

    _login(cb, VOTER, VOTER_PW)
    assert _submit(cb, cid, [{"question_id": qid, "text": "第一份"}]).status_code == 200
    assert _submit(cb, cid, [{"question_id": qid, "text": "第二份"}]).status_code == 200

    info = cb.client.get(
        f"/api/collections/share/{collection['share_token']}/info"
    ).json()
    assert [item["seq_no"] for item in info["my_responses"]] == [1, 2]

    # 撤掉第 1 份，现有最大序号还是 2
    first_id = info["my_responses"][0]["id"]
    assert cb.client.post(
        f"/api/collections/{cid}/responses/{first_id}/withdraw"
    ).status_code == 200

    # 再交一份应拿到第 3 份，而不是和现有的第 2 份撞唯一约束
    again = _submit(cb, cid, [{"question_id": qid, "text": "补交"}])
    assert again.status_code == 200, again.text
    assert again.json()["seq_no"] == 3

    info = cb.client.get(
        f"/api/collections/share/{collection['share_token']}/info"
    ).json()
    assert [item["seq_no"] for item in info["my_responses"]] == [2, 3]


def test_withdraw_releases_competitive_slot(cb):
    """竞争收集：本人撤回后名额释放，后来的人能补上"""
    collection = _create(
        cb,
        max_responses=1,
        targets=[VOTER, OTHER],
        allow_withdraw=True,
    )
    cid = collection["id"]
    qid = collection["questions"][0]["id"]

    _login(cb, VOTER, VOTER_PW)
    first = _submit(cb, cid, [{"question_id": qid, "text": "占位"}])
    assert first.status_code == 200, first.text
    response_id = first.json()["id"]

    _login(cb, OTHER, OTHER_PW)
    full = _submit(cb, cid, [{"question_id": qid, "text": "补位"}])
    assert full.status_code == 400 and "名额已满" in full.json()["detail"]

    # 别人不能撤回我的答卷
    assert cb.client.post(
        f"/api/collections/{cid}/responses/{response_id}/withdraw"
    ).status_code == 404

    _login(cb, VOTER, VOTER_PW)
    assert cb.client.post(
        f"/api/collections/{cid}/responses/{response_id}/withdraw"
    ).status_code == 200

    _login(cb, OTHER, OTHER_PW)
    assert _submit(cb, cid, [{"question_id": qid, "text": "补位"}]).status_code == 200
