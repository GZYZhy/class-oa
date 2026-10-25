"""考勤记录的备注

覆盖：写、改、清掉、超长被拒、重复提交不刷日志、权限、列表带上、导出带上。

两处最容易出错的地方：

1. **空串即清掉**。备注只有这一个接口（PATCH 整条替换），没有单独的「删除备注」。
   谁要是给 remark 加上 `min_length=1` 之类的校验，用户就再也没法把写错的备注去掉 ——
   而他只会看到「保存失败」，不知道是自己没填。

2. **没变就不写日志**。备注是高频修改的东西（一句话措辞来回改），每次保存都记一条
   「A → A」的话，操作日志会被淹掉，真正要查的事反而翻不到。

cd backend && python -m pytest tests/test_attendance.py -q
"""
from io import BytesIO

import httpx
import openpyxl
import pytest

PW = "initpass123"

# 2 个学生：甲（管理员，能点名）、乙（普通学生，只配看）
ROSTER = [
    ("2024000001", "甲", "010001"),
    ("2024000002", "乙", "010002"),
]
ADMIN_NO = ROSTER[0][0]


@pytest.fixture
def b(backend):
    """初始化完毕、管理员已登录的后端"""
    inst = backend(db="attendance.db", BOOTSTRAP_PASSWORD=PW)
    assert inst.login(password=PW).status_code == 200

    for no, name, tail in ROSTER:
        r = inst.client.post("/api/students", json={
            "student_no": no, "name": name,
            "id_card": f"11010120060101{tail}",
        })
        assert r.status_code == 200, r.text

    r = inst.client.post("/api/setup/finish", json={"admin_student_no": ADMIN_NO})
    assert r.status_code == 200, r.text

    inst.logout()
    assert inst.login(ADMIN_NO, ROSTER[0][2]).status_code == 200
    return inst


def as_student(b, idx):
    """以第 idx 个学生身份登录的新客户端"""
    no, _, tail = ROSTER[idx]
    c = httpx.Client(base_url=str(b.client.base_url), timeout=10)
    r = c.post("/api/auth/login", json={"student_no": no, "password": tail})
    assert r.status_code == 200, r.text
    return c


def make_record(b, scope="全班"):
    """上报一次点名，返回记录 id。后面所有备注用例都挂在这条上"""
    r = b.client.post("/api/attendance", json={
        "scope": scope,
        "records": [
            {"student_no": ADMIN_NO, "name": "甲", "status": "present"},
            {"student_no": ROSTER[1][0], "name": "乙", "status": "absent"},
        ],
    })
    assert r.status_code < 300, r.text
    return r.json()["id"]


def remark_logs(b):
    """操作日志里「考勤备注」那几条的 detail"""
    return [r[0] for r in b.rows(
        "SELECT detail FROM operation_logs WHERE action='考勤备注' ORDER BY id"
    )]


# ---------- 写 / 改 / 清 ----------

def test_my_attendance_lists_only_records_with_my_row(b):
    """学生端只看到有自己那一行的点名，状态是点名时记下的"""
    make_record(b)
    c = as_student(b, 1)
    r = c.get("/api/attendance/mine")
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["total"] == 1
    row = data["items"][0]
    assert (row["type"], row["category"], row["scope"]) == ("attendance", "点名", "全班")
    assert (row["status"], row["status_label"]) == ("absent", "无应答")

    # 管理员自己那条是「到」
    row = b.client.get("/api/attendance/mine").json()["items"][0]
    assert row["status_label"] == "到"

    # 另点一次只有管理员在范围里的名，普通学生不该看到
    r = b.client.post("/api/attendance", json={
        "scope": "共青团员",
        "records": [{"student_no": ADMIN_NO, "name": "甲", "status": "present"}],
    })
    assert r.status_code < 300, r.text
    assert c.get("/api/attendance/mine").json()["total"] == 1
    # 「共青团员」在学生端统一叫「团员」
    scopes = {x["scope"] for x in b.client.get("/api/attendance/mine").json()["items"]}
    assert "团员" in scopes


def test_committee_can_view_another_students_attendance(b):
    """学生详情卡片：班委能看某学生的考勤，普通学生不能借这个接口看别人"""
    make_record(b)
    sid = b.rows("SELECT id FROM students WHERE student_no=?", ROSTER[1][0])[0][0]

    r = b.client.get(f"/api/attendance/student/{sid}")
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["total"] == 1
    assert data["items"][0]["status_label"] == "无应答"

    c = as_student(b, 1)
    assert c.get(f"/api/attendance/student/{sid}").status_code == 403
    assert b.client.get("/api/attendance/student/99999").status_code == 404


def test_student_detail_can_list_submitted_cases(b):
    """学生详情卡片：班委能看某学生的办件记录，普通学生不能"""
    from datetime import datetime

    sid = b.rows("SELECT id FROM students WHERE student_no=?", ROSTER[1][0])[0][0]
    b.run(
        "INSERT INTO feedbacks "
        "(student_id, student_no, student_name, title, content, status, created_at) "
        "VALUES (?,?,?,?,?,?,?)",
        sid, ROSTER[1][0], "乙", "宿舍建议", "内容", "unread",
        datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f"),
    )

    r = b.client.get(f"/api/workspace/student/{sid}/submitted")
    assert r.status_code == 200, r.text
    items = r.json()
    assert len(items) == 1
    assert (items[0]["type"], items[0]["title"]) == ("feedback", "宿舍建议")

    c = as_student(b, 1)
    assert c.get(f"/api/workspace/student/{sid}/submitted").status_code == 403
    assert b.client.get("/api/workspace/student/99999/submitted").status_code == 404


def test_writes_and_updates_remark(b):
    """写上、再改成别的，库里和返回值都得跟着变"""
    rid = make_record(b)

    r = b.client.patch(f"/api/attendance/{rid}", json={"remark": "学院临时集合"})
    assert r.status_code == 200, r.text
    assert r.json()["remark"] == "学院临时集合", "PATCH 的返回值要带上新备注"
    assert b.rows("SELECT remark FROM attendance_records WHERE id=?", rid)[0][0] == "学院临时集合"

    r = b.client.patch(f"/api/attendance/{rid}", json={"remark": "学院临时集合，迟到 5 人"})
    assert r.status_code == 200, r.text
    assert b.rows("SELECT remark FROM attendance_records WHERE id=?", rid)[0][0] == "学院临时集合，迟到 5 人"


def test_new_record_starts_with_empty_remark(b):
    """刚点完名的记录没有备注，不是 None —— 前端 `row.remark` 直接当字符串用"""
    rid = make_record(b)
    r = b.client.get(f"/api/attendance/{rid}")
    assert r.status_code == 200, r.text
    assert r.json()["remark"] == ""


def test_empty_remark_clears_it(b):
    """传空串 = 清掉。这是唯一的「删备注」途径，不能被校验挡住"""
    rid = make_record(b)
    b.client.patch(f"/api/attendance/{rid}", json={"remark": "写错了"})

    r = b.client.patch(f"/api/attendance/{rid}", json={"remark": ""})
    assert r.status_code == 200, r.text
    assert r.json()["remark"] == ""
    assert b.rows("SELECT remark FROM attendance_records WHERE id=?", rid)[0][0] == ""


def test_whitespace_only_remark_clears_it(b):
    """只打空格也当清掉：不然存进去一串看不见的空格，页面上显示成「有备注」"""
    rid = make_record(b)
    b.client.patch(f"/api/attendance/{rid}", json={"remark": "先写上"})

    r = b.client.patch(f"/api/attendance/{rid}", json={"remark": "   "})
    assert r.status_code == 200, r.text
    assert b.rows("SELECT remark FROM attendance_records WHERE id=?", rid)[0][0] == ""


def test_surrounding_spaces_are_trimmed(b):
    """首尾空格顺手去掉：从别处复制过来的话很容易带一个"""
    rid = make_record(b)
    r = b.client.patch(f"/api/attendance/{rid}", json={"remark": "  学院集合  "})
    assert r.json()["remark"] == "学院集合"


def test_overlong_remark_is_rejected(b):
    """超过 200 字要被拒，而且不能把已经写好的那条覆盖掉"""
    rid = make_record(b)
    b.client.patch(f"/api/attendance/{rid}", json={"remark": "原有备注"})

    r = b.client.patch(f"/api/attendance/{rid}", json={"remark": "长" * 201})
    assert r.status_code == 422, f"超长备注应当被拒，实际 {r.status_code}"
    assert b.rows("SELECT remark FROM attendance_records WHERE id=?", rid)[0][0] == "原有备注"


# ---------- 日志 ----------

def test_same_remark_writes_no_extra_log(b):
    """没变就不记日志 —— 同一句话保存两遍只该留一条"""
    rid = make_record(b)
    b.client.patch(f"/api/attendance/{rid}", json={"remark": "学院集合"})
    b.client.patch(f"/api/attendance/{rid}", json={"remark": "学院集合"})

    logs = remark_logs(b)
    assert len(logs) == 1, f"重复保存不该多记一条日志：{logs}"


def test_log_keeps_before_and_after(b):
    """日志里要有改动的前后文，否则事后只看到「改过备注」等于没记"""
    rid = make_record(b)
    b.client.patch(f"/api/attendance/{rid}", json={"remark": "学院集合"})
    b.client.patch(f"/api/attendance/{rid}", json={"remark": "学院集合，迟到 5 人"})
    b.client.patch(f"/api/attendance/{rid}", json={"remark": ""})

    logs = remark_logs(b)
    assert len(logs) == 3, logs
    assert "写上" in logs[0] and "学院集合" in logs[0]
    assert "学院集合" in logs[1] and "迟到 5 人" in logs[1], "改的时候要看得见改前是什么"
    assert "清掉" in logs[2], "删备注也要留痕"


def test_remark_log_does_not_claim_a_student(b):
    """备注不属于任何一个学生，不该挂 target_student_nos ——
    挂上去的话，每个人的「被操作记录」里都会凭空多一条和他无关的事
    """
    rid = make_record(b)
    b.client.patch(f"/api/attendance/{rid}", json={"remark": "学院集合"})

    row = b.rows(
        "SELECT target_student_nos FROM operation_logs WHERE action='考勤备注' ORDER BY id DESC LIMIT 1"
    )[0][0]
    assert not row, f"备注日志不该挂具体学生，实际挂了：{row!r}"


# ---------- 权限 ----------

def test_student_cannot_write_remark(b):
    """普通学生改不了（考勤页本来就是班委及以上才看得见）"""
    rid = make_record(b)
    c = as_student(b, 1)
    try:
        r = c.patch(f"/api/attendance/{rid}", json={"remark": "我是学生"})
        assert r.status_code == 403, f"学生不该改得动备注，实际 {r.status_code}"
    finally:
        c.close()
    assert b.rows("SELECT remark FROM attendance_records WHERE id=?", rid)[0][0] == ""


def test_missing_record_is_404(b):
    r = b.client.patch("/api/attendance/999999", json={"remark": "查无此记录"})
    assert r.status_code == 404, r.text


# ---------- 列表与导出 ----------

def test_remark_comes_back_in_list(b):
    """列表接口也要带上：备注就显示在列表那一列上"""
    rid = make_record(b)
    b.client.patch(f"/api/attendance/{rid}", json={"remark": "学院集合"})

    r = b.client.get("/api/attendance", params={"page": 1, "page_size": 20})
    assert r.status_code == 200, r.text
    row = next(x for x in r.json()["items"] if x["id"] == rid)
    assert row["remark"] == "学院集合"


def test_remark_shows_up_in_export(b):
    """导出给人看的那份也得带上 —— 列表里叫标题，导出表头也用同一个名字"""
    rid = make_record(b)
    b.client.patch(f"/api/attendance/{rid}", json={"remark": "学院临时集合"})

    r = b.client.get("/api/attendance/export")
    assert r.status_code == 200, r.text
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    assert "标题" in header, f"导出表头里没有标题列：{header}"
    assert "备注" not in header, f"导出表头还留着旧叫法：{header}"

    col = header.index("标题")
    values = [row[col].value for row in ws.iter_rows(min_row=2)]
    assert set(values) == {"学院临时集合"}, f"明细行上没带标题：{values}"


def test_unwritten_remark_exports_blank(b):
    """没写标题的记录导出时是空的，不能漏出 None 字面量"""
    make_record(b)

    r = b.client.get("/api/attendance/export")
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    col = [c.value for c in ws[1]].index("标题")
    values = [row[col].value for row in ws.iter_rows(min_row=2)]
    assert all(v is None or v == "" for v in values), f"没写备注的那列应当是空的：{values}"
    assert not any(isinstance(v, str) for v in values), "别把 Python 的 None 当字符串打进单元格"


# ---------- 迟到 ----------

def test_late_is_manual_only_and_shows_everywhere(b):
    """「迟到」只能由班委/管理员事后改状态产生，点名上报收不了它

    这条测试同时盯着三个落点：创建接口的闸门、改状态后的计数、学生端和导出
    的中文标签。少一处就会出现「后台看着是迟到，学生看到无应答/导出是 late」。
    """
    # 1) 点名上报不能直接写 late（界面上也没有这个键）
    r = b.client.post("/api/attendance", json={
        "scope": "全班",
        "records": [{"student_no": ADMIN_NO, "name": "甲", "status": "late"}],
    })
    assert r.status_code == 422, f"点名上报不该收迟到，实际 {r.status_code}"

    # 2) 正常点名后，在明细里手动改
    rid = make_record(b)
    r = b.client.patch(f"/api/attendance/{rid}/items/{ADMIN_NO}", json={"status": "late"})
    assert r.status_code == 200, r.text
    data = r.json()
    assert (data["present"], data["late"], data["leave"], data["absent"], data["total"]) == (
        0, 1, 0, 1, 2,
    )

    # 3) 列表与详情一致
    row = next(
        x for x in b.client.get("/api/attendance").json()["items"] if x["id"] == rid
    )
    assert (row["present"], row["late"], row["absent"]) == (0, 1, 1)

    # 4) 学生端标签是「迟到」，不是英文枚举
    mine = b.client.get("/api/attendance/mine").json()["items"]
    assert [x["status_label"] for x in mine] == ["迟到"]

    # 5) 导出状态列也显示中文
    r = b.client.get("/api/attendance/export")
    assert r.status_code == 200, r.text
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    col = header.index("状态")
    values = [row[col].value for row in ws.iter_rows(min_row=2)]
    assert "迟到" in values, f"导出里没有「迟到」：{values}"

    # 6) 操作日志留痕：从「到」改成「迟到」
    detail = b.rows(
        "SELECT detail FROM operation_logs WHERE action='修改考勤' ORDER BY id DESC LIMIT 1"
    )[0][0]
    assert "到 → 迟到" in detail, detail


def test_overview_scope_uses_standard_short_names(b):
    """列表「范围」只出现全班 / 团员 / 班委 / N人，不直接露原始选项"""
    make_record(b, scope="全班")
    ids = {}
    for scope in ("共青团员", "男生", "班委"):
        r = b.client.post("/api/attendance", json={
            "scope": scope,
            "records": [{"student_no": ADMIN_NO, "name": "甲", "status": "present"}],
        })
        assert r.status_code == 200, r.text
        ids[scope] = r.json()["id"]

    items = b.client.get("/api/attendance/overview").json()["items"]
    labels = {i["title"]: i["scope_label"] for i in items if i["type"] == "attendance"}
    assert labels["全班"] == "全班"
    assert labels["共青团员"] == "团员"
    assert labels["班委"] == "班委"
    assert labels["男生"] == "1人"

    # 列表 / 详情 / 导出都用同一套简语
    row = next(
        x for x in b.client.get("/api/attendance").json()["items"]
        if x["id"] == ids["共青团员"]
    )
    assert row["scope_label"] == "团员"
    detail = b.client.get(f"/api/attendance/{ids['共青团员']}").json()
    assert detail["scope_label"] == "团员"

    r = b.client.get("/api/attendance/export")
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    col = header.index("范围")
    values = {row[col].value for row in ws.iter_rows(min_row=2)}
    assert {"全班", "团员", "班委", "1人"} <= values, values

    # 学生端「考勤情况」的参与范围用同一套口径
    mine = b.client.get("/api/attendance/mine").json()["items"]
    scopes = {i["scope"] for i in mine}
    assert "团员" in scopes, scopes
    assert "1人" in scopes, scopes
